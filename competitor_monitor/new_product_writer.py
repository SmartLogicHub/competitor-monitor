from __future__ import annotations

import re
from copy import copy
from dataclasses import dataclass
from typing import Iterable

from openpyxl.worksheet.cell_range import CellRange
from openpyxl.worksheet.worksheet import Worksheet

from new_product_service import NewProduct


@dataclass(frozen=True)
class NewProductLayout:
    start_column: int
    end_column: int
    header_row: int
    data_start_row: int
    date_column: int
    model_column: int
    shape_column: int
    price_column: int
    live_price_columns: list[int]
    selling_point_column: int


@dataclass(frozen=True)
class BrandRange:
    brand: str
    start_row: int
    end_row: int


@dataclass(frozen=True)
class NewProductWriteStats:
    written: int = 0
    skipped_duplicates: int = 0
    no_data_written: int = 0
    failed: int = 0
    failure_details: tuple[str, ...] = ()


def detect_new_product_layout(worksheet: Worksheet) -> NewProductLayout:
    """Locate the right-side 上新监控 region from headers instead of fixed columns."""
    monitor_cell = _find_header(worksheet, "上新监控", max_row=3, min_col=1, max_col=worksheet.max_column)
    min_col = monitor_cell[1] if monitor_cell else 1
    headers = {
        "date_column": _require_header(worksheet, "上架日期", min_col),
        "model_column": _require_header(worksheet, "型号", min_col),
        "shape_column": _require_header(worksheet, "形态", min_col),
        "price_column": _require_header(worksheet, "价格", min_col),
        "live_price": _require_header(worksheet, "直播价格", min_col),
        "selling_point_column": _require_header(worksheet, "新卖点", min_col),
    }
    live_price_columns = _detect_live_price_columns(worksheet, headers["live_price"])
    header_row = max(_header_bottom_row(worksheet, cell) for cell in headers.values())
    end_column = max(headers["selling_point_column"][1], *(live_price_columns or [headers["live_price"][1]]))
    return NewProductLayout(
        start_column=headers["date_column"][1],
        end_column=end_column,
        header_row=header_row,
        data_start_row=header_row + 1,
        date_column=headers["date_column"][1],
        model_column=headers["model_column"][1],
        shape_column=headers["shape_column"][1],
        price_column=headers["price_column"][1],
        live_price_columns=live_price_columns,
        selling_point_column=headers["selling_point_column"][1],
    )


def write_new_products(
    worksheet: Worksheet,
    layout: NewProductLayout,
    products_by_brand: dict[str, list[NewProduct]],
    no_data_text: str = "无上新",
    force_overwrite: bool = False,
) -> NewProductWriteStats:
    written = 0
    skipped_duplicates = 0
    no_data_written = 0
    failed = 0
    failure_details: list[str] = []

    brand_ranges = _detect_brand_ranges(worksheet, layout.data_start_row)
    right_ranges = _detect_right_brand_ranges(worksheet, layout, brand_ranges, no_data_text)
    for incoming_brand, products in products_by_brand.items():
        matched_range = _match_brand_range(incoming_brand, brand_ranges)
        if matched_range is None:
            failed += max(len(products), 1)
            if products:
                for product in products:
                    failure_details.append(f"{incoming_brand} {product.model}: Excel 上新区域未找到品牌行")
            else:
                failure_details.append(f"{incoming_brand}: Excel 上新区域未找到品牌行")
            continue
        brand_range = right_ranges.get(matched_range.brand, matched_range)
        if not products:
            changed = _write_no_new_if_needed(worksheet, layout, brand_range, no_data_text, force_overwrite)
            no_data_written += 1 if changed else 0
            continue
        if _range_has_no_data_marker(worksheet, layout, brand_range, no_data_text) or force_overwrite:
            _clear_right_area(worksheet, layout, brand_range, keep_styles=True)
            write_range = brand_range
        else:
            write_range = _effective_product_range(worksheet, layout, brand_range, no_data_text)
            _unmerge_right_area(worksheet, layout, write_range, keep_styles=True)
        existing_keys = _existing_product_keys(worksheet, layout, write_range)
        for product in products:
            product_keys = _product_keys(product)
            if product_keys and product_keys.intersection(existing_keys):
                skipped_duplicates += 1
                continue
            row = _first_empty_product_row(worksheet, layout, write_range)
            appended_row = False
            if row is None:
                write_range = _append_empty_row_to_brand_range(worksheet, layout, write_range, no_data_text)
                right_ranges[matched_range.brand] = write_range
                appended_row = True
                row = _first_empty_product_row(worksheet, layout, write_range)
                if row is None:
                    failed += 1
                    failure_details.append(
                        f"{incoming_brand} {product.model}: 品牌区域 {write_range.start_row}-{write_range.end_row} 没有空行"
                    )
                    continue
            _write_product_row(worksheet, layout, row, product)
            if appended_row:
                right_ranges = _detect_right_brand_ranges(worksheet, layout, brand_ranges, no_data_text)
            written += 1
            existing_keys.update(product_keys)

    return NewProductWriteStats(
        written=written,
        skipped_duplicates=skipped_duplicates,
        no_data_written=no_data_written,
        failed=failed,
        failure_details=tuple(failure_details),
    )


def list_new_product_brands(worksheet: Worksheet, layout: NewProductLayout) -> list[str]:
    return [brand_range.brand for brand_range in _detect_brand_ranges(worksheet, layout.data_start_row)]


def match_new_product_brand(
    brand: str,
    excel_brands: Iterable[str],
    brand_shop_aliases: dict[str, list[str]] | None = None,
) -> str | None:
    brands = list(excel_brands)
    aliases = brand_shop_aliases or {}
    for excel_brand in brands:
        alias_values = aliases.get(excel_brand) or []
        if _matches_alias(brand, alias_values):
            return excel_brand
    strict_brands = {brand_name for brand_name, alias_values in aliases.items() if alias_values}
    return _best_text_match(brand, [excel_brand for excel_brand in brands if excel_brand not in strict_brands])


def _write_product_row(worksheet: Worksheet, layout: NewProductLayout, row: int, product: NewProduct) -> None:
    worksheet.cell(row=row, column=layout.date_column).value = f"{product.on_sale_date.month}月{product.on_sale_date.day}"
    model_cell = worksheet.cell(row=row, column=layout.model_column)
    model_cell.value = product.model
    if product.link:
        model_cell.hyperlink = product.link
        _apply_hyperlink_font(model_cell)
        _ensure_model_cell_table_border(worksheet, layout, row, model_cell)
    worksheet.cell(row=row, column=layout.shape_column).value = product.shape
    worksheet.cell(row=row, column=layout.price_column).value = product.price
    for column in layout.live_price_columns:
        worksheet.cell(row=row, column=column).value = "/"
    worksheet.cell(row=row, column=layout.selling_point_column).value = product.selling_point or "无"


def _write_no_new_if_needed(
    worksheet: Worksheet,
    layout: NewProductLayout,
    brand_range: BrandRange,
    no_data_text: str,
    force_overwrite: bool,
) -> bool:
    if _range_has_any_product_data(worksheet, layout, brand_range) and not force_overwrite:
        return False
    if _range_has_no_data_marker(worksheet, layout, brand_range, no_data_text):
        return False
    _clear_right_area(worksheet, layout, brand_range, keep_styles=True)
    worksheet.merge_cells(
        start_row=brand_range.start_row,
        start_column=layout.start_column,
        end_row=brand_range.end_row,
        end_column=layout.end_column,
    )
    worksheet.cell(row=brand_range.start_row, column=layout.start_column).value = no_data_text
    return True


def _clear_right_area(worksheet: Worksheet, layout: NewProductLayout, brand_range: BrandRange, keep_styles: bool) -> None:
    _unmerge_right_area(worksheet, layout, brand_range, keep_styles=keep_styles)
    for row in range(brand_range.start_row, brand_range.end_row + 1):
        for column in range(layout.start_column, layout.end_column + 1):
            cell = worksheet.cell(row=row, column=column)
            cell.value = None
            cell.hyperlink = None


def _unmerge_right_area(worksheet: Worksheet, layout: NewProductLayout, brand_range: BrandRange, keep_styles: bool) -> None:
    style = _style_snapshot(worksheet.cell(row=brand_range.start_row, column=layout.start_column))
    ranges = [
        merged_range
        for merged_range in list(worksheet.merged_cells.ranges)
        if _ranges_intersect(
            merged_range,
            brand_range.start_row,
            layout.start_column,
            brand_range.end_row,
            layout.end_column,
        )
    ]
    for merged_range in ranges:
        worksheet.unmerge_cells(str(merged_range))
    if keep_styles and ranges:
        for row in range(brand_range.start_row, brand_range.end_row + 1):
            for column in range(layout.start_column, layout.end_column + 1):
                _apply_style_snapshot(worksheet.cell(row=row, column=column), style)


def _detect_brand_ranges(worksheet: Worksheet, data_start_row: int) -> list[BrandRange]:
    ranges: list[BrandRange] = []
    covered_rows: set[int] = set()
    for merged_range in worksheet.merged_cells.ranges:
        if merged_range.min_col <= 1 <= merged_range.max_col and merged_range.max_row >= data_start_row:
            brand = _cell_text(worksheet.cell(row=merged_range.min_row, column=1))
            if brand:
                start_row = max(merged_range.min_row, data_start_row)
                end_row = merged_range.max_row
                ranges.append(BrandRange(brand=brand, start_row=start_row, end_row=end_row))
                covered_rows.update(range(start_row, end_row + 1))
    brand_rows = [
        row
        for row in range(data_start_row, worksheet.max_row + 1)
        if row not in covered_rows and _cell_text(worksheet.cell(row=row, column=1))
    ]
    for index, row in enumerate(brand_rows):
        end_row = (brand_rows[index + 1] - 1) if index + 1 < len(brand_rows) else row
        ranges.append(BrandRange(brand=_cell_text(worksheet.cell(row=row, column=1)), start_row=row, end_row=end_row))
    return sorted(ranges, key=lambda item: item.start_row)


def _match_brand_range(brand: str, brand_ranges: Iterable[BrandRange]) -> BrandRange | None:
    ranges = list(brand_ranges)
    matched_brand = _best_text_match(brand, [brand_range.brand for brand_range in ranges])
    if matched_brand is None:
        return None
    for brand_range in ranges:
        if brand_range.brand == matched_brand:
            return brand_range
    return None


def _detect_right_brand_ranges(
    worksheet: Worksheet,
    layout: NewProductLayout,
    brand_ranges: Iterable[BrandRange],
    no_data_text: str,
) -> dict[str, BrandRange]:
    ordered_ranges = sorted(brand_ranges, key=lambda item: item.start_row)
    ranges: dict[str, BrandRange] = {}
    for brand_range in ordered_ranges:
        height = brand_range.end_row - brand_range.start_row + 1
        no_data_range = _find_no_data_merged_range_near_brand(worksheet, layout, brand_range, 0, no_data_text)
        if no_data_range is not None:
            right_range = BrandRange(brand=brand_range.brand, start_row=no_data_range.min_row, end_row=no_data_range.max_row)
        else:
            right_range = _effective_product_range(
                worksheet,
                layout,
                BrandRange(
                    brand=brand_range.brand,
                    start_row=brand_range.start_row,
                    end_row=brand_range.start_row + height - 1,
                ),
                no_data_text,
            )
        ranges[brand_range.brand] = right_range
    return ranges


def _existing_product_keys(worksheet: Worksheet, layout: NewProductLayout, brand_range: BrandRange) -> set[str]:
    keys: set[str] = set()
    for row in range(brand_range.start_row, brand_range.end_row + 1):
        model_cell = worksheet.cell(row=row, column=layout.model_column)
        model = _cell_text(model_cell)
        if model:
            keys.add(f"model:{_normalize_key(model)}")
        if model_cell.hyperlink and model_cell.hyperlink.target:
            keys.add(f"link:{model_cell.hyperlink.target.strip().lower()}")
    return keys


def _product_keys(product: NewProduct) -> set[str]:
    keys: set[str] = set()
    if product.link:
        keys.add(f"link:{product.link.strip().lower()}")
    if product.model:
        keys.add(f"model:{_normalize_key(product.model)}")
    return keys


def _first_empty_product_row(worksheet: Worksheet, layout: NewProductLayout, brand_range: BrandRange) -> int | None:
    for row in range(brand_range.start_row, brand_range.end_row + 1):
        if not _cell_text(worksheet.cell(row=row, column=layout.model_column)):
            return row
    return None


def _effective_product_range(
    worksheet: Worksheet,
    layout: NewProductLayout,
    brand_range: BrandRange,
    no_data_text: str,
) -> BrandRange:
    end_row = brand_range.end_row
    row = end_row + 1
    while row <= worksheet.max_row:
        if _row_has_no_data_marker(worksheet, layout, row, no_data_text):
            break
        if not _right_row_has_any_data(worksheet, layout, row):
            break
        end_row = row
        row += 1
    return BrandRange(brand=brand_range.brand, start_row=brand_range.start_row, end_row=end_row)


def _append_empty_row_to_brand_range(
    worksheet: Worksheet,
    layout: NewProductLayout,
    brand_range: BrandRange,
    no_data_text: str,
) -> BrandRange:
    insert_row = brand_range.end_row + 1
    source_row = brand_range.end_row

    claimed_from_no_data = _claim_row_from_no_data_marker(worksheet, layout, insert_row, no_data_text)
    if not claimed_from_no_data and _right_row_has_any_data(worksheet, layout, insert_row):
        _shift_right_area_down(worksheet, layout, insert_row)
    for column in range(layout.start_column, layout.end_column + 1):
        source_cell = worksheet.cell(row=source_row, column=column)
        target_cell = worksheet.cell(row=insert_row, column=column)
        _apply_style_snapshot(target_cell, _style_snapshot(source_cell))
        target_cell.value = None
        target_cell.hyperlink = None
    return BrandRange(brand=brand_range.brand, start_row=brand_range.start_row, end_row=insert_row)


def _claim_row_from_no_data_marker(
    worksheet: Worksheet,
    layout: NewProductLayout,
    row: int,
    no_data_text: str,
) -> bool:
    no_data_range = _find_no_data_merged_range_containing_row(worksheet, layout, row, no_data_text)
    if no_data_range is None:
        return False
    style = _style_snapshot(worksheet.cell(row=no_data_range.min_row, column=no_data_range.min_col))
    worksheet.unmerge_cells(str(no_data_range))
    for target_row in range(no_data_range.min_row, no_data_range.max_row + 1):
        for column in range(layout.start_column, layout.end_column + 1):
            cell = worksheet.cell(row=target_row, column=column)
            _apply_style_snapshot(cell, style)
            cell.value = None
            cell.hyperlink = None
    _merge_no_data_remainder(
        worksheet,
        layout,
        no_data_range.min_row,
        row - 1,
        no_data_text,
        style,
    )
    _merge_no_data_remainder(
        worksheet,
        layout,
        row + 1,
        no_data_range.max_row,
        no_data_text,
        style,
    )
    return True


def _merge_no_data_remainder(
    worksheet: Worksheet,
    layout: NewProductLayout,
    start_row: int,
    end_row: int,
    no_data_text: str,
    style: dict,
) -> None:
    if start_row > end_row:
        return
    for row in range(start_row, end_row + 1):
        for column in range(layout.start_column, layout.end_column + 1):
            _apply_style_snapshot(worksheet.cell(row=row, column=column), style)
    worksheet.merge_cells(
        start_row=start_row,
        start_column=layout.start_column,
        end_row=end_row,
        end_column=layout.end_column,
    )
    marker_cell = worksheet.cell(row=start_row, column=layout.start_column)
    _apply_style_snapshot(marker_cell, style)
    marker_cell.value = no_data_text


def _shift_right_area_down(worksheet: Worksheet, layout: NewProductLayout, insert_row: int) -> None:
    shifted_merged_ranges = _remove_and_shift_right_merged_ranges_at_or_below(worksheet, layout, insert_row)
    bottom_row = max(
        [worksheet.max_row]
        + [max_row for _, _, max_row, _ in shifted_merged_ranges]
        + [insert_row]
    )
    for row in range(bottom_row, insert_row - 1, -1):
        for column in range(layout.start_column, layout.end_column + 1):
            _copy_cell(worksheet.cell(row=row, column=column), worksheet.cell(row=row + 1, column=column))
    for column in range(layout.start_column, layout.end_column + 1):
        cell = worksheet.cell(row=insert_row, column=column)
        cell.value = None
        cell.hyperlink = None
    for min_row, min_col, max_row, max_col in shifted_merged_ranges:
        worksheet.merge_cells(start_row=min_row, start_column=min_col, end_row=max_row, end_column=max_col)


def _remove_and_shift_right_merged_ranges_at_or_below(
    worksheet: Worksheet,
    layout: NewProductLayout,
    insert_row: int,
) -> list[tuple[int, int, int, int]]:
    shifted: list[tuple[int, int, int, int]] = []
    for merged_range in list(worksheet.merged_cells.ranges):
        if not (
            merged_range.min_col >= layout.start_column
            and merged_range.max_col <= layout.end_column
            and merged_range.max_row >= insert_row
        ):
            continue
        if merged_range.min_row >= insert_row:
            shifted.append(
                (
                    merged_range.min_row + 1,
                    merged_range.min_col,
                    merged_range.max_row + 1,
                    merged_range.max_col,
                )
            )
            worksheet.unmerge_cells(str(merged_range))
        elif merged_range.min_row < insert_row <= merged_range.max_row:
            shifted.append(
                (
                    merged_range.min_row,
                    merged_range.min_col,
                    merged_range.max_row + 1,
                    merged_range.max_col,
                )
            )
            worksheet.unmerge_cells(str(merged_range))
    return shifted


def _range_has_no_data_marker(
    worksheet: Worksheet,
    layout: NewProductLayout,
    brand_range: BrandRange,
    no_data_text: str,
) -> bool:
    for row in range(brand_range.start_row, brand_range.end_row + 1):
        for column in range(layout.start_column, layout.end_column + 1):
            if _cell_text(worksheet.cell(row=row, column=column)) == no_data_text:
                return True
    return False


def _row_has_no_data_marker(
    worksheet: Worksheet,
    layout: NewProductLayout,
    row: int,
    no_data_text: str,
) -> bool:
    for column in range(layout.start_column, layout.end_column + 1):
        if _cell_text(worksheet.cell(row=row, column=column)) == no_data_text:
            return True
    return False


def _find_no_data_merged_range_near_brand(
    worksheet: Worksheet,
    layout: NewProductLayout,
    brand_range: BrandRange,
    row_offset: int,
    no_data_text: str,
) -> CellRange | None:
    min_row = brand_range.start_row + row_offset
    max_row = brand_range.end_row + row_offset
    for merged_range in worksheet.merged_cells.ranges:
        if not (
            min_row <= merged_range.min_row <= max_row
            and merged_range.min_col >= layout.start_column
            and merged_range.max_col <= layout.end_column
        ):
            continue
        if _cell_text(worksheet.cell(row=merged_range.min_row, column=merged_range.min_col)) == no_data_text:
            return merged_range
    return None


def _find_no_data_merged_range_containing_row(
    worksheet: Worksheet,
    layout: NewProductLayout,
    row: int,
    no_data_text: str,
) -> CellRange | None:
    for merged_range in worksheet.merged_cells.ranges:
        if not (
            merged_range.min_row <= row <= merged_range.max_row
            and merged_range.min_col >= layout.start_column
            and merged_range.max_col <= layout.end_column
        ):
            continue
        if _cell_text(worksheet.cell(row=merged_range.min_row, column=merged_range.min_col)) == no_data_text:
            return merged_range
    return None


def _range_has_any_product_data(worksheet: Worksheet, layout: NewProductLayout, brand_range: BrandRange) -> bool:
    for row in range(brand_range.start_row, brand_range.end_row + 1):
        if _right_row_has_any_data(worksheet, layout, row):
            return True
    return False


def _right_row_has_any_data(worksheet: Worksheet, layout: NewProductLayout, row: int) -> bool:
    for column in range(layout.start_column, layout.end_column + 1):
        if _cell_text(worksheet.cell(row=row, column=column)):
            return True
    return False


def _detect_live_price_columns(worksheet: Worksheet, header_cell: tuple[int, int]) -> list[int]:
    row, column = header_cell
    for merged_range in worksheet.merged_cells.ranges:
        if merged_range.min_row <= row <= merged_range.max_row and merged_range.min_col <= column <= merged_range.max_col:
            return list(range(merged_range.min_col, merged_range.max_col + 1))
    return [column]


def _header_bottom_row(worksheet: Worksheet, header_cell: tuple[int, int]) -> int:
    row, column = header_cell
    for merged_range in worksheet.merged_cells.ranges:
        if merged_range.min_row <= row <= merged_range.max_row and merged_range.min_col <= column <= merged_range.max_col:
            return merged_range.max_row
    return row


def _require_header(worksheet: Worksheet, label: str, min_col: int) -> tuple[int, int]:
    header = _find_header(worksheet, label, max_row=6, min_col=min_col, max_col=worksheet.max_column)
    if header is None:
        raise ValueError(f"{worksheet.title} 未能定位上新监控表头：{label}")
    return header


def _find_header(
    worksheet: Worksheet,
    label: str,
    max_row: int,
    min_col: int,
    max_col: int,
) -> tuple[int, int] | None:
    normalized = _normalize_header(label)
    for row in range(1, min(max_row, worksheet.max_row) + 1):
        for column in range(min_col, min(max_col, worksheet.max_column) + 1):
            if _normalize_header(worksheet.cell(row=row, column=column).value) == normalized:
                return row, column
    return None


def _ranges_intersect(cell_range: CellRange, min_row: int, min_col: int, max_row: int, max_col: int) -> bool:
    return not (
        cell_range.max_row < min_row
        or cell_range.min_row > max_row
        or cell_range.max_col < min_col
        or cell_range.min_col > max_col
    )


def _style_snapshot(cell):
    return {
        "_style": copy(cell._style),
        "font": copy(cell.font),
        "fill": copy(cell.fill),
        "border": copy(cell.border),
        "number_format": cell.number_format,
        "alignment": copy(cell.alignment),
        "protection": copy(cell.protection),
    }


def _apply_style_snapshot(cell, snapshot: dict) -> None:
    cell._style = copy(snapshot["_style"])
    cell.font = copy(snapshot["font"])
    cell.fill = copy(snapshot["fill"])
    cell.border = copy(snapshot["border"])
    cell.number_format = snapshot["number_format"]
    cell.alignment = copy(snapshot["alignment"])
    cell.protection = copy(snapshot["protection"])


def _copy_cell(source_cell, target_cell) -> None:
    _apply_style_snapshot(target_cell, _style_snapshot(source_cell))
    target_cell.value = source_cell.value
    target_cell.hyperlink = copy(source_cell.hyperlink) if source_cell.hyperlink else None


def _apply_hyperlink_font(cell) -> None:
    font = copy(cell.font)
    font.color = "0563C1"
    font.underline = "single"
    cell.font = font


def _ensure_model_cell_table_border(worksheet: Worksheet, layout: NewProductLayout, row: int, model_cell) -> None:
    if _border_has_visible_side(model_cell.border):
        return
    for column in (
        layout.date_column,
        layout.shape_column,
        layout.price_column,
        layout.selling_point_column,
    ):
        border = worksheet.cell(row=row, column=column).border
        if _border_has_visible_side(border):
            model_cell.border = copy(border)
            return


def _border_has_visible_side(border) -> bool:
    return any(
        side.style
        for side in (
            border.left,
            border.right,
            border.top,
            border.bottom,
        )
    )


def _normalize_header(value) -> str:
    return re.sub(r"\s+", "", str(value or ""))


def _normalize_key(value: str) -> str:
    return re.sub(r"\s+", "", str(value or "")).lower()


def _best_text_match(incoming: str, candidates: list[str]) -> str | None:
    incoming_key = _normalize_key(incoming)
    if not incoming_key:
        return None

    direct_matches = [
        candidate
        for candidate in candidates
        if _normalize_key(candidate)
        and (_normalize_key(candidate) == incoming_key or _normalize_key(candidate) in incoming_key or incoming_key in _normalize_key(candidate))
    ]
    if direct_matches:
        return max(direct_matches, key=lambda candidate: len(_normalize_key(candidate)))

    normalized_incoming = _normalize_brand(incoming)
    normalized_matches = [
        candidate
        for candidate in candidates
        if _normalize_brand(candidate)
        and (
            _normalize_brand(candidate) == normalized_incoming
            or _normalize_brand(candidate) in normalized_incoming
            or normalized_incoming in _normalize_brand(candidate)
        )
    ]
    if normalized_matches:
        return max(normalized_matches, key=lambda candidate: len(_normalize_key(candidate)))
    return None


def _matches_alias(incoming: str, aliases: list[str]) -> bool:
    incoming_key = _normalize_key(incoming)
    for alias in aliases:
        alias_key = _normalize_key(alias)
        if alias_key and (alias_key == incoming_key or alias_key in incoming_key or incoming_key in alias_key):
            return True
    return False


def _normalize_brand(value: str) -> str:
    text = _normalize_key(value)
    for token in ("官方", "旗舰店", "专卖店", "天猫", "淘宝", "耳机", "影音", "数码"):
        text = text.replace(token, "")
    return text


def _cell_text(cell) -> str:
    return str(cell.value).strip() if cell.value is not None else ""
