from __future__ import annotations

import ctypes
from ctypes import wintypes
import os
import re
import shutil
import tempfile
import time
import xml.etree.ElementTree as ET
import posixpath
from dataclasses import dataclass
from datetime import date, datetime, timedelta
from pathlib import Path
from typing import Callable
from zipfile import ZIP_DEFLATED, BadZipFile, ZipFile

from openpyxl import load_workbook
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

from date_service import (
    build_sheet_name,
    date_header_label,
    date_header_labels,
    get_work_week,
    is_workday,
    looks_like_period_sheet,
    sheet_name_contains_date,
)


URL_PATTERN = re.compile(r"https?://\S+|www\.\S+", re.IGNORECASE)
SPREADSHEET_NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
OFFICE_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
PACKAGE_REL_NS = "http://schemas.openxmlformats.org/package/2006/relationships"
CONTENT_TYPES_NS = "http://schemas.openxmlformats.org/package/2006/content-types"


class WorkbookLayoutError(RuntimeError):
    pass


class WorkbookLockedError(PermissionError):
    pass


@dataclass(frozen=True)
class MatureProduct:
    row: int
    brand: str | None
    name: str
    url: str


@dataclass(frozen=True)
class SheetLayout:
    model_column: int
    price_column: int
    price_status_column: int
    activity_column: int
    price_header_row: int
    data_start_row: int
    right_start_column: int | None
    date_price_columns: list[int]


@dataclass(frozen=True)
class TemplateDayCompleteness:
    date: date
    status: str
    filled_count: int
    total_count: int
    missing_rows: tuple[int, ...]
    sheet_name: str | None = None
    error: str | None = None

    @property
    def is_complete(self) -> bool:
        return self.status == "complete"


@dataclass(frozen=True)
class TemplateCompleteness:
    sheet_name: str | None
    week_start: date
    week_end: date
    days: tuple[TemplateDayCompleteness, ...]

    @property
    def is_complete(self) -> bool:
        return all(day.is_complete for day in self.days)

    @property
    def missing_dates(self) -> tuple[date, ...]:
        return tuple(day.date for day in self.days if not day.is_complete)

    @property
    def backfillable_dates(self) -> tuple[date, ...]:
        return tuple(
            day.date
            for day in self.days
            if day.status in {"missing_sheet", "missing", "partial", "empty_template"}
        )


def backup_excel(excel_path: Path | str, backup_dir: Path | str) -> Path:
    source = Path(excel_path)
    target_dir = Path(backup_dir)
    target_dir.mkdir(parents=True, exist_ok=True)
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    target = target_dir / f"{source.stem}_{timestamp}{source.suffix}"
    shutil.copy2(source, target)
    return target


class _WindowsExclusiveFile:
    GENERIC_READ = 0x80000000
    GENERIC_WRITE = 0x40000000
    OPEN_EXISTING = 3
    FILE_ATTRIBUTE_NORMAL = 0x80
    INVALID_HANDLE_VALUE = ctypes.c_void_p(-1).value

    def __init__(self, path: Path):
        kernel32 = ctypes.WinDLL("kernel32", use_last_error=True)
        kernel32.CreateFileW.argtypes = [
            wintypes.LPCWSTR,
            wintypes.DWORD,
            wintypes.DWORD,
            wintypes.LPVOID,
            wintypes.DWORD,
            wintypes.DWORD,
            wintypes.HANDLE,
        ]
        kernel32.CreateFileW.restype = wintypes.HANDLE
        kernel32.CloseHandle.argtypes = [wintypes.HANDLE]
        kernel32.CloseHandle.restype = wintypes.BOOL
        handle = kernel32.CreateFileW(
            str(path),
            self.GENERIC_READ | self.GENERIC_WRITE,
            0,
            None,
            self.OPEN_EXISTING,
            self.FILE_ATTRIBUTE_NORMAL,
            None,
        )
        if handle in (self.INVALID_HANDLE_VALUE, -1, None):
            raise PermissionError(ctypes.WinError(ctypes.get_last_error()))
        self._kernel32 = kernel32
        self._handle = handle

    def close(self) -> None:
        handle = getattr(self, "_handle", None)
        if handle not in (None, self.INVALID_HANDLE_VALUE, -1):
            self._kernel32.CloseHandle(handle)
            self._handle = None


SaveFallback = Callable[[Path, Path, Exception], Path]


class ExcelService:
    def __init__(
        self,
        excel_path: Path | str,
        save_fallback: SaveFallback | None = None,
        fallback_replace_attempts: int = 3,
    ):
        self.excel_path = Path(excel_path)
        self._compatible_copy_path: Path | None = None
        self._save_fallback = save_fallback
        self._fallback_replace_attempts = max(1, int(fallback_replace_attempts))
        self._using_fallback_path = False

    def load_workbook(self):
        """Load workbook with a compatibility fallback for WPS-style empty fill nodes."""
        try:
            return load_workbook(self.excel_path)
        except TypeError as exc:
            if "Fill() takes no arguments" not in str(exc) and "expected <class 'openpyxl.styles.fills.Fill'>" not in str(exc):
                raise
            self._compatible_copy_path = self._create_openpyxl_compatible_copy()
            return load_workbook(self._compatible_copy_path)

    @staticmethod
    def assert_workbook_writable(path: Path | str) -> None:
        target = Path(path)
        if not target.exists():
            raise FileNotFoundError(f"当前 Excel 文件不存在：{target}")
        handle = None
        try:
            handle = ExcelService._open_file_exclusive(target)
        except PermissionError as exc:
            raise WorkbookLockedError(
                f"当前 Excel 文件被占用，请关闭 Excel/WPS/预览窗口后再运行：{target}"
            ) from exc
        finally:
            if handle is not None:
                handle.close()

    @staticmethod
    def wait_workbook_writable(
        path: Path | str,
        timeout_seconds: float = 30.0,
        delay_seconds: float = 1.0,
        on_wait=None,
    ) -> None:
        delay = delay_seconds if delay_seconds > 0 else 1.0
        attempts = max(1, int(timeout_seconds / delay) + 1)
        last_error: WorkbookLockedError | None = None
        for attempt_index in range(attempts):
            try:
                ExcelService.assert_workbook_writable(path)
                return
            except WorkbookLockedError as exc:
                last_error = exc
                if attempt_index == attempts - 1:
                    break
                remaining = max(0, int(round(timeout_seconds - ((attempt_index + 1) * delay))))
                if on_wait is not None:
                    on_wait(attempt_index + 1, remaining)
                time.sleep(delay)
        if last_error is not None:
            raise last_error

    @staticmethod
    def _open_file_exclusive(path: Path):
        if os.name == "nt":
            return _WindowsExclusiveFile(path)
        return path.open("r+b")

    def save_workbook(self, workbook, output_path: Path | str | None = None) -> Path:
        requested_target = Path(output_path) if output_path else self.excel_path
        target = self.excel_path if self._using_fallback_path and requested_target != self.excel_path else requested_target
        target.parent.mkdir(parents=True, exist_ok=True)
        suffix = target.suffix or ".xlsx"
        format_state = self._read_workbook_format_state(target) if target.exists() else {}
        tab_color_states = self._read_sheet_tab_color_states(target) if target.exists() else {}
        drawing_state = self._read_drawing_package_state(target) if target.exists() else {}
        temp = tempfile.NamedTemporaryFile(delete=False, suffix=suffix, dir=target.parent)
        temp_path = Path(temp.name)
        temp.close()
        try:
            workbook.save(temp_path)
            self._restore_workbook_format_state(temp_path, format_state)
            self._apply_mature_product_hyperlink_visual_state(temp_path)
            self._restore_sheet_tab_color_states(temp_path, tab_color_states)
            self._restore_drawing_package_state(temp_path, drawing_state)
            replace_kwargs = {}
            if self._save_fallback is not None:
                replace_kwargs["attempts"] = self._fallback_replace_attempts
            self._replace_file_with_retry(temp_path, target, **replace_kwargs)
        except WorkbookLockedError as exc:
            if self._save_fallback is None:
                raise
            fallback_path = Path(self._save_fallback(temp_path, target, exc))
            self.excel_path = fallback_path
            self._using_fallback_path = True
            return fallback_path
        finally:
            if temp_path.exists():
                temp_path.unlink()
        return target

    @staticmethod
    def _replace_file_with_retry(source: Path, target: Path, attempts: int = 30, delay_seconds: float = 1.0) -> None:
        last_error: PermissionError | None = None
        for attempt in range(attempts):
            try:
                source.replace(target)
                return
            except PermissionError as exc:
                last_error = exc
                if attempt == attempts - 1:
                    break
                time.sleep(delay_seconds)
        if last_error is not None:
            raise WorkbookLockedError(
                f"当前 Excel 文件被占用，请关闭 Excel/WPS/预览窗口后再运行：{target}"
            ) from last_error

    def get_or_create_current_sheet(self, workbook, target_date: date) -> Worksheet:
        existing = self.find_sheet_for_date(workbook, target_date)
        if existing is not None:
            return existing
        if not is_workday(target_date):
            raise WorkbookLayoutError(
                f"{target_date.isoformat()} is not a Monday-Friday monitoring date; "
                "refusing to create a new sheet automatically."
            )
        target_name = build_sheet_name(get_work_week(target_date))
        if target_name in workbook.sheetnames:
            return workbook[target_name]
        template = self.find_latest_template_sheet(workbook)
        new_sheet = workbook.copy_worksheet(template)
        new_sheet.title = target_name
        self.prepare_copied_sheet(new_sheet, target_date)
        return new_sheet

    def find_sheet_for_date(self, workbook, target_date: date) -> Worksheet | None:
        for sheet_name in workbook.sheetnames:
            if sheet_name_contains_date(sheet_name, target_date):
                return workbook[sheet_name]
        return None

    def check_weekly_price_completeness(self, workbook, reference_date: date) -> TemplateCompleteness:
        week = get_work_week(reference_date)
        worksheet = self.find_sheet_for_date(workbook, week.monday)
        if worksheet is None:
            days = tuple(
                TemplateDayCompleteness(
                    date=week.monday + timedelta(days=day_offset),
                    status="missing_sheet",
                    filled_count=0,
                    total_count=0,
                    missing_rows=(),
                    sheet_name=None,
                )
                for day_offset in range(5)
            )
            return TemplateCompleteness(
                sheet_name=None,
                week_start=week.monday,
                week_end=week.friday,
                days=days,
            )

        day_states: list[TemplateDayCompleteness] = []
        for day_offset in range(5):
            current_date = week.monday + timedelta(days=day_offset)
            day_states.append(self._check_price_day_completeness(worksheet, current_date))
        return TemplateCompleteness(
            sheet_name=worksheet.title,
            week_start=week.monday,
            week_end=week.friday,
            days=tuple(day_states),
        )

    def _check_price_day_completeness(self, worksheet: Worksheet, target_date: date) -> TemplateDayCompleteness:
        try:
            layout = self.detect_layout(worksheet, target_date)
            products = self.list_mature_products(worksheet, layout.model_column, layout.data_start_row)
        except WorkbookLayoutError as exc:
            return TemplateDayCompleteness(
                date=target_date,
                status="template_error",
                filled_count=0,
                total_count=0,
                missing_rows=(),
                sheet_name=worksheet.title,
                error=str(exc),
            )

        missing_rows: list[int] = []
        for product in products:
            value = worksheet.cell(row=product.row, column=layout.price_column).value
            if self._is_blank_cell_value(value):
                missing_rows.append(product.row)

        total_count = len(products)
        filled_count = total_count - len(missing_rows)
        if total_count == 0:
            status = "empty_template"
        elif filled_count == total_count:
            status = "complete"
        elif filled_count == 0:
            status = "missing"
        else:
            status = "partial"

        return TemplateDayCompleteness(
            date=target_date,
            status=status,
            filled_count=filled_count,
            total_count=total_count,
            missing_rows=tuple(missing_rows),
            sheet_name=worksheet.title,
        )

    @staticmethod
    def _is_blank_cell_value(value) -> bool:
        return value is None or str(value).strip() == ""

    def find_latest_template_sheet(self, workbook) -> Worksheet:
        for worksheet in reversed(workbook.worksheets):
            if looks_like_period_sheet(worksheet.title):
                return worksheet
        raise WorkbookLayoutError("未找到可复制的 5 天周期 sheet")

    def prepare_copied_sheet(self, worksheet: Worksheet, target_date: date) -> None:
        week = get_work_week(target_date)
        right_start = self._find_header(worksheet, "上新监控", max_row=3, max_col=worksheet.max_column)
        right_start_col = right_start[1] if right_start else None
        date_cells = self._find_mature_date_header_cells(worksheet, right_start_col)
        if len(date_cells) < 5:
            raise WorkbookLayoutError(f"{worksheet.title} 未能识别 5 个成熟单品日期表头")
        for (row, col), label in zip(date_cells[:5], date_header_labels(week)):
            worksheet.cell(row=row, column=col).value = label

        layout = self.detect_layout(worksheet, week.monday)
        clear_columns = set(layout.date_price_columns)
        clear_columns.add(layout.price_status_column)
        clear_columns.add(layout.activity_column)
        for row in range(layout.data_start_row, worksheet.max_row + 1):
            for col in clear_columns:
                worksheet.cell(row=row, column=col).value = None
            if layout.right_start_column:
                for col in range(layout.right_start_column, worksheet.max_column + 1):
                    worksheet.cell(row=row, column=col).value = None

    def detect_layout(self, worksheet: Worksheet, target_date: date) -> SheetLayout:
        right_header = self._find_header(worksheet, "上新监控", max_row=3, max_col=worksheet.max_column)
        right_start_col = right_header[1] if right_header else None
        mature_max_col = (right_start_col - 1) if right_start_col else worksheet.max_column
        price_header = self._find_header(worksheet, date_header_label(target_date), max_row=6, max_col=mature_max_col)
        activity_header = self._find_header(worksheet, "活动", max_row=6, max_col=mature_max_col)
        price_status_header = self._find_header(worksheet, "价格情况", max_row=6, max_col=mature_max_col)
        model_header = self._find_header(worksheet, "型号", max_row=6, max_col=mature_max_col)

        missing = []
        if not price_header:
            missing.append(f"当天价格表头 {date_header_label(target_date)}")
        if not activity_header:
            missing.append("活动")
        if not price_status_header:
            missing.append("价格情况")
        if missing:
            raise WorkbookLayoutError(f"{worksheet.title} 未能定位表头：{', '.join(missing)}")

        date_cols = [col for _, col in self._find_mature_date_header_cells(worksheet, right_start_col)]
        data_start = max(price_header[0], activity_header[0], price_status_header[0]) + 1
        return SheetLayout(
            model_column=model_header[1] if model_header else 2,
            price_column=price_header[1],
            price_status_column=price_status_header[1],
            activity_column=activity_header[1],
            price_header_row=price_header[0],
            data_start_row=data_start,
            right_start_column=right_start_col,
            date_price_columns=date_cols,
        )

    def list_mature_products(self, worksheet: Worksheet, model_column: int = 2, start_row: int = 4) -> list[MatureProduct]:
        products: list[MatureProduct] = []
        current_brand: str | None = None
        for row in range(start_row, worksheet.max_row + 1):
            brand_value = worksheet.cell(row=row, column=1).value
            if brand_value:
                current_brand = str(brand_value).strip()
            model_cell = worksheet.cell(row=row, column=model_column)
            name = str(model_cell.value).strip() if model_cell.value is not None else ""
            if not name or name in {"型号", "商品型号"}:
                continue
            url = self._extract_url(model_cell)
            if not url:
                continue
            products.append(MatureProduct(row=row, brand=current_brand, name=name, url=url))
        return products

    def write_product_result(
        self,
        worksheet: Worksheet,
        layout: SheetLayout,
        product: MatureProduct,
        price: str | None,
        activity_value: str,
        force_overwrite: bool = False,
    ) -> bool:
        wrote_price = False
        price_cell = worksheet.cell(row=product.row, column=layout.price_column)
        if price and (force_overwrite or price_cell.value in (None, "")):
            price_cell.value = price
            wrote_price = True
        activity_cell = worksheet.cell(row=product.row, column=layout.activity_column)
        activity_cell.value = activity_value
        return wrote_price

    def _create_openpyxl_compatible_copy(self) -> Path:
        ns_uri = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
        ET.register_namespace("", ns_uri)
        temp = tempfile.NamedTemporaryFile(delete=False, suffix=".xlsx")
        temp.close()
        fixed_path = Path(temp.name)
        with ZipFile(self.excel_path, "r") as source, ZipFile(fixed_path, "w", ZIP_DEFLATED) as target:
            for item in source.infolist():
                data = source.read(item.filename)
                if item.filename == "xl/styles.xml":
                    data = self._fix_empty_fill_nodes(data, ns_uri)
                target.writestr(item, data)
        return fixed_path

    @classmethod
    def _read_sheet_tab_color_states(cls, path: Path) -> dict[str, bytes | None]:
        states: dict[str, bytes | None] = {}
        try:
            with ZipFile(path, "r") as workbook_zip:
                worksheet_paths = cls._worksheet_paths_by_title(workbook_zip)
                for title, member_name in worksheet_paths.items():
                    root = ET.fromstring(workbook_zip.read(member_name))
                    tab_color = root.find(f"{{{SPREADSHEET_NS}}}sheetPr/{{{SPREADSHEET_NS}}}tabColor")
                    states[title] = ET.tostring(tab_color, encoding="utf-8") if tab_color is not None else None
        except (BadZipFile, KeyError, ET.ParseError, FileNotFoundError):
            return {}
        return states

    @classmethod
    def _restore_sheet_tab_color_states(cls, path: Path, states: dict[str, bytes | None]) -> None:
        if not states:
            return
        rewritten_path: Path | None = None
        try:
            with ZipFile(path, "r") as source:
                worksheet_paths = cls._worksheet_paths_by_title(source)
                state_by_member = {
                    worksheet_paths[title]: state
                    for title, state in states.items()
                    if title in worksheet_paths
                }
                if not state_by_member:
                    return
                temp = tempfile.NamedTemporaryFile(delete=False, suffix=path.suffix or ".xlsx", dir=path.parent)
                rewritten_path = Path(temp.name)
                temp.close()
                with ZipFile(rewritten_path, "w", ZIP_DEFLATED) as target:
                    for item in source.infolist():
                        data = source.read(item.filename)
                        if item.filename in state_by_member:
                            data = cls._apply_tab_color_state(data, state_by_member[item.filename])
                        target.writestr(item, data)
            if rewritten_path is not None:
                cls._replace_file_with_retry(rewritten_path, path)
        except (BadZipFile, KeyError, ET.ParseError):
            return
        finally:
            if rewritten_path is not None and rewritten_path.exists():
                rewritten_path.unlink()

    @classmethod
    def _read_drawing_package_state(cls, path: Path) -> dict:
        state: dict = {"members": {}, "worksheets": {}, "content_types": []}
        try:
            with ZipFile(path, "r") as workbook_zip:
                names = set(workbook_zip.namelist())
                for member_name in names:
                    if member_name.startswith(("xl/drawings/", "xl/media/")):
                        state["members"][member_name] = workbook_zip.read(member_name)
                if not state["members"]:
                    return {}
                for _title, worksheet_member in cls._worksheet_paths_by_title(workbook_zip).items():
                    worksheet_root = ET.fromstring(workbook_zip.read(worksheet_member))
                    drawing_nodes = [
                        ET.tostring(child, encoding="utf-8")
                        for child in worksheet_root
                        if child.tag == f"{{{SPREADSHEET_NS}}}drawing"
                    ]
                    rels_member = cls._worksheet_rels_member_name(worksheet_member)
                    drawing_rels = []
                    if rels_member in names:
                        rels_root = ET.fromstring(workbook_zip.read(rels_member))
                        drawing_rels = [
                            dict(rel.attrib)
                            for rel in rels_root
                            if rel.attrib.get("Type") == "http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing"
                        ]
                    if drawing_nodes or drawing_rels:
                        state["worksheets"][worksheet_member] = {
                            "drawings": drawing_nodes,
                            "rels_member": rels_member,
                            "rels": drawing_rels,
                        }
                content_root = ET.fromstring(workbook_zip.read("[Content_Types].xml"))
                drawing_part_names = {f"/{name}" for name in state["members"] if name.startswith("xl/drawings/")}
                media_extensions = {
                    name.rsplit(".", 1)[-1].lower()
                    for name in state["members"]
                    if name.startswith("xl/media/") and "." in name
                }
                state["content_types"] = [
                    ET.tostring(child, encoding="utf-8")
                    for child in content_root
                    if (
                        child.tag == f"{{{CONTENT_TYPES_NS}}}Override"
                        and child.attrib.get("PartName") in drawing_part_names
                    )
                    or (
                        child.tag == f"{{{CONTENT_TYPES_NS}}}Default"
                        and child.attrib.get("Extension", "").lower() in media_extensions
                    )
                ]
        except (BadZipFile, KeyError, ET.ParseError, FileNotFoundError):
            return {}
        return state if state["members"] else {}

    @classmethod
    def _restore_drawing_package_state(cls, path: Path, state: dict) -> None:
        if not state:
            return
        rewritten_path: Path | None = None
        try:
            with ZipFile(path, "r") as source:
                source_names = set(source.namelist())
                temp = tempfile.NamedTemporaryFile(delete=False, suffix=path.suffix or ".xlsx", dir=path.parent)
                rewritten_path = Path(temp.name)
                temp.close()
                written_names: set[str] = set()
                with ZipFile(rewritten_path, "w", ZIP_DEFLATED) as target:
                    for item in source.infolist():
                        data = source.read(item.filename)
                        if item.filename == "[Content_Types].xml":
                            data = cls._merge_content_type_elements(data, state.get("content_types", []))
                        if item.filename in state.get("worksheets", {}):
                            data = cls._merge_worksheet_drawing_elements(data, state["worksheets"][item.filename].get("drawings", []))
                        for worksheet_state in state.get("worksheets", {}).values():
                            if item.filename == worksheet_state.get("rels_member"):
                                data = cls._merge_relationship_elements(data, worksheet_state.get("rels", []))
                        target.writestr(item, data)
                        written_names.add(item.filename)
                    for worksheet_state in state.get("worksheets", {}).values():
                        rels_member = worksheet_state.get("rels_member")
                        if rels_member and rels_member not in written_names and worksheet_state.get("rels"):
                            target.writestr(rels_member, cls._build_relationships_xml(worksheet_state["rels"]))
                            written_names.add(rels_member)
                    for member_name, member_data in state.get("members", {}).items():
                        if member_name not in source_names and member_name not in written_names:
                            target.writestr(member_name, member_data)
                            written_names.add(member_name)
            if rewritten_path is not None:
                cls._replace_file_with_retry(rewritten_path, path)
        except (BadZipFile, KeyError, ET.ParseError):
            return
        finally:
            if rewritten_path is not None and rewritten_path.exists():
                rewritten_path.unlink()

    @staticmethod
    def _worksheet_rels_member_name(worksheet_member: str) -> str:
        directory = posixpath.dirname(worksheet_member)
        filename = posixpath.basename(worksheet_member)
        return posixpath.join(directory, "_rels", f"{filename}.rels")

    @staticmethod
    def _merge_content_type_elements(content_types_xml: bytes, elements: list[bytes]) -> bytes:
        if not elements:
            return content_types_xml
        root = ET.fromstring(content_types_xml)
        existing_defaults = {
            child.attrib.get("Extension")
            for child in root
            if child.tag == f"{{{CONTENT_TYPES_NS}}}Default"
        }
        existing_overrides = {
            child.attrib.get("PartName")
            for child in root
            if child.tag == f"{{{CONTENT_TYPES_NS}}}Override"
        }
        for element_xml in elements:
            element = ET.fromstring(element_xml)
            if element.tag == f"{{{CONTENT_TYPES_NS}}}Default":
                extension = element.attrib.get("Extension")
                if extension not in existing_defaults:
                    root.append(element)
                    existing_defaults.add(extension)
            elif element.tag == f"{{{CONTENT_TYPES_NS}}}Override":
                part_name = element.attrib.get("PartName")
                if part_name not in existing_overrides:
                    root.append(element)
                    existing_overrides.add(part_name)
        return ET.tostring(root, encoding="utf-8", xml_declaration=True)

    @staticmethod
    def _merge_worksheet_drawing_elements(worksheet_xml: bytes, drawing_elements: list[bytes]) -> bytes:
        if not drawing_elements:
            return worksheet_xml
        root = ET.fromstring(worksheet_xml)
        existing_ids = {
            child.attrib.get(f"{{{OFFICE_REL_NS}}}id")
            for child in root
            if child.tag == f"{{{SPREADSHEET_NS}}}drawing"
        }
        changed = False
        for drawing_xml in drawing_elements:
            drawing = ET.fromstring(drawing_xml)
            drawing_id = drawing.attrib.get(f"{{{OFFICE_REL_NS}}}id")
            if drawing_id not in existing_ids:
                root.append(drawing)
                existing_ids.add(drawing_id)
                changed = True
        return ET.tostring(root, encoding="utf-8", xml_declaration=True) if changed else worksheet_xml

    @staticmethod
    def _merge_relationship_elements(relationships_xml: bytes, relationships: list[dict]) -> bytes:
        if not relationships:
            return relationships_xml
        root = ET.fromstring(relationships_xml)
        existing = {
            (rel.attrib.get("Id"), rel.attrib.get("Type"), rel.attrib.get("Target"))
            for rel in root
        }
        existing_ids = {rel.attrib.get("Id") for rel in root}
        for relationship in relationships:
            key = (relationship.get("Id"), relationship.get("Type"), relationship.get("Target"))
            if key in existing:
                continue
            if relationship.get("Id") in existing_ids:
                continue
            rel = ET.Element(f"{{{PACKAGE_REL_NS}}}Relationship")
            rel.attrib.update(relationship)
            root.append(rel)
            existing.add(key)
            existing_ids.add(relationship.get("Id"))
        return ET.tostring(root, encoding="utf-8", xml_declaration=True)

    @staticmethod
    def _build_relationships_xml(relationships: list[dict]) -> bytes:
        root = ET.Element(f"{{{PACKAGE_REL_NS}}}Relationships")
        for relationship in relationships:
            rel = ET.Element(f"{{{PACKAGE_REL_NS}}}Relationship")
            rel.attrib.update(relationship)
            root.append(rel)
        return ET.tostring(root, encoding="utf-8", xml_declaration=True)

    @classmethod
    def _read_workbook_format_state(cls, path: Path) -> dict:
        state: dict = {"styles": None, "worksheets": {}}
        try:
            with ZipFile(path, "r") as workbook_zip:
                if "xl/styles.xml" in workbook_zip.namelist():
                    state["styles"] = workbook_zip.read("xl/styles.xml")
                for title, member_name in cls._worksheet_paths_by_title(workbook_zip).items():
                    if member_name in workbook_zip.namelist():
                        state["worksheets"][title] = cls._extract_worksheet_format_state(
                            workbook_zip.read(member_name)
                        )
        except (BadZipFile, KeyError, ET.ParseError, FileNotFoundError):
            return {}
        return state

    @classmethod
    def _apply_mature_product_hyperlink_visual_state(cls, path: Path) -> None:
        rewritten_path: Path | None = None
        try:
            with ZipFile(path, "r") as source:
                names = set(source.namelist())
                if "xl/styles.xml" not in names:
                    return
                worksheet_paths = cls._worksheet_paths_by_title(source)
                if not worksheet_paths:
                    return

                styles_root = ET.fromstring(source.read("xl/styles.xml"))
                style_cache: dict[str, str] = {}
                rewritten_worksheets: dict[str, bytes] = {}
                for member_name in worksheet_paths.values():
                    if member_name not in names:
                        continue
                    root = ET.fromstring(source.read(member_name))
                    hyperlink_refs = cls._mature_product_hyperlink_refs(root)
                    if not hyperlink_refs:
                        continue
                    if cls._apply_hyperlink_visual_styles(root, styles_root, style_cache, hyperlink_refs):
                        rewritten_worksheets[member_name] = ET.tostring(root, encoding="utf-8", xml_declaration=True)

                if not rewritten_worksheets:
                    return

                updated_styles = ET.tostring(styles_root, encoding="utf-8", xml_declaration=True)
                temp = tempfile.NamedTemporaryFile(delete=False, suffix=path.suffix or ".xlsx", dir=path.parent)
                rewritten_path = Path(temp.name)
                temp.close()
                with ZipFile(rewritten_path, "w", ZIP_DEFLATED) as target:
                    for item in source.infolist():
                        data = source.read(item.filename)
                        if item.filename == "xl/styles.xml":
                            data = updated_styles
                        elif item.filename in rewritten_worksheets:
                            data = rewritten_worksheets[item.filename]
                        target.writestr(item, data)
            if rewritten_path is not None:
                cls._replace_file_with_retry(rewritten_path, path)
        except (BadZipFile, KeyError, ET.ParseError):
            return
        finally:
            if rewritten_path is not None and rewritten_path.exists():
                rewritten_path.unlink()

    @classmethod
    def _mature_product_hyperlink_refs(cls, worksheet_root: ET.Element) -> set[str]:
        refs: set[str] = set()
        hyperlink_parent = worksheet_root.find(f"{{{SPREADSHEET_NS}}}hyperlinks")
        if hyperlink_parent is None:
            return refs
        for hyperlink in hyperlink_parent.findall(f"{{{SPREADSHEET_NS}}}hyperlink"):
            refs.update(cls._expand_b_column_hyperlink_ref(hyperlink.attrib.get("ref", "")))
        return refs

    @staticmethod
    def _expand_b_column_hyperlink_ref(ref: str) -> set[str]:
        single = re.fullmatch(r"B(\d+)", ref or "")
        if single:
            row = int(single.group(1))
            return {ref} if row >= 4 else set()
        ranged = re.fullmatch(r"B(\d+):B(\d+)", ref or "")
        if not ranged:
            return set()
        start, end = int(ranged.group(1)), int(ranged.group(2))
        if start > end:
            start, end = end, start
        return {f"B{row}" for row in range(max(start, 4), end + 1)}

    @classmethod
    def _apply_hyperlink_visual_styles(
        cls,
        worksheet_root: ET.Element,
        styles_root: ET.Element,
        style_cache: dict[str, str],
        refs: set[str],
    ) -> bool:
        changed = False
        for row in worksheet_root.findall(f".//{{{SPREADSHEET_NS}}}sheetData/{{{SPREADSHEET_NS}}}row"):
            for cell in row.findall(f"{{{SPREADSHEET_NS}}}c"):
                ref = cell.attrib.get("r")
                if ref not in refs:
                    continue
                base_style_id = cell.attrib.get("s", "0")
                hyperlink_style_id = cls._hyperlink_visual_style_id(styles_root, style_cache, base_style_id)
                if hyperlink_style_id is None:
                    continue
                if cell.attrib.get("s") != hyperlink_style_id:
                    cell.attrib["s"] = hyperlink_style_id
                    changed = True
        return changed

    @classmethod
    def _hyperlink_visual_style_id(
        cls,
        styles_root: ET.Element,
        style_cache: dict[str, str],
        base_style_id: str,
    ) -> str | None:
        if base_style_id in style_cache:
            return style_cache[base_style_id]
        fonts = styles_root.find(f"{{{SPREADSHEET_NS}}}fonts")
        cell_xfs = styles_root.find(f"{{{SPREADSHEET_NS}}}cellXfs")
        if fonts is None or cell_xfs is None:
            return None
        xfs = list(cell_xfs)
        try:
            base_index = int(base_style_id)
        except ValueError:
            base_index = 0
        if base_index < 0 or base_index >= len(xfs):
            base_index = 0
        base_xf = xfs[base_index]
        font_id = cls._safe_int(base_xf.attrib.get("fontId"), 0)
        font_nodes = list(fonts)
        if font_id < 0 or font_id >= len(font_nodes):
            font_id = 0
        base_font = font_nodes[font_id] if font_nodes else ET.Element(f"{{{SPREADSHEET_NS}}}font")
        if cls._font_has_hyperlink_visual(base_font):
            style_cache[base_style_id] = str(base_index)
            return str(base_index)

        hyperlink_font = ET.fromstring(ET.tostring(base_font))
        for child in list(hyperlink_font):
            local_name = child.tag.rsplit("}", 1)[-1]
            if local_name in {"color", "u"}:
                hyperlink_font.remove(child)
        ET.SubElement(hyperlink_font, f"{{{SPREADSHEET_NS}}}color", {"rgb": "FF0563C1"})
        ET.SubElement(hyperlink_font, f"{{{SPREADSHEET_NS}}}u")
        fonts.append(hyperlink_font)
        new_font_id = len(fonts) - 1
        fonts.attrib["count"] = str(len(fonts))

        hyperlink_xf = ET.fromstring(ET.tostring(base_xf))
        hyperlink_xf.attrib["fontId"] = str(new_font_id)
        hyperlink_xf.attrib["applyFont"] = "1"
        cell_xfs.append(hyperlink_xf)
        new_style_id = len(cell_xfs) - 1
        cell_xfs.attrib["count"] = str(len(cell_xfs))
        style_cache[base_style_id] = str(new_style_id)
        return str(new_style_id)

    @staticmethod
    def _safe_int(value: str | None, default: int) -> int:
        try:
            return int(value) if value is not None else default
        except ValueError:
            return default

    @staticmethod
    def _font_has_hyperlink_visual(font: ET.Element) -> bool:
        underline = font.find(f"{{{SPREADSHEET_NS}}}u")
        if underline is None:
            return False
        underline_value = underline.attrib.get("val", "single")
        if underline_value in {"none", "0", "false"}:
            return False
        color = font.find(f"{{{SPREADSHEET_NS}}}color")
        if color is None:
            return False
        return color.attrib.get("rgb", "").upper().endswith("0563C1")

    @classmethod
    def _restore_workbook_format_state(cls, path: Path, state: dict) -> None:
        if not state:
            return
        rewritten_path: Path | None = None
        try:
            with ZipFile(path, "r") as source:
                worksheet_paths = cls._worksheet_paths_by_title(source)
                worksheet_state_by_member = {
                    worksheet_paths[title]: worksheet_state
                    for title, worksheet_state in state.get("worksheets", {}).items()
                    if title in worksheet_paths
                }
                if not worksheet_state_by_member and not state.get("styles"):
                    return
                preserved_style_refs = cls._preserved_cell_style_refs(source, worksheet_state_by_member)
                can_restore_styles = cls._can_restore_style_table(
                    source,
                    worksheet_state_by_member,
                    state.get("styles"),
                    preserved_style_refs,
                )
                temp = tempfile.NamedTemporaryFile(delete=False, suffix=path.suffix or ".xlsx", dir=path.parent)
                rewritten_path = Path(temp.name)
                temp.close()
                with ZipFile(rewritten_path, "w", ZIP_DEFLATED) as target:
                    for item in source.infolist():
                        data = source.read(item.filename)
                        if item.filename == "xl/styles.xml" and state.get("styles") and can_restore_styles:
                            data = state["styles"]
                        elif item.filename in worksheet_state_by_member:
                            data = cls._apply_worksheet_format_state(
                                data,
                                worksheet_state_by_member[item.filename],
                                restore_cell_styles=can_restore_styles,
                                preserve_style_refs=preserved_style_refs.get(item.filename, set()),
                            )
                        target.writestr(item, data)
            if rewritten_path is not None:
                cls._replace_file_with_retry(rewritten_path, path)
        except (BadZipFile, KeyError, ET.ParseError):
            return
        finally:
            if rewritten_path is not None and rewritten_path.exists():
                rewritten_path.unlink()

    @classmethod
    def _extract_worksheet_format_state(cls, data: bytes) -> dict:
        root = ET.fromstring(data)
        return {
            "sections": {
                local_name: cls._first_child_xml(root, local_name)
                for local_name in ("sheetPr", "sheetViews", "sheetFormatPr", "cols")
            },
            "row_attrs": {
                row.attrib["r"]: dict(row.attrib)
                for row in root.findall(f".//{{{SPREADSHEET_NS}}}sheetData/{{{SPREADSHEET_NS}}}row")
                if "r" in row.attrib
            },
            "cell_styles": {
                cell.attrib["r"]: cell.attrib.get("s")
                for cell in root.findall(f".//{{{SPREADSHEET_NS}}}sheetData/{{{SPREADSHEET_NS}}}row/{{{SPREADSHEET_NS}}}c")
                if "r" in cell.attrib
            },
        }

    @classmethod
    def _preserved_cell_style_refs(
        cls,
        workbook_zip: ZipFile,
        worksheet_state_by_member: dict[str, dict],
    ) -> dict[str, set[str]]:
        preserved: dict[str, set[str]] = {}
        for item in workbook_zip.infolist():
            if not item.filename.startswith("xl/worksheets/") or not item.filename.endswith(".xml"):
                continue
            original_cell_styles = worksheet_state_by_member.get(item.filename, {}).get("cell_styles", {})
            if not original_cell_styles:
                continue
            root = ET.fromstring(workbook_zip.read(item.filename))
            for cell in root.findall(f".//{{{SPREADSHEET_NS}}}sheetData/{{{SPREADSHEET_NS}}}row/{{{SPREADSHEET_NS}}}c"):
                ref = cell.attrib.get("r")
                if ref not in original_cell_styles:
                    continue
                if cell.attrib.get("s") != original_cell_styles[ref]:
                    preserved.setdefault(item.filename, set()).add(ref)
        return preserved

    @classmethod
    def _can_restore_style_table(
        cls,
        workbook_zip: ZipFile,
        worksheet_state_by_member: dict[str, dict],
        styles_xml: bytes | None,
        preserved_style_refs: dict[str, set[str]] | None = None,
    ) -> bool:
        if not styles_xml:
            return False
        preserved_style_refs = preserved_style_refs or {}
        style_count = cls._cell_style_count(styles_xml)
        for item in workbook_zip.infolist():
            if not item.filename.startswith("xl/worksheets/") or not item.filename.endswith(".xml"):
                continue
            covered_cell_refs = worksheet_state_by_member.get(item.filename, {}).get("cell_styles", {})
            member_preserved_refs = preserved_style_refs.get(item.filename, set())
            root = ET.fromstring(workbook_zip.read(item.filename))
            for cell in root.findall(f".//{{{SPREADSHEET_NS}}}sheetData/{{{SPREADSHEET_NS}}}row/{{{SPREADSHEET_NS}}}c"):
                ref = cell.attrib.get("r")
                style_id = cell.attrib.get("s")
                if style_id is None or (ref in covered_cell_refs and ref not in member_preserved_refs):
                    continue
                try:
                    if int(style_id) >= style_count:
                        return False
                except ValueError:
                    return False
        return True

    @staticmethod
    def _cell_style_count(styles_xml: bytes) -> int:
        root = ET.fromstring(styles_xml)
        cell_xfs = root.find(f"{{{SPREADSHEET_NS}}}cellXfs")
        return len(list(cell_xfs)) if cell_xfs is not None else 0

    @staticmethod
    def _first_child_xml(root: ET.Element, local_name: str) -> bytes | None:
        element = root.find(f"{{{SPREADSHEET_NS}}}{local_name}")
        return ET.tostring(element, encoding="utf-8") if element is not None else None

    @classmethod
    def _apply_worksheet_format_state(
        cls,
        data: bytes,
        state: dict,
        restore_cell_styles: bool = True,
        preserve_style_refs: set[str] | None = None,
    ) -> bytes:
        ET.register_namespace("", SPREADSHEET_NS)
        root = ET.fromstring(data)
        preserve_style_refs = preserve_style_refs or set()
        for local_name, section_xml in state.get("sections", {}).items():
            cls._replace_worksheet_section(root, local_name, section_xml)
        row_attrs = state.get("row_attrs", {})
        cell_styles = state.get("cell_styles", {})
        for row in root.findall(f".//{{{SPREADSHEET_NS}}}sheetData/{{{SPREADSHEET_NS}}}row"):
            row_id = row.attrib.get("r")
            if row_id in row_attrs:
                row.attrib.clear()
                row.attrib.update(row_attrs[row_id])
            if not restore_cell_styles:
                continue
            for cell in row.findall(f"{{{SPREADSHEET_NS}}}c"):
                ref = cell.attrib.get("r")
                if ref in preserve_style_refs:
                    continue
                if ref not in cell_styles:
                    continue
                style_id = cell_styles[ref]
                if style_id is None:
                    cell.attrib.pop("s", None)
                else:
                    cell.attrib["s"] = style_id
        if restore_cell_styles:
            cls._restore_missing_style_only_cells(root, row_attrs, cell_styles, preserve_style_refs)
        return ET.tostring(root, encoding="utf-8", xml_declaration=True)

    @classmethod
    def _restore_missing_style_only_cells(
        cls,
        root: ET.Element,
        row_attrs: dict,
        cell_styles: dict[str, str | None],
        preserve_style_refs: set[str] | None = None,
    ) -> None:
        sheet_data = root.find(f"{{{SPREADSHEET_NS}}}sheetData")
        if sheet_data is None:
            return
        preserve_style_refs = preserve_style_refs or set()
        rows = {
            row.attrib.get("r"): row
            for row in sheet_data.findall(f"{{{SPREADSHEET_NS}}}row")
            if row.attrib.get("r")
        }
        existing_refs = {
            cell.attrib.get("r")
            for row in rows.values()
            for cell in row.findall(f"{{{SPREADSHEET_NS}}}c")
            if cell.attrib.get("r")
        }
        for ref, style_id in cell_styles.items():
            if ref in preserve_style_refs:
                continue
            if style_id is None or ref in existing_refs:
                continue
            row_id = cls._row_number_from_cell_ref(ref)
            row = rows.get(row_id)
            if row is None:
                row = ET.Element(f"{{{SPREADSHEET_NS}}}row", row_attrs.get(row_id, {"r": row_id}))
                cls._insert_row_sorted(sheet_data, row)
                rows[row_id] = row
            cell = ET.Element(f"{{{SPREADSHEET_NS}}}c", {"r": ref, "s": style_id})
            cls._insert_cell_sorted(row, cell)
            existing_refs.add(ref)

    @staticmethod
    def _row_number_from_cell_ref(cell_ref: str) -> str:
        match = re.search(r"\d+", cell_ref)
        return match.group(0) if match else "1"

    @staticmethod
    def _column_number_from_cell_ref(cell_ref: str) -> int:
        match = re.match(r"[A-Z]+", cell_ref)
        if not match:
            return 0
        value = 0
        for char in match.group(0):
            value = value * 26 + (ord(char) - 64)
        return value

    @classmethod
    def _insert_row_sorted(cls, sheet_data: ET.Element, row: ET.Element) -> None:
        row_number = int(row.attrib.get("r", "0"))
        for index, existing in enumerate(list(sheet_data)):
            try:
                existing_row_number = int(existing.attrib.get("r", "0"))
            except ValueError:
                continue
            if existing_row_number > row_number:
                sheet_data.insert(index, row)
                return
        sheet_data.append(row)

    @classmethod
    def _insert_cell_sorted(cls, row: ET.Element, cell: ET.Element) -> None:
        column_number = cls._column_number_from_cell_ref(cell.attrib.get("r", ""))
        for index, existing in enumerate(list(row)):
            existing_ref = existing.attrib.get("r", "")
            if cls._column_number_from_cell_ref(existing_ref) > column_number:
                row.insert(index, cell)
                return
        row.append(cell)

    @classmethod
    def _replace_worksheet_section(cls, root: ET.Element, local_name: str, section_xml: bytes | None) -> None:
        tag = f"{{{SPREADSHEET_NS}}}{local_name}"
        for child in list(root):
            if child.tag == tag:
                root.remove(child)
        if section_xml is None:
            return
        element = ET.fromstring(section_xml)
        if local_name == "sheetPr":
            root.insert(0, element)
        elif local_name in {"sheetViews", "sheetFormatPr", "cols"}:
            root.insert(cls._index_before(root, {"sheetData"}), element)
        elif local_name == "mergeCells":
            root.insert(cls._index_after(root, "sheetData"), element)

    @staticmethod
    def _index_before(root: ET.Element, local_names: set[str]) -> int:
        for index, child in enumerate(list(root)):
            if child.tag.rsplit("}", 1)[-1] in local_names:
                return index
        return len(root)

    @staticmethod
    def _index_after(root: ET.Element, local_name: str) -> int:
        for index, child in enumerate(list(root)):
            if child.tag.rsplit("}", 1)[-1] == local_name:
                return index + 1
        return len(root)

    @staticmethod
    def _worksheet_paths_by_title(workbook_zip: ZipFile) -> dict[str, str]:
        workbook_root = ET.fromstring(workbook_zip.read("xl/workbook.xml"))
        rels_root = ET.fromstring(workbook_zip.read("xl/_rels/workbook.xml.rels"))
        rel_targets = {
            rel.attrib["Id"]: ExcelService._normalize_workbook_rel_target(rel.attrib["Target"])
            for rel in rels_root
            if "Id" in rel.attrib and "Target" in rel.attrib
        }
        paths: dict[str, str] = {}
        for sheet in workbook_root.findall(f"{{{SPREADSHEET_NS}}}sheets/{{{SPREADSHEET_NS}}}sheet"):
            title = sheet.attrib.get("name")
            rel_id = sheet.attrib.get(f"{{{OFFICE_REL_NS}}}id")
            if title and rel_id in rel_targets:
                paths[title] = rel_targets[rel_id]
        return paths

    @staticmethod
    def _normalize_workbook_rel_target(target: str) -> str:
        cleaned = target.lstrip("/")
        normalized = cleaned if cleaned.startswith("xl/") else f"xl/{cleaned}"
        return posixpath.normpath(normalized)

    @staticmethod
    def _apply_tab_color_state(data: bytes, tab_color_state: bytes | None) -> bytes:
        ET.register_namespace("", SPREADSHEET_NS)
        root = ET.fromstring(data)
        sheet_pr_tag = f"{{{SPREADSHEET_NS}}}sheetPr"
        tab_color_tag = f"{{{SPREADSHEET_NS}}}tabColor"
        sheet_pr = root.find(sheet_pr_tag)
        if sheet_pr is not None:
            for child in list(sheet_pr):
                if child.tag == tab_color_tag:
                    sheet_pr.remove(child)
        if tab_color_state is not None:
            if sheet_pr is None:
                sheet_pr = ET.Element(sheet_pr_tag)
                root.insert(0, sheet_pr)
            sheet_pr.insert(0, ET.fromstring(tab_color_state))
        return ET.tostring(root, encoding="utf-8", xml_declaration=True)

    @staticmethod
    def _fix_empty_fill_nodes(data: bytes, ns_uri: str) -> bytes:
        root = ET.fromstring(data)
        fills = root.find(f"{{{ns_uri}}}fills")
        if fills is not None:
            for fill in fills.findall(f"{{{ns_uri}}}fill"):
                if len(fill) == 0:
                    ET.SubElement(fill, f"{{{ns_uri}}}patternFill", {"patternType": "none"})
        return ET.tostring(root, encoding="utf-8", xml_declaration=True)

    @staticmethod
    def _extract_url(cell) -> str | None:
        if cell.hyperlink and cell.hyperlink.target:
            return cell.hyperlink.target
        value = str(cell.value or "")
        match = URL_PATTERN.search(value)
        return match.group(0) if match else None

    @staticmethod
    def _normalize_header(value) -> str:
        return re.sub(r"\s+", "", str(value or ""))

    def _find_header(self, worksheet: Worksheet, label: str, max_row: int, max_col: int) -> tuple[int, int] | None:
        normalized = self._normalize_header(label)
        for row in range(1, min(max_row, worksheet.max_row) + 1):
            for col in range(1, min(max_col, worksheet.max_column) + 1):
                if self._normalize_header(worksheet.cell(row=row, column=col).value) == normalized:
                    return row, col
        return None

    def _find_mature_date_header_cells(self, worksheet: Worksheet, right_start_col: int | None) -> list[tuple[int, int]]:
        max_col = (right_start_col - 1) if right_start_col else worksheet.max_column
        cells: list[tuple[int, int]] = []
        for row in range(1, min(6, worksheet.max_row) + 1):
            row_cells: list[tuple[int, int]] = []
            for col in range(1, min(max_col, worksheet.max_column) + 1):
                value = self._normalize_header(worksheet.cell(row=row, column=col).value)
                if re.fullmatch(r"\d{1,2}月\d{1,2}", value):
                    row_cells.append((row, col))
            if len(row_cells) >= 5:
                return row_cells
            cells.extend(row_cells)
        return cells


def column_name(column_index: int) -> str:
    return get_column_letter(column_index)
