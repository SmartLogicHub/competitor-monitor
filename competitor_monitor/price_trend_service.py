from __future__ import annotations

import re
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from typing import Iterable

from openpyxl.worksheet.worksheet import Worksheet

from excel_service import SheetLayout


TREND_STABLE = "价格稳定"
TREND_DOWN = "价格下降"
TREND_UP = "价格上涨"
TREND_DOWN_THEN_UP = "先降后升"
TREND_UP_THEN_DOWN = "先升后降"


@dataclass(frozen=True)
class PriceTrendWriteStats:
    total: int = 0
    written: int = 0
    skipped_existing: int = 0
    skipped_incomplete: int = 0
    skipped_unclassified: int = 0


def extract_lowest_price(value) -> Decimal | None:
    if value is None:
        return None
    if isinstance(value, Decimal):
        return value
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    text = str(value).strip()
    if not text or text == "/":
        return None
    prices: list[Decimal] = []
    for match in re.finditer(r"(?<![A-Za-z0-9])(\d+(?:\.\d+)?)(?![A-Za-z0-9])", text):
        try:
            prices.append(Decimal(match.group(1)))
        except InvalidOperation:
            continue
    return min(prices) if prices else None


def classify_price_trend(values: Iterable[object]) -> str | None:
    prices = _extract_trend_prices(values)
    if prices is None:
        return None

    compressed = [prices[0]]
    for price in prices[1:]:
        if price != compressed[-1]:
            compressed.append(price)
    if len(compressed) == 1:
        return TREND_STABLE

    signs = [1 if right > left else -1 for left, right in zip(compressed, compressed[1:])]
    if all(sign > 0 for sign in signs):
        return TREND_UP
    if all(sign < 0 for sign in signs):
        return TREND_DOWN

    first_positive = next((index for index, sign in enumerate(signs) if sign > 0), None)
    first_negative = next((index for index, sign in enumerate(signs) if sign < 0), None)
    if first_positive is not None and all(sign < 0 for sign in signs[:first_positive]) and all(sign > 0 for sign in signs[first_positive:]):
        return TREND_DOWN_THEN_UP
    if first_negative is not None and all(sign > 0 for sign in signs[:first_negative]) and all(sign < 0 for sign in signs[first_negative:]):
        return TREND_UP_THEN_DOWN
    return _describe_direction_sequence(signs)


def update_price_trends(
    worksheet: Worksheet,
    layout: SheetLayout,
    force_overwrite: bool = False,
) -> PriceTrendWriteStats:
    total = 0
    written = 0
    skipped_existing = 0
    skipped_incomplete = 0
    skipped_unclassified = 0
    date_columns = layout.date_price_columns[:5]

    for row in range(layout.data_start_row, worksheet.max_row + 1):
        if not _cell_text(worksheet.cell(row=row, column=layout.model_column)):
            continue
        total += 1
        status_cell = worksheet.cell(row=row, column=layout.price_status_column)
        if _cell_text(status_cell) and not force_overwrite:
            skipped_existing += 1
            continue
        values = [worksheet.cell(row=row, column=column).value for column in date_columns]
        trend = classify_price_trend(values)
        if trend is None:
            skipped_incomplete += 1
            continue
        status_cell.value = trend
        written += 1

    return PriceTrendWriteStats(
        total=total,
        written=written,
        skipped_existing=skipped_existing,
        skipped_incomplete=skipped_incomplete,
        skipped_unclassified=skipped_unclassified,
    )


def _cell_text(cell) -> str:
    return str(cell.value).strip() if cell.value is not None else ""


def _extract_trend_prices(values: Iterable[object]) -> list[Decimal] | None:
    parsed = [price for price in (extract_lowest_price(value) for value in values) if price is not None]
    if len(parsed) < 2:
        return None
    return parsed


def _describe_direction_sequence(signs: Iterable[int]) -> str:
    words = ["涨" if sign > 0 else "跌" for sign in signs]
    return "、".join(("先" if index == 0 else "再") + word for index, word in enumerate(words))
