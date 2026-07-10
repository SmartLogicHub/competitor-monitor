import csv
import json
import subprocess
import tempfile
import threading
import time
import unittest
from datetime import date
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

from credential_service import load_bi_credentials, load_taobao_credentials, save_account_credentials
from excel_service import WorkbookLockedError
from web_api import RunContext, WebApiService


def create_uploadable_template(path: Path, model: str = "S6S proII") -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "7.6-7.10"
    ws["A1"] = "品牌"
    ws["B1"] = "竞争品牌核心单品监控"
    ws["J1"] = "上新监控"
    ws["B2"] = "型号"
    ws["C2"] = "7月6"
    ws["D2"] = "7月7"
    ws["E2"] = "7月8"
    ws["F2"] = "7月9"
    ws["G2"] = "7月10"
    ws["H2"] = "价格情况"
    ws["I2"] = "活动"
    ws["J2"] = "上架日期"
    ws["K2"] = "型号"
    ws["L2"] = "形态"
    ws["M2"] = "价格"
    ws["S2"] = "新卖点"
    ws["A4"] = "塞那"
    ws["B4"] = model
    ws["B4"].hyperlink = "https://detail.tmall.com/item.htm?id=1"
    wb.save(path)


class WebApiServiceTest(unittest.TestCase):
    def test_idle_status_previews_target_sheet_from_today_not_stale_config_date(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            base_dir = Path(tmpdir)
            config_path = base_dir / "config.yaml"
            config_path.write_text('run_date: "2026-07-06"\n', encoding="utf-8")
            service = WebApiService(
                base_dir=base_dir,
                config_path=config_path,
                today_provider=lambda: date(2026, 7, 14),
            )

            status = service.get_tasks_status()

            self.assertEqual(status["target_sheet_name"], "7.13-7.17")
            self.assertEqual(status["target_period_range"], "2026-07-13 至 2026-07-17")
            self.assertEqual(status["tasks"][0]["template_status"], "7.13-7.17")

    def test_save_config_syncs_windows_tasks_when_schedule_changes(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            base_dir = Path(tmpdir)
            config_path = base_dir / "config.yaml"
            config_path.write_text(
                "\n".join(
                    [
                        'excel_path: "竞品监控.xlsx"',
                        'daily_run_time: "09:30"',
                        "new_product_allowed_weekdays:",
                        "  - 5",
                    ]
                ),
                encoding="utf-8",
            )
            sync_calls = []

            def syncer(config):
                sync_calls.append(dict(config))

            service = WebApiService(base_dir=base_dir, config_path=config_path, scheduler_syncer=syncer)

            service.save_config(
                {
                    "daily_run_time": "11:08",
                    "weekly_new_time": "11:12",
                    "price_trend_time": "18:45",
                    "bi_run_saturday": True,
                    "bi_run_sunday": True,
                }
            )

            self.assertEqual(len(sync_calls), 1)
            self.assertEqual(sync_calls[0]["daily_run_time"], "11:08")
            self.assertEqual(sync_calls[0]["weekly_new_time"], "11:12")
            self.assertEqual(sync_calls[0]["price_trend_time"], "18:45")
            self.assertEqual(sync_calls[0]["new_product_allowed_weekdays"], [5, 6])
            public_config = service.get_config()
            self.assertEqual(public_config["weekly_new_time"], "11:12")
            self.assertEqual(public_config["price_trend_time"], "18:45")
            self.assertIn("系统计划任务已同步", [item["message"] for item in service.get_logs(level="success")])

    def test_save_config_returns_before_slow_schedule_sync_finishes(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            base_dir = Path(tmpdir)
            config_path = base_dir / "config.yaml"
            config_path.write_text('daily_run_time: "09:30"\n', encoding="utf-8")
            release = threading.Event()
            sync_started = threading.Event()

            def slow_syncer(_config):
                sync_started.set()
                release.wait(1)

            service = WebApiService(base_dir=base_dir, config_path=config_path, scheduler_syncer=slow_syncer)

            response = service.save_config({"daily_run_time": "11:08"})

            self.assertEqual(response["message"], "配置已保存")
            self.assertTrue(sync_started.wait(1))
            self.assertFalse(service.wait_for_schedule_sync(timeout=0.01))
            self.assertIn("系统计划任务正在后台同步", [item["message"] for item in service.get_logs()])
            release.set()
            self.assertTrue(service.wait_for_schedule_sync(timeout=1))

    def test_config_masks_and_preserves_sensitive_values(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            base_dir = Path(tmpdir)
            config_path = base_dir / "config.yaml"
            taobao_path = base_dir / "secrets" / "taobao.json"
            bi_path = base_dir / "secrets" / "bi.json"
            save_account_credentials(taobao_path, "test-taobao-user", "old-pass")
            save_account_credentials(bi_path, "test-bi-user", "old-bi-pass")
            config_path.write_text(
                "\n".join(
                    [
                        'excel_path: "竞品监控.xlsx"',
                        f'taobao_credentials_path: "{taobao_path.as_posix()}"',
                        f'bi_credentials_path: "{bi_path.as_posix()}"',
                        'wecom_webhook: "https://qyapi.weixin.qq.com/cgi-bin/webhook/send?key=old"',
                        "wecom_enabled: true",
                        'run_date: "2026-07-06"',
                        "dry_run: true",
                        "test_one: true",
                    ]
                ),
                encoding="utf-8",
            )
            service = WebApiService(base_dir=base_dir, config_path=config_path)

            public_config = service.get_config()

            self.assertEqual(public_config["excel_path"], "竞品监控.xlsx")
            self.assertTrue(public_config["wecom_webhook_set"])
            self.assertNotIn("wecom_webhook", public_config)
            self.assertEqual(public_config["run_date"], "2026-07-06")
            self.assertTrue(public_config["dry_run"])
            self.assertTrue(public_config["test_one"])
            self.assertEqual(public_config["taobao_login_status"], "运行时检测")
            self.assertEqual(public_config["taobao_credential_status"], "已保存凭据")
            self.assertEqual(public_config["taobao_username_masked"], "已保存账号")
            self.assertTrue(public_config["taobao_password_set"])
            self.assertEqual(public_config["bi_login_status"], "运行时检测")
            self.assertEqual(public_config["bi_credential_status"], "已保存凭据")
            self.assertEqual(public_config["bi_username_masked"], "已保存账号")
            self.assertTrue(public_config["bi_password_set"])
            self.assertNotIn("test-taobao-user", json.dumps(public_config, ensure_ascii=False))
            self.assertNotIn("test-bi-user", json.dumps(public_config, ensure_ascii=False))

            service.save_config(
                {
                    "excel_path": "竞品监控_验收测试.xlsx",
                    "wecom_webhook": "******",
                    "taobao_username": "new-taobao",
                    "taobao_password": "new-pass",
                    "bi_username": "new-bi",
                    "bi_password": "new-bi-pass",
                    "run_date": "2026-07-07",
                    "dry_run": False,
                    "test_one": False,
                }
            )

            saved_text = config_path.read_text(encoding="utf-8")
            self.assertIn('excel_path: "竞品监控_验收测试.xlsx"', saved_text)
            self.assertIn('run_date: "2026-07-07"', saved_text)
            self.assertIn("dry_run: false", saved_text)
            self.assertIn("test_one: false", saved_text)
            self.assertIn("key=old", saved_text)
            self.assertNotIn("new-pass", saved_text)
            self.assertEqual(load_taobao_credentials({"taobao_credentials_path": str(taobao_path)}).username, "new-taobao")
            self.assertEqual(load_bi_credentials({"bi_credentials_path": str(bi_path)}).password, "new-bi-pass")

    def test_config_masks_non_phone_taobao_account_without_exposing_fragments(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            base_dir = Path(tmpdir)
            config_path = base_dir / "config.yaml"
            taobao_path = base_dir / "secrets" / "taobao.json"
            save_account_credentials(taobao_path, "漫步者官方旗舰店:金", "old-pass")
            config_path.write_text(
                "\n".join(
                    [
                        'excel_path: "竞品监控.xlsx"',
                        f'taobao_credentials_path: "{taobao_path.as_posix()}"',
                    ]
                ),
                encoding="utf-8",
            )
            service = WebApiService(base_dir=base_dir, config_path=config_path)

            public_config = service.get_config()
            public_json = json.dumps(public_config, ensure_ascii=False)

            self.assertEqual(public_config["taobao_username_masked"], "已保存账号")
            self.assertNotIn("漫步者", public_json)
            self.assertNotIn("旗舰店", public_json)

    def test_get_config_refreshes_persisted_template_upload_status_after_restart(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            base_dir = Path(tmpdir)
            config_path = base_dir / "config.yaml"
            config_path.write_text('excel_path: "竞品监控.xlsx"\n', encoding="utf-8")
            status_dir = base_dir / "competitor_monitor" / "web_state"
            status_dir.mkdir(parents=True)
            (status_dir / "status.json").write_text(
                json.dumps(
                    {
                        "template_uploaded_at": "2026-07-07 12:57:29",
                        "template_validation_status": "模板可用",
                        "state_updated_at": "2026-07-07 12:57:29",
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            service = WebApiService(base_dir=base_dir, config_path=config_path)

            public_config = service.get_config()

            self.assertEqual(public_config["template_uploaded_at"], "2026-07-07 12:57:29")
            self.assertEqual(public_config["template_validation_status"], "模板可用")

    def test_runtime_config_resolves_task_paths_under_web_runtime(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            base_dir = Path(tmpdir)
            service = WebApiService(base_dir=base_dir, config_path=base_dir / "competitor_monitor" / "config.yaml")
            config = {
                "excel_path": "竞品监控.xlsx",
                "backup_dir": "competitor_monitor/backup",
                "log_dir": "competitor_monitor/logs",
                "browser_user_data_dir": "competitor_monitor/browser_profile",
                "browser_storage_state_path": "competitor_monitor/browser_profile/storage_state.json",
                "taobao_credentials_path": "competitor_monitor/secrets/taobao_credentials.json",
                "bi_browser_user_data_dir": "competitor_monitor/browser_profile_bi",
                "bi_credentials_path": "competitor_monitor/secrets/bi_credentials.json",
            }

            runtime_config = service._runtime_config(config)

            for key in config:
                self.assertEqual(Path(runtime_config[key]), base_dir / config[key])

    def test_scheduled_tasks_pass_web_config_and_hide_task_runner_window(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            base_dir = Path(tmpdir) / "CompetitorMonitorWeb"
            base_dir.mkdir()
            runner_dir = base_dir.parent / "CompetitorMonitorTaskRunner"
            runner_dir.mkdir()
            (runner_dir / "CompetitorMonitorTaskRunner.exe").write_bytes(b"fake")
            config_path = base_dir / "competitor_monitor" / "config.yaml"
            config_path.parent.mkdir(parents=True)
            config_path.write_text('excel_path: "竞品监控.xlsx"\n', encoding="utf-8")
            captured = {}

            def fake_run(command, **kwargs):
                captured["command"] = command
                captured["script"] = command[-1]
                captured["kwargs"] = kwargs

                class Completed:
                    returncode = 0
                    stdout = ""
                    stderr = ""

                return Completed()

            service = WebApiService(base_dir=base_dir, config_path=config_path)
            original_run = subprocess.run
            try:
                subprocess.run = fake_run
                service._sync_windows_scheduled_tasks(
                    {
                        "daily_run_time": "09:30",
                        "weekly_new_time": "09:35",
                        "price_trend_time": "18:30",
                        "new_product_allowed_weekdays": [5],
                    }
                )
            finally:
                subprocess.run = original_run

            script = captured["script"]
            self.assertIn("--config", script)
            self.assertIn(str(config_path), script)
            self.assertIn('$action = New-ScheduledTaskAction -Execute "powershell.exe"', script)
            self.assertIn("-WindowStyle Hidden", script)
            self.assertIn("-EncodedCommand $encodedCommand", script)
            self.assertIn("Start-Process -FilePath", script)
            self.assertIn("-WindowStyle Hidden -Wait -PassThru", script)
            self.assertNotIn('$action = New-ScheduledTaskAction -Execute $runnerExe', script)
            if hasattr(subprocess, "CREATE_NO_WINDOW"):
                self.assertEqual(captured["kwargs"].get("creationflags"), subprocess.CREATE_NO_WINDOW)

    def test_run_task_reports_progress_results_and_honors_stop_request(self):
        started = threading.Event()

        def runner(context: RunContext) -> None:
            context.set_total(2)
            context.set_step("正在采集")
            context.record_result(
                {
                    "mode": context.payload["mode"],
                    "sheet": "6.29-7.3",
                    "brand": "塞那",
                    "model": "S6S Ultra",
                    "status": "success",
                    "output": "268.52",
                }
            )
            started.set()
            while not context.stop_requested():
                time.sleep(0.01)
            context.log("warning", "收到停止请求", "下一条商品开始前停止")

        with tempfile.TemporaryDirectory() as tmpdir:
            service = WebApiService(base_dir=Path(tmpdir), runner=runner)

            response = service.run_task({"mode": "daily_price", "run_date": "2026-07-03"})
            self.assertEqual(response["message"], "任务已启动")
            self.assertTrue(started.wait(1))

            running_status = service.get_tasks_status()
            self.assertEqual(running_status["system_status"], "running")
            self.assertEqual(running_status["progress_done"], 1)
            self.assertEqual(running_status["success_count"], 1)
            self.assertEqual(running_status["run_date"], "2026-07-03")

            stop_response = service.stop_task()
            self.assertEqual(stop_response["message"], "已请求停止当前任务")
            service.wait_for_current_task(timeout=1)

            stopped_status = service.get_tasks_status()
            self.assertEqual(stopped_status["system_status"], "stopped")
            self.assertEqual(stopped_status["last_error"], None)
            self.assertEqual(service.get_results()[0]["model"], "S6S Ultra")
            self.assertEqual(service.get_logs(level="warning")[0]["message"], "收到停止请求")

    def test_collection_phase_progress_updates_status_without_result_row(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            service = WebApiService(base_dir=Path(tmpdir))

            service._record_collection_progress(
                "daily_price",
                {
                    "event": "phase",
                    "processed": 0,
                    "total": 100,
                    "step": "已识别 100 条商品，正在启动浏览器",
                },
            )

            status = service.get_tasks_status()
            self.assertEqual(status["progress_total"], 100)
            self.assertEqual(status["progress_done"], 0)
            self.assertEqual(status["current_step"], "已识别 100 条商品，正在启动浏览器")
            self.assertEqual(service.get_results(), [])

    def test_run_task_persists_failure_traceback_for_scheduled_runner(self):
        def runner(context: RunContext) -> None:
            raise RuntimeError("browser launch failed")

        with tempfile.TemporaryDirectory() as tmpdir:
            base_dir = Path(tmpdir)
            config_path = base_dir / "config.yaml"
            config_path.write_text('log_dir: "competitor_monitor/logs"\n', encoding="utf-8")
            service = WebApiService(base_dir=base_dir, config_path=config_path, runner=runner)

            service.run_task({"mode": "daily_price", "run_date": "2026-07-07"})
            service.wait_for_current_task(timeout=1)

            failure_log = base_dir / "competitor_monitor" / "logs" / "web_task_failures.log"
            self.assertTrue(failure_log.exists())
            failure_text = failure_log.read_text(encoding="utf-8")
            self.assertIn("browser launch failed", failure_text)
            self.assertIn("RuntimeError", failure_text)

    def test_default_runner_warns_and_continues_when_startup_lock_check_is_strict(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            base_dir = Path(tmpdir)
            excel_path = base_dir / "current.xlsx"
            excel_path.write_bytes(b"fake workbook")
            config_path = base_dir / "config.yaml"
            config_path.write_text(
                'excel_path: "current.xlsx"\nlog_dir: "competitor_monitor/logs"\n',
                encoding="utf-8",
            )
            service = WebApiService(base_dir=base_dir, config_path=config_path)
            context = RunContext(service, {"mode": "daily_price", "run_date": "2026-07-07"}, service.stop_event)
            backup_called = {"value": False}
            wait_called = {"value": False}
            collection_called = {"value": False}

            def fake_backup(_source, _backup_dir):
                backup_called["value"] = True

            def fake_collection(**_kwargs):
                collection_called["value"] = True

            def fake_wait(_path, timeout_seconds=30, delay_seconds=1, on_wait=None):
                wait_called["value"] = True
                self.assertEqual(timeout_seconds, 3)
                if on_wait:
                    on_wait(1, 2)
                raise WorkbookLockedError("当前 Excel 文件被占用，请关闭 Excel/WPS/预览窗口后再运行：current.xlsx")

            with patch("web_api.backup_excel", fake_backup), patch(
                "web_api.ExcelService.wait_workbook_writable",
                fake_wait,
            ), patch("web_api.ExcelService.load_workbook", return_value=Workbook()), patch(
                "main.run_daily_price_job",
                side_effect=fake_collection,
            ):
                service._default_runner(context)

            self.assertTrue(wait_called["value"])
            self.assertTrue(backup_called["value"])
            self.assertTrue(collection_called["value"])
            warnings = service.get_logs(level="warning")
            self.assertTrue(any("启动前检测到 Excel 可能被占用" in item["message"] for item in warnings))

    def test_scheduled_runner_status_is_visible_to_web_console_process(self):
        started = threading.Event()
        release = threading.Event()

        def runner(context: RunContext) -> None:
            context.set_total(3)
            context.set_step("正在采集每日价格")
            started.set()
            release.wait(1)
            context.record_result({"mode": "daily_price", "model": "S6S proII", "status": "success"})

        with tempfile.TemporaryDirectory() as tmpdir:
            base_dir = Path(tmpdir)
            scheduled_service = WebApiService(base_dir=base_dir, runner=runner)
            web_console_service = WebApiService(base_dir=base_dir)

            scheduled_service.run_task({"mode": "daily_price", "run_date": "2026-07-07"})
            self.assertTrue(started.wait(1))

            running_status = web_console_service.get_tasks_status()
            self.assertEqual(running_status["system_status"], "running")
            self.assertEqual(running_status["active_mode"], "daily_price")
            self.assertEqual(running_status["current_step"], "正在采集每日价格")
            self.assertEqual(running_status["progress_total"], 3)

            release.set()
            scheduled_service.wait_for_current_task(timeout=1)
            finished_status = web_console_service.get_tasks_status()
            self.assertEqual(finished_status["system_status"], "success")
            self.assertEqual(finished_status["progress_done"], 1)

    def test_web_console_stop_request_reaches_scheduled_runner_process(self):
        started = threading.Event()

        def runner(context: RunContext) -> None:
            context.set_total(1)
            context.set_step("正在启动")
            started.set()
            deadline = time.time() + 1
            while time.time() < deadline and not context.stop_requested():
                time.sleep(0.01)

        with tempfile.TemporaryDirectory() as tmpdir:
            base_dir = Path(tmpdir)
            scheduled_service = WebApiService(base_dir=base_dir, runner=runner)
            web_console_service = WebApiService(base_dir=base_dir)

            scheduled_service.run_task({"mode": "daily_price", "run_date": "2026-07-07"})
            self.assertTrue(started.wait(1))

            stop_response = web_console_service.stop_task()
            scheduled_service.wait_for_current_task(timeout=2)
            status = web_console_service.get_tasks_status()

            self.assertEqual(stop_response["message"], "已请求停止当前任务")
            self.assertEqual(status["system_status"], "stopped")
            self.assertEqual(status["current_step"], "已停止")

    def test_scheduled_runner_preserves_template_upload_status_when_updating_shared_status(self):
        def runner(context: RunContext) -> None:
            context.set_total(1)
            context.record_result({"mode": "daily_price", "model": "S6S proII", "status": "skipped"})

        with tempfile.TemporaryDirectory() as tmpdir:
            base_dir = Path(tmpdir)
            status_dir = base_dir / "competitor_monitor" / "web_state"
            status_dir.mkdir(parents=True)
            (status_dir / "status.json").write_text(
                json.dumps(
                    {
                        "template_uploaded_at": "2026-07-07 12:57:29",
                        "template_validation_status": "模板可用",
                        "state_updated_at": "2026-07-07 12:57:29",
                    },
                    ensure_ascii=False,
                ),
                encoding="utf-8",
            )
            scheduled_service = WebApiService(base_dir=base_dir, runner=runner)
            web_console_service = WebApiService(base_dir=base_dir)

            scheduled_service.run_task({"mode": "daily_price", "run_date": "2026-07-07", "dry_run": True})
            scheduled_service.wait_for_current_task(timeout=1)
            status = web_console_service.get_tasks_status()

            self.assertEqual(status["template_uploaded_at"], "2026-07-07 12:57:29")
            self.assertEqual(status["template_validation_status"], "模板可用")

    def test_template_completeness_and_export_results_are_frontend_ready(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            base_dir = Path(tmpdir)
            config_path = base_dir / "config.yaml"
            config_path.write_text('excel_path: "missing.xlsx"\n', encoding="utf-8")
            service = WebApiService(base_dir=base_dir, config_path=config_path)
            service.record_result(
                {
                    "mode": "weekly_new",
                    "sheet": "6.29-7.3",
                    "brand": "绿联",
                    "model": "S6PRO",
                    "status": "skipped",
                    "error_reason": "重复商品",
                }
            )

            completeness = service.get_template_completeness("2026-07-03")
            self.assertEqual(completeness["sheet_name"], None)
            self.assertEqual(completeness["period_start"], "2026-06-29")
            self.assertFalse(completeness["complete"])
            self.assertEqual(completeness["missing_dates"], ["2026-06-29", "2026-06-30", "2026-07-01", "2026-07-02", "2026-07-03"])

            rows = list(csv.DictReader(service.export_results_csv().splitlines()))

            self.assertEqual(rows[0]["任务"], "weekly_new")
            self.assertEqual(rows[0]["品牌"], "绿联")
            self.assertEqual(rows[0]["状态"], "跳过")

    def test_results_are_persisted_and_reloaded_without_full_product_urls(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            base_dir = Path(tmpdir)
            service = WebApiService(base_dir=base_dir)

            service.record_result(
                {
                    "mode": "daily_price",
                    "sheet": "7.6-7.10",
                    "brand": "塞那",
                    "model": "S6S proII",
                    "status": "success",
                    "output": "268.52",
                    "url": "https://detail.tmall.com/item.htm?id=800417757444&pisk=secret-tracking&spm=abc",
                }
            )
            reloaded = WebApiService(base_dir=base_dir)
            results = reloaded.get_results()

            self.assertEqual(results[0]["model"], "S6S proII")
            self.assertEqual(results[0]["output"], "268.52")
            self.assertIn("id=800417757444", results[0]["url"])
            self.assertNotIn("pisk", results[0]["url"])
            self.assertTrue((base_dir / "competitor_monitor" / "web_state" / "results.jsonl").exists())

    def test_cleanup_results_clears_persisted_web_results_without_touching_business_files(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            base_dir = Path(tmpdir)
            service = WebApiService(base_dir=base_dir)
            excel_path = base_dir / "竞品监控.xlsx"
            log_path = base_dir / "competitor_monitor" / "logs" / "run.log"
            backup_path = base_dir / "competitor_monitor" / "backup" / "backup.xlsx"
            log_path.parent.mkdir(parents=True)
            backup_path.parent.mkdir(parents=True)
            excel_path.write_text("excel-placeholder", encoding="utf-8")
            log_path.write_text("log-placeholder", encoding="utf-8")
            backup_path.write_text("backup-placeholder", encoding="utf-8")

            service.record_result(
                {
                    "mode": "daily_price",
                    "sheet": "7.6-7.10",
                    "brand": "漫步者",
                    "model": "A5",
                    "status": "success",
                    "output": "109.65",
                }
            )
            results_path = base_dir / "competitor_monitor" / "web_state" / "results.jsonl"
            self.assertTrue(results_path.exists())

            response = service.cleanup_maintenance("results")

            self.assertEqual(service.get_results(), [])
            self.assertFalse(results_path.exists())
            self.assertIn("结果记录已清空", response["message"])
            self.assertIn("Excel、日志和备份文件未删除", response["message"])
            self.assertTrue(excel_path.exists())
            self.assertTrue(log_path.exists())
            self.assertTrue(backup_path.exists())

    def test_logs_include_recent_real_log_files_and_can_clear_view_without_deleting_files(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            base_dir = Path(tmpdir)
            log_dir = base_dir / "competitor_monitor" / "logs"
            log_dir.mkdir(parents=True)
            log_path = log_dir / "run_20260706_093000.log"
            log_path.write_text(
                "2026-07-06 09:30:01,123 [INFO] Saved Excel: demo.xlsx\n"
                "2026-07-06 09:30:02,123 [WARNING] 价格异常波动：Excel 第 24 行 A5\n"
                "2026-07-06 09:30:03,123 [INFO] Price trend sheet: 6.29-7.3\n",
                encoding="utf-8",
            )
            config_path = base_dir / "config.yaml"
            config_path.write_text('log_dir: "competitor_monitor/logs"\n', encoding="utf-8")
            service = WebApiService(base_dir=base_dir, config_path=config_path)

            logs = service.get_logs(query="A5")
            self.assertEqual(logs[0]["level"], "warning")
            self.assertIn("A5", logs[0]["message"])

            service.cleanup_maintenance("logs")
            self.assertTrue(log_path.exists())
            self.assertEqual(service.get_logs(query="A5"), [])

    def test_web_log_view_translates_known_internal_english_messages(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            base_dir = Path(tmpdir)
            log_dir = base_dir / "competitor_monitor" / "logs"
            log_dir.mkdir(parents=True)
            log_path = log_dir / "run_20260706_093000.log"
            log_path.write_text(
                "2026-07-06 09:30:03,123 [INFO] Price trend sheet: 6.29-7.3\n",
                encoding="utf-8",
            )
            config_path = base_dir / "config.yaml"
            config_path.write_text('log_dir: "competitor_monitor/logs"\n', encoding="utf-8")
            service = WebApiService(base_dir=base_dir, config_path=config_path)

            logs = service.get_logs(query="价格情况")

            self.assertEqual(logs[0]["message"], "价格情况 Sheet：6.29-7.3")

    def test_template_upload_validates_backs_up_and_replaces_current_template(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            base_dir = Path(tmpdir)
            current = base_dir / "current.xlsx"
            incoming = base_dir / "incoming.xlsx"
            create_uploadable_template(current, model="OLD")
            create_uploadable_template(incoming, model="NEW")
            config_path = base_dir / "config.yaml"
            config_path.write_text(
                'excel_path: "current.xlsx"\nbackup_dir: "competitor_monitor/backup"\n',
                encoding="utf-8",
            )
            service = WebApiService(base_dir=base_dir, config_path=config_path)

            response = service.upload_template("同事模板.xlsx", incoming.read_bytes())
            saved = load_workbook(current)

            self.assertEqual(response["message"], "模板已上传并替换当前主 Excel")
            self.assertEqual(saved["7.6-7.10"]["B4"].value, "NEW")
            self.assertTrue(list((base_dir / "competitor_monitor" / "backup").glob("current_*.xlsx")))
            self.assertEqual(service.get_config()["template_validation_status"], "模板可用")

    def test_template_upload_rejects_invalid_file_without_replacing_current_template(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            base_dir = Path(tmpdir)
            current = base_dir / "current.xlsx"
            create_uploadable_template(current, model="OLD")
            config_path = base_dir / "config.yaml"
            config_path.write_text('excel_path: "current.xlsx"\n', encoding="utf-8")
            service = WebApiService(base_dir=base_dir, config_path=config_path)

            response = service.upload_template("bad.xlsx", b"not an xlsx")
            saved = load_workbook(current)

            self.assertEqual(response["message"], "模板上传失败")
            self.assertIn("error", response)
            self.assertEqual(saved["7.6-7.10"]["B4"].value, "OLD")

    def test_template_upload_ignores_stale_cross_machine_pending_state(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            base_dir = Path(tmpdir)
            current = base_dir / "current.xlsx"
            incoming = base_dir / "incoming.xlsx"
            create_uploadable_template(current, model="OLD")
            create_uploadable_template(incoming, model="NEW")
            config_path = base_dir / "config.yaml"
            config_path.write_text('excel_path: "current.xlsx"\n', encoding="utf-8")
            service = WebApiService(base_dir=base_dir, config_path=config_path)
            service.workbook_state.mark_pending_sync(
                Path("G:/old-machine/output/latest/current.xlsx"),
                current,
                "old machine path",
            )

            response = service.upload_template("incoming.xlsx", incoming.read_bytes())
            saved = load_workbook(current)

            self.assertEqual(response["message"], "模板已上传并替换当前主 Excel")
            self.assertEqual(saved["7.6-7.10"]["B4"].value, "NEW")
            self.assertEqual(service.get_config()["workbook_sync_status"], "synced")

    def test_template_upload_ignores_stale_pending_state_when_primary_missing(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            base_dir = Path(tmpdir)
            current = base_dir / "current.xlsx"
            incoming = base_dir / "incoming.xlsx"
            create_uploadable_template(incoming, model="NEW")
            config_path = base_dir / "config.yaml"
            config_path.write_text('excel_path: "current.xlsx"\n', encoding="utf-8")
            service = WebApiService(base_dir=base_dir, config_path=config_path)
            service.workbook_state.mark_pending_sync(
                Path("G:/old-machine/output/latest/current.xlsx"),
                current,
                "old machine path",
            )

            response = service.upload_template("incoming.xlsx", incoming.read_bytes())
            saved = load_workbook(current)

            self.assertEqual(response["message"], "模板已上传并替换当前主 Excel")
            self.assertEqual(saved["7.6-7.10"]["B4"].value, "NEW")
            self.assertEqual(service.get_config()["workbook_sync_status"], "synced")

    def test_get_template_file_returns_current_real_excel_path(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            base_dir = Path(tmpdir)
            current = base_dir / "current.xlsx"
            create_uploadable_template(current)
            config_path = base_dir / "config.yaml"
            config_path.write_text('excel_path: "current.xlsx"\n', encoding="utf-8")
            service = WebApiService(base_dir=base_dir, config_path=config_path)

            path = service.get_template_file()

            self.assertEqual(path, current)

    def test_status_and_config_include_workbook_sync_state(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            base_dir = Path(tmpdir)
            current = base_dir / "current.xlsx"
            latest = base_dir / "output" / "latest" / "current.xlsx"
            create_uploadable_template(current, model="OLD")
            latest.parent.mkdir(parents=True)
            create_uploadable_template(latest, model="NEW")
            config_path = base_dir / "config.yaml"
            config_path.write_text('excel_path: "current.xlsx"\n', encoding="utf-8")
            service = WebApiService(base_dir=base_dir, config_path=config_path)
            service.workbook_state.mark_pending_sync(latest, current, "locked")

            status = service.get_tasks_status()
            config = service.get_config()

            self.assertEqual(status["workbook_sync_status"], "pending_sync")
            self.assertEqual(config["workbook_sync_status"], "pending_sync")
            self.assertEqual(Path(status["active_workbook_path"]), latest.resolve())
            self.assertEqual(Path(config["primary_workbook_path"]), current.resolve())
            self.assertIn("workbook_sync_status_text", status)
            self.assertIn("run_workbook_count", config)

    def test_template_download_and_notify_use_active_workbook_when_pending_sync(self):
        notify_calls = []

        def fake_notifier(config, excel_path, summary):
            notify_calls.append(excel_path)
            return {"file_name": excel_path.name, "text_status": "success", "file_status": "success"}

        with tempfile.TemporaryDirectory() as tmpdir:
            base_dir = Path(tmpdir)
            current = base_dir / "current.xlsx"
            latest = base_dir / "output" / "latest" / "current.xlsx"
            create_uploadable_template(current, model="OLD")
            latest.parent.mkdir(parents=True)
            create_uploadable_template(latest, model="NEW")
            config_path = base_dir / "config.yaml"
            config_path.write_text(
                "\n".join(
                    [
                        'excel_path: "current.xlsx"',
                        "wecom_enabled: true",
                        'wecom_webhook: "https://qyapi.weixin.qq.com/cgi-bin/webhook/send?key=secret"',
                    ]
                ),
                encoding="utf-8",
            )
            service = WebApiService(base_dir=base_dir, config_path=config_path, notifier=fake_notifier)
            service.workbook_state.mark_pending_sync(latest, current, "locked")

            template_path = service.get_template_file()
            service.send_template()

            self.assertEqual(template_path, latest.resolve())
            self.assertEqual(notify_calls, [latest.resolve()])

    def test_template_upload_rejects_pending_sync_without_replacing_current_template(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            base_dir = Path(tmpdir)
            current = base_dir / "current.xlsx"
            latest = base_dir / "output" / "latest" / "current.xlsx"
            incoming = base_dir / "incoming.xlsx"
            create_uploadable_template(current, model="OLD")
            latest.parent.mkdir(parents=True)
            create_uploadable_template(latest, model="NEW")
            create_uploadable_template(incoming, model="INCOMING")
            config_path = base_dir / "config.yaml"
            config_path.write_text('excel_path: "current.xlsx"\n', encoding="utf-8")
            service = WebApiService(base_dir=base_dir, config_path=config_path)
            service.workbook_state.mark_pending_sync(latest, current, "locked")

            response = service.upload_template("incoming.xlsx", incoming.read_bytes())
            saved = load_workbook(current)

            self.assertEqual(response["message"], "主模板待同步，暂不能上传新模板")
            self.assertIn("error", response)
            self.assertEqual(saved["7.6-7.10"]["B4"].value, "OLD")

    def test_sync_latest_template_copies_active_to_primary(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            base_dir = Path(tmpdir)
            current = base_dir / "current.xlsx"
            latest = base_dir / "output" / "latest" / "current.xlsx"
            create_uploadable_template(current, model="OLD")
            latest.parent.mkdir(parents=True)
            create_uploadable_template(latest, model="NEW")
            config_path = base_dir / "config.yaml"
            config_path.write_text('excel_path: "current.xlsx"\n', encoding="utf-8")
            service = WebApiService(base_dir=base_dir, config_path=config_path)
            service.workbook_state.mark_pending_sync(latest, current, "locked")

            response = service.sync_latest_template()
            saved = load_workbook(current)

            self.assertEqual(response["sync_status"], "synced")
            self.assertEqual(response["message"], "主模板已同步")
            self.assertEqual(saved["7.6-7.10"]["B4"].value, "NEW")

    def test_backfill_template_starts_real_background_task_and_blocks_duplicates(self):
        started = threading.Event()

        def runner(context: RunContext) -> None:
            started.set()
            while not context.stop_requested():
                time.sleep(0.01)

        with tempfile.TemporaryDirectory() as tmpdir:
            service = WebApiService(base_dir=Path(tmpdir), runner=runner)

            response = service.backfill_template({"dates": ["2026-07-06"]})
            self.assertEqual(response["message"], "补跑任务已启动")
            self.assertTrue(started.wait(1))
            duplicate = service.backfill_template({"dates": ["2026-07-07"]})
            self.assertEqual(duplicate["message"], "已有任务正在运行")
            service.stop_task()
            service.wait_for_current_task(timeout=1)

    def test_summary_results_are_visible_for_non_itemized_web_tasks(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            service = WebApiService(base_dir=Path(tmpdir))
            service.status["system_status"] = "running"

            service.record_task_summary("price_trend", "价格情况分析", "共 100 行，写入 97 行")

            status = service.get_tasks_status()
            results = service.get_results()

            self.assertEqual(status["progress_total"], 1)
            self.assertEqual(status["progress_done"], 1)
            self.assertEqual(status["success_count"], 1)
            self.assertEqual(results[0]["mode"], "price_trend")
            self.assertEqual(results[0]["model"], "价格情况分析")
            self.assertEqual(results[0]["output"], "共 100 行，写入 97 行")

    def test_dry_run_task_does_not_send_wecom_template(self):
        notify_calls = []

        def runner(context: RunContext) -> None:
            context.set_total(1)
            context.record_result({"mode": "daily_price", "model": "S6S proII", "status": "skipped"})

        def notifier(config, excel_path, summary):
            notify_calls.append((config, excel_path, summary))
            return {"file_name": excel_path.name, "text_status": "success", "file_status": "success"}

        with tempfile.TemporaryDirectory() as tmpdir:
            base_dir = Path(tmpdir)
            config_path = base_dir / "config.yaml"
            config_path.write_text(
                "\n".join(
                    [
                        'excel_path: "missing.xlsx"',
                        "wecom_enabled: true",
                        'wecom_webhook: "https://qyapi.weixin.qq.com/cgi-bin/webhook/send?key=secret"',
                    ]
                ),
                encoding="utf-8",
            )
            service = WebApiService(base_dir=base_dir, config_path=config_path, runner=runner, notifier=notifier)

            service.run_task({"mode": "daily_price", "run_date": "2026-07-06", "dry_run": True})
            service.wait_for_current_task(timeout=1)

            status = service.get_tasks_status()
            self.assertEqual(status["system_status"], "success")
            self.assertEqual(notify_calls, [])
            self.assertEqual(status["last_notify_file_status"], "not_sent")

    def test_notify_and_session_endpoints_update_public_state_without_exposing_secrets(self):
        def fake_notifier(config, excel_path, summary):
            return {"file_name": excel_path.name, "text_status": "success", "file_status": "success"}

        with tempfile.TemporaryDirectory() as tmpdir:
            base_dir = Path(tmpdir)
            excel_path = base_dir / "竞品监控.xlsx"
            excel_path.write_bytes(b"fake excel")
            config_path = base_dir / "config.yaml"
            config_path.write_text(
                "\n".join(
                    [
                        f'excel_path: "{excel_path.as_posix()}"',
                        'wecom_enabled: true',
                        'wecom_webhook: "https://qyapi.weixin.qq.com/cgi-bin/webhook/send?key=secret"',
                    ]
                ),
                encoding="utf-8",
            )
            service = WebApiService(base_dir=base_dir, config_path=config_path, notifier=fake_notifier)

            notify = service.send_template()
            session = service.reset_session("taobao")
            status = service.get_tasks_status()
            config = service.get_config()

            self.assertEqual(notify["message"], "当前模板已发送")
            self.assertEqual(status["last_notify_file_status"], "success")
            self.assertEqual(status["last_sent_file_name"], "竞品监控.xlsx")
            self.assertEqual(session["message"], "登录状态已重置")
            self.assertEqual(config["taobao_status"], "需重新登录")
            self.assertNotIn("secret", json.dumps(status, ensure_ascii=False))

    def test_send_template_reports_not_sent_when_wecom_is_disabled(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            base_dir = Path(tmpdir)
            excel_path = base_dir / "竞品监控.xlsx"
            excel_path.write_bytes(b"fake excel")
            config_path = base_dir / "config.yaml"
            config_path.write_text(
                "\n".join(
                    [
                        f'excel_path: "{excel_path.as_posix()}"',
                        "wecom_enabled: false",
                        'wecom_webhook: ""',
                    ]
                ),
                encoding="utf-8",
            )
            service = WebApiService(base_dir=base_dir, config_path=config_path)

            response = service.send_template()
            status = service.get_tasks_status()

            self.assertEqual(response["message"], "企业微信未启用，未发送")
            self.assertEqual(status["last_notify_text_status"], "not_sent")
            self.assertEqual(status["last_notify_file_status"], "not_sent")
            self.assertIsNone(status["last_notify_time"])
            self.assertIsNone(status["last_sent_file_name"])

    def test_wecom_summary_includes_price_alerts_for_manual_review(self):
        summaries = []

        def fake_notifier(config, excel_path, summary):
            summaries.append(summary)
            return {"file_name": excel_path.name, "text_status": "success", "file_status": "success"}

        with tempfile.TemporaryDirectory() as tmpdir:
            base_dir = Path(tmpdir)
            excel_path = base_dir / "竞品监控.xlsx"
            excel_path.write_bytes(b"fake excel")
            config_path = base_dir / "config.yaml"
            config_path.write_text(
                "\n".join(
                    [
                        f'excel_path: "{excel_path.as_posix()}"',
                        'wecom_enabled: true',
                        'wecom_webhook: "https://qyapi.weixin.qq.com/cgi-bin/webhook/send?key=secret"',
                    ]
                ),
                encoding="utf-8",
            )
            service = WebApiService(base_dir=base_dir, config_path=config_path, notifier=fake_notifier)
            service.log("warning", "价格异常波动，需人工复核", "row=24 商品=A5 历史价=109.65 本次写入=79.9")

            service.send_template()

            self.assertTrue(summaries)
            self.assertIn("需人工复核的价格异常", summaries[0])
            self.assertIn("A5", summaries[0])
            self.assertIn("109.65", summaries[0])


if __name__ == "__main__":
    unittest.main()
