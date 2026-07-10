import json
import tempfile
import threading
import unittest
from pathlib import Path
from uuid import uuid4
from urllib.request import Request, urlopen
from unittest.mock import patch

from openpyxl import Workbook

from web_api import WebApiService
from web_server import _safe_print, _write_startup_error, create_server, run_server


def create_template(path: Path, model: str = "S6S proII") -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "7.6-7.10"
    ws["A1"] = "品牌"
    ws["B1"] = "竞争品牌核心单品监控"
    ws["J1"] = "上新监控"
    ws["B2"] = "型号"
    ws["C2"] = "7月6"
    ws["D2"] = "7月7"
    ws["E2"] = "7月8"
    ws["F2"] = "7月9"
    ws["G2"] = "7月10"
    ws["H2"] = "价格情况"
    ws["I2"] = "活动"
    ws["J2"] = "上架日期"
    ws["K2"] = "型号"
    ws["L2"] = "形态"
    ws["M2"] = "价格"
    ws["S2"] = "新卖点"
    ws["A4"] = "塞那"
    ws["B4"] = model
    ws["B4"].hyperlink = "https://detail.tmall.com/item.htm?id=1"
    wb.save(path)


def multipart_body(field_name: str, filename: str, content: bytes) -> tuple[bytes, str]:
    boundary = f"----codex-{uuid4().hex}"
    body = b"".join(
        [
            f"--{boundary}\r\n".encode("utf-8"),
            f'Content-Disposition: form-data; name="{field_name}"; filename="{filename}"\r\n'.encode("utf-8"),
            b"Content-Type: application/vnd.openxmlformats-officedocument.spreadsheetml.sheet\r\n\r\n",
            content,
            b"\r\n",
            f"--{boundary}--\r\n".encode("utf-8"),
        ]
    )
    return body, f"multipart/form-data; boundary={boundary}"


class WebServerTest(unittest.TestCase):
    def test_runtime_base_dir_can_be_overridden_for_packaged_runtime_data(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            with patch.dict("os.environ", {"COMPETITOR_MONITOR_RUNTIME_DIR": tmpdir}):
                from web_server import runtime_base_dir

                self.assertEqual(runtime_base_dir(), Path(tmpdir))

    def test_safe_print_does_not_crash_when_stdout_is_unavailable(self):
        with patch("builtins.print", side_effect=AttributeError("stdout unavailable")):
            _safe_print("Web 控制台已启动")

    def test_startup_errors_are_written_to_runtime_log(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            base_dir = Path(tmpdir)
            with patch("web_server.runtime_base_dir", return_value=base_dir):
                _write_startup_error(RuntimeError("bind failed"))

            log_path = base_dir / "competitor_monitor" / "logs" / "web_server_startup.log"
            self.assertTrue(log_path.exists())
            self.assertIn("bind failed", log_path.read_text(encoding="utf-8"))

    def test_run_server_opens_browser_when_existing_service_is_available(self):
        opened = []

        with patch("web_server._health_is_available", return_value=True), patch(
            "web_server._open_web_console",
            side_effect=lambda url: opened.append(url),
        ):
            run_server("127.0.0.1", 8765)

        self.assertEqual(opened, ["http://127.0.0.1:8765/"])

    def test_run_server_opens_browser_after_local_server_starts(self):
        opened = []

        class FakeServer:
            def serve_forever(self):
                raise KeyboardInterrupt

            def server_close(self):
                pass

        with tempfile.TemporaryDirectory() as tmpdir:
            with patch("web_server.runtime_base_dir", return_value=Path(tmpdir)), patch(
                "web_server.ensure_local_config",
                return_value=Path(tmpdir) / "competitor_monitor" / "config.yaml",
            ), patch("web_server.create_server", return_value=FakeServer()), patch(
                "web_server._open_web_console",
                side_effect=lambda url: opened.append(url),
            ), patch("web_server._health_is_available", return_value=False):
                run_server("127.0.0.1", 8765)

        self.assertEqual(opened, ["http://127.0.0.1:8765/"])

    def test_serves_static_frontend_and_json_api(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            base_dir = Path(tmpdir)
            web_root = base_dir / "web_frontend"
            web_root.mkdir()
            (web_root / "index.html").write_text("<!doctype html><title>竞品监控</title>", encoding="utf-8")
            service = WebApiService(base_dir=base_dir)
            server = create_server("127.0.0.1", 0, service, web_root)
            thread = threading.Thread(target=server.serve_forever, daemon=True)
            thread.start()
            try:
                root = f"http://127.0.0.1:{server.server_address[1]}"
                html = urlopen(f"{root}/", timeout=2).read().decode("utf-8")
                status = json.loads(urlopen(f"{root}/api/tasks/status", timeout=2).read().decode("utf-8"))

                self.assertIn("竞品监控", html)
                self.assertEqual(status["system_status"], "idle")
            finally:
                server.shutdown()
                server.server_close()

    def test_serves_health_endpoint_for_startup_checks(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            service = WebApiService(base_dir=Path(tmpdir))
            server = create_server("127.0.0.1", 0, service, Path(tmpdir))
            thread = threading.Thread(target=server.serve_forever, daemon=True)
            thread.start()
            try:
                root = f"http://127.0.0.1:{server.server_address[1]}"
                health = json.loads(urlopen(f"{root}/api/health", timeout=2).read().decode("utf-8"))

                self.assertEqual(health["status"], "ok")
                self.assertEqual(health["service"], "competitor-monitor-web")
            finally:
                server.shutdown()
                server.server_close()

    def test_routes_post_json_to_service(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            service = WebApiService(base_dir=Path(tmpdir), runner=lambda context: None)
            server = create_server("127.0.0.1", 0, service, Path(tmpdir))
            thread = threading.Thread(target=server.serve_forever, daemon=True)
            thread.start()
            try:
                root = f"http://127.0.0.1:{server.server_address[1]}"
                request = Request(
                    f"{root}/api/tasks/run",
                    data=json.dumps({"mode": "daily_price", "run_date": "2026-07-03"}).encode("utf-8"),
                    headers={"Content-Type": "application/json"},
                    method="POST",
                )
                response = json.loads(urlopen(request, timeout=2).read().decode("utf-8"))

                self.assertEqual(response["message"], "任务已启动")
            finally:
                server.shutdown()
                server.server_close()

    def test_template_download_and_upload_routes_use_real_excel_files(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            base_dir = Path(tmpdir)
            current = base_dir / "current.xlsx"
            incoming = base_dir / "incoming.xlsx"
            create_template(current, "OLD")
            create_template(incoming, "NEW")
            config_path = base_dir / "config.yaml"
            config_path.write_text('excel_path: "current.xlsx"\n', encoding="utf-8")
            service = WebApiService(base_dir=base_dir, config_path=config_path)
            server = create_server("127.0.0.1", 0, service, Path(tmpdir))
            thread = threading.Thread(target=server.serve_forever, daemon=True)
            thread.start()
            try:
                root = f"http://127.0.0.1:{server.server_address[1]}"
                downloaded = urlopen(f"{root}/api/template/download", timeout=2).read()
                self.assertGreater(len(downloaded), 1000)

                body, content_type = multipart_body("template", "同事模板.xlsx", incoming.read_bytes())
                request = Request(
                    f"{root}/api/template/upload",
                    data=body,
                    headers={"Content-Type": content_type, "Content-Length": str(len(body))},
                    method="POST",
                )
                response = json.loads(urlopen(request, timeout=2).read().decode("utf-8"))

                self.assertEqual(response["message"], "模板已上传并替换当前主 Excel")
                self.assertEqual(service.get_template_file(), current)
            finally:
                server.shutdown()
                server.server_close()

    def test_sync_latest_route_copies_pending_workbook_to_primary(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            base_dir = Path(tmpdir)
            current = base_dir / "current.xlsx"
            latest = base_dir / "output" / "latest" / "current.xlsx"
            create_template(current, "OLD")
            latest.parent.mkdir(parents=True)
            create_template(latest, "NEW")
            config_path = base_dir / "config.yaml"
            config_path.write_text('excel_path: "current.xlsx"\n', encoding="utf-8")
            service = WebApiService(base_dir=base_dir, config_path=config_path)
            service.workbook_state.mark_pending_sync(latest, current, "locked")
            server = create_server("127.0.0.1", 0, service, Path(tmpdir))
            thread = threading.Thread(target=server.serve_forever, daemon=True)
            thread.start()
            try:
                root = f"http://127.0.0.1:{server.server_address[1]}"
                request = Request(
                    f"{root}/api/template/sync-latest",
                    data=b"{}",
                    headers={"Content-Type": "application/json"},
                    method="POST",
                )
                response = json.loads(urlopen(request, timeout=2).read().decode("utf-8"))

                self.assertEqual(response["sync_status"], "synced")
                self.assertEqual(response["message"], "主模板已同步")
            finally:
                server.shutdown()
                server.server_close()


if __name__ == "__main__":
    unittest.main()
