import unittest
from copy import copy
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Border, Side

from new_product_service import NewProduct
from new_product_writer import (
    NewProductWriteStats,
    detect_new_product_layout,
    match_new_product_brand,
    write_new_products,
)


def create_new_product_sheet():
    wb = Workbook()
    ws = wb.active
    ws.title = "6.15-6.19"
    ws["A1"] = "品牌"
    ws["B2"] = "型号"
    ws["J1"] = "上新监控"
    ws.merge_cells("J1:S1")
    ws["J2"] = "上架日期"
    ws.merge_cells("J2:J3")
    ws["K2"] = "型号"
    ws.merge_cells("K2:K3")
    ws["L2"] = "形态"
    ws.merge_cells("L2:L3")
    ws["M2"] = "价格"
    ws.merge_cells("M2:M3")
    ws["N2"] = "直播价格"
    ws.merge_cells("N2:R2")
    ws["S2"] = "新卖点"
    ws.merge_cells("S2:S3")
    for cell, value in zip(["N3", "O3", "P3", "Q3", "R3"], ["6月15", "6月16", "6月17", "6月18", "6月19"]):
        ws[cell] = value

    ws["A4"] = "绿联"
    ws.merge_cells("A4:A12")
    for row in range(4, 13):
        ws.cell(row=row, column=2).value = f"old-{row}"
    ws.merge_cells("J4:S12")
    ws["J4"] = "无上新"

    ws["A13"] = "华为"
    ws.merge_cells("A13:A18")
    for row in range(13, 19):
        ws.cell(row=row, column=2).value = f"huawei-{row}"
    ws.merge_cells("J13:S18")
    ws["J13"] = "无上新"
    return wb, ws


class NewProductWriterTest(unittest.TestCase):
    def test_detects_new_product_layout_from_headers(self):
        _, ws = create_new_product_sheet()

        layout = detect_new_product_layout(ws)

        self.assertEqual(layout.date_column, 10)
        self.assertEqual(layout.model_column, 11)
        self.assertEqual(layout.shape_column, 12)
        self.assertEqual(layout.price_column, 13)
        self.assertEqual(layout.live_price_columns, [14, 15, 16, 17, 18])
        self.assertEqual(layout.selling_point_column, 19)
        self.assertEqual(layout.data_start_row, 4)

    def test_writes_product_like_existing_green_union_example(self):
        _, ws = create_new_product_sheet()
        layout = detect_new_product_layout(ws)
        date_style = copy(ws["J4"]._style)
        product = NewProduct(
            brand="绿联",
            on_sale_date=date(2026, 6, 18),
            model="S6PRO",
            shape="耳夹式",
            price="199",
            link="https://example.com/s6pro",
            selling_point="无",
            source_id="s6pro",
            raw_name="绿联 S6PRO 耳夹式蓝牙耳机",
        )

        stats = write_new_products(ws, layout, {"绿联": [product]}, no_data_text="无上新")

        self.assertEqual(stats, NewProductWriteStats(written=1, skipped_duplicates=0, no_data_written=0, failed=0))
        self.assertEqual(ws["J4"].value, "6月18")
        self.assertEqual(ws["K4"].value, "S6PRO")
        self.assertEqual(ws["K4"].hyperlink.target, "https://example.com/s6pro")
        self.assertEqual(ws["L4"].value, "耳夹式")
        self.assertEqual(ws["M4"].value, "199")
        self.assertEqual([ws.cell(row=4, column=col).value for col in range(14, 19)], ["/", "/", "/", "/", "/"])
        self.assertEqual(ws["S4"].value, "无")
        self.assertEqual(copy(ws["J4"]._style), date_style)
        self.assertNotIn("J4:S12", {str(rng) for rng in ws.merged_cells.ranges})

    def test_writes_no_new_when_brand_has_no_products(self):
        _, ws = create_new_product_sheet()
        layout = detect_new_product_layout(ws)

        stats = write_new_products(ws, layout, {"华为": []}, no_data_text="无上新")

        self.assertEqual(stats.no_data_written, 0)
        self.assertEqual(ws["J13"].value, "无上新")
        self.assertIn("J13:S18", {str(rng) for rng in ws.merged_cells.ranges})

    def test_skips_duplicate_existing_link(self):
        _, ws = create_new_product_sheet()
        layout = detect_new_product_layout(ws)
        product = NewProduct(
            brand="绿联",
            on_sale_date=date(2026, 6, 18),
            model="S6PRO",
            shape="耳夹式",
            price="199",
            link="https://example.com/s6pro",
            selling_point="无",
            source_id="s6pro",
            raw_name="绿联 S6PRO 耳夹式蓝牙耳机",
        )
        write_new_products(ws, layout, {"绿联": [product]}, no_data_text="无上新")

        stats = write_new_products(ws, layout, {"绿联": [product]}, no_data_text="无上新")

        self.assertEqual(stats.skipped_duplicates, 1)
        self.assertEqual(ws["K5"].value, None)

    def test_expands_brand_area_when_no_empty_new_product_row_exists(self):
        wb = Workbook()
        ws = wb.active
        ws["J1"] = "上新监控"
        ws.merge_cells("J1:S1")
        ws["J2"] = "上架日期"
        ws.merge_cells("J2:J3")
        ws["K2"] = "型号"
        ws.merge_cells("K2:K3")
        ws["L2"] = "形态"
        ws.merge_cells("L2:L3")
        ws["M2"] = "价格"
        ws.merge_cells("M2:M3")
        ws["N2"] = "直播价格"
        ws.merge_cells("N2:R2")
        ws["S2"] = "新卖点"
        ws.merge_cells("S2:S3")
        ws["A4"] = "索尼影音"
        ws.merge_cells("A4:A5")
        ws["B4"] = "sony-left-4"
        ws["B5"] = "sony-left-5"
        ws["K4"] = "existing-1"
        ws["K5"] = "existing-2"
        ws["A6"] = "华为"
        ws.merge_cells("A6:A7")
        ws["B6"] = "huawei-left-6"
        ws["B6"].hyperlink = "https://example.com/huawei-left-6"
        ws["B7"] = "huawei-left-7"
        ws.merge_cells("J6:S7")
        ws["J6"] = "无上新"
        ws["K5"].style = "Hyperlink"
        thin = Side(style="thin")
        ws["K5"].border = Border(left=thin, right=thin, top=thin, bottom=thin)
        original_following_brand_merges = {
            str(rng)
            for rng in ws.merged_cells.ranges
            if rng.min_col <= 1 <= rng.max_col and rng.min_row >= 6
        }
        layout = detect_new_product_layout(ws)
        product = NewProduct(
            brand="索尼影音",
            on_sale_date=date(2026, 7, 3),
            model="H9",
            shape="头戴式",
            price="1969",
            link="https://example.com/h9",
            selling_point="/",
            source_id="h9",
            raw_name="Sony H9",
        )

        stats = write_new_products(ws, layout, {"索尼影音": [product]}, no_data_text="无上新")

        self.assertEqual(stats.failed, 0)
        self.assertEqual(stats.written, 1)
        self.assertEqual(stats.failure_details, ())
        self.assertEqual(ws["K6"].value, "H9")
        self.assertEqual(ws["K6"].hyperlink.target, "https://example.com/h9")
        self.assertEqual(ws["A4"].value, "索尼影音")
        self.assertIn("A4:A5", {str(rng) for rng in ws.merged_cells.ranges})
        self.assertEqual(ws["B4"].value, "sony-left-4")
        self.assertEqual(ws["B5"].value, "sony-left-5")
        self.assertEqual(ws["A6"].value, "华为")
        self.assertIn("A6:A7", {str(rng) for rng in ws.merged_cells.ranges})
        self.assertEqual(ws["B6"].value, "huawei-left-6")
        self.assertEqual(ws["B6"].hyperlink.target, "https://example.com/huawei-left-6")
        self.assertEqual(ws["B7"].value, "huawei-left-7")
        self.assertEqual(
            {
                str(rng)
                for rng in ws.merged_cells.ranges
                if rng.min_col <= 1 <= rng.max_col and rng.min_row >= 6
            },
            original_following_brand_merges,
        )
        self.assertIn("J7:S7", {str(rng) for rng in ws.merged_cells.ranges})
        self.assertEqual(ws["J7"].value, "无上新")
        self.assertFalse(
            any(
                rng.min_row <= 6 <= rng.max_row and rng.min_col <= 10 and rng.max_col >= 19
                for rng in ws.merged_cells.ranges
            )
        )
        self.assertEqual(ws["K6"].style, "Hyperlink")
        self.assertEqual(ws["K6"].border.top.style, "thin")
        self.assertEqual(ws["K6"].border.bottom.style, "thin")

        duplicate_stats = write_new_products(ws, layout, {"索尼影音": [product]}, no_data_text="无上新")

        self.assertEqual(duplicate_stats.skipped_duplicates, 1)
        self.assertEqual(ws["K6"].value, "H9")
        self.assertEqual(ws["K7"].value, None)
        self.assertIn("J7:S7", {str(rng) for rng in ws.merged_cells.ranges})

    def test_right_side_expansion_keeps_following_brand_ranges_shifted(self):
        wb = Workbook()
        ws = wb.active
        ws["J1"] = "上新监控"
        ws.merge_cells("J1:S1")
        ws["J2"] = "上架日期"
        ws.merge_cells("J2:J3")
        ws["K2"] = "型号"
        ws.merge_cells("K2:K3")
        ws["L2"] = "形态"
        ws.merge_cells("L2:L3")
        ws["M2"] = "价格"
        ws.merge_cells("M2:M3")
        ws["N2"] = "直播价格"
        ws.merge_cells("N2:R2")
        ws["S2"] = "新卖点"
        ws.merge_cells("S2:S3")
        ws["A4"] = "索尼影音"
        ws.merge_cells("A4:A5")
        ws["K4"] = "existing-1"
        ws["K5"] = "existing-2"
        ws["A6"] = "绿联"
        ws.merge_cells("A6:A7")
        ws.merge_cells("J6:S7")
        ws["J6"] = "无上新"
        ws["A8"] = "华为"
        ws.merge_cells("A8:A9")
        ws.merge_cells("J8:S9")
        ws["J8"] = "无上新"
        ws["A10"] = "小米"
        ws.merge_cells("A10:A11")
        layout = detect_new_product_layout(ws)

        stats = write_new_products(
            ws,
            layout,
            {
                "索尼影音": [
                    NewProduct(
                        brand="索尼影音",
                        on_sale_date=date(2026, 7, 3),
                        model="H9",
                        shape="头戴式",
                        price="1969",
                        link="https://example.com/h9",
                        selling_point="/",
                        source_id="h9",
                        raw_name="Sony H9",
                    )
                ],
                "小米": [
                    NewProduct(
                        brand="小米",
                        on_sale_date=date(2026, 7, 2),
                        model="Buds8",
                        shape="未知",
                        price="139",
                        link="https://example.com/buds8",
                        selling_point="/",
                        source_id="buds8",
                        raw_name="Xiaomi Buds8",
                    )
                ],
            },
            no_data_text="无上新",
        )

        self.assertEqual(stats.failed, 0)
        self.assertEqual(stats.written, 2)
        self.assertEqual(ws["K6"].value, "H9")
        self.assertIn("J7:S7", {str(rng) for rng in ws.merged_cells.ranges})
        self.assertIn("J8:S9", {str(rng) for rng in ws.merged_cells.ranges})
        self.assertEqual(ws["K10"].value, "Buds8")
        self.assertEqual(ws["A6"].value, "绿联")
        self.assertIn("A6:A7", {str(rng) for rng in ws.merged_cells.ranges})
        self.assertEqual(ws["A8"].value, "华为")
        self.assertIn("A8:A9", {str(rng) for rng in ws.merged_cells.ranges})
        self.assertEqual(ws["A10"].value, "小米")
        self.assertIn("A10:A11", {str(rng) for rng in ws.merged_cells.ranges})

    def test_appended_linked_model_cell_keeps_table_border_when_source_link_style_has_no_border(self):
        wb = Workbook()
        ws = wb.active
        ws["J1"] = "上新监控"
        ws.merge_cells("J1:S1")
        ws["J2"] = "上架日期"
        ws.merge_cells("J2:J3")
        ws["K2"] = "型号"
        ws.merge_cells("K2:K3")
        ws["L2"] = "形态"
        ws.merge_cells("L2:L3")
        ws["M2"] = "价格"
        ws.merge_cells("M2:M3")
        ws["N2"] = "直播价格"
        ws.merge_cells("N2:R2")
        ws["S2"] = "新卖点"
        ws.merge_cells("S2:S3")
        ws["A4"] = "索尼影音"
        ws.merge_cells("A4:A5")
        ws["K4"] = "existing-1"
        ws["K5"] = "existing-2"
        thin = Side(style="thin")
        table_border = Border(left=thin, right=thin, top=thin, bottom=thin)
        for row in range(4, 8):
            for column in range(10, 20):
                ws.cell(row=row, column=column).border = table_border
        ws["K5"].style = "Hyperlink"
        self.assertIsNone(ws["K5"].border.top.style)
        ws["A6"] = "绿联"
        ws.merge_cells("A6:A7")
        ws.merge_cells("J6:S7")
        ws["J6"] = "无上新"
        layout = detect_new_product_layout(ws)
        product = NewProduct(
            brand="索尼影音",
            on_sale_date=date(2026, 7, 3),
            model="H9",
            shape="头戴式",
            price="1969",
            link="https://example.com/h9",
            selling_point="/",
            source_id="h9",
            raw_name="Sony H9",
        )

        stats = write_new_products(ws, layout, {"索尼影音": [product]}, no_data_text="无上新")

        self.assertEqual(stats.written, 1)
        self.assertEqual(ws["K6"].value, "H9")
        self.assertEqual(ws["K6"].hyperlink.target, "https://example.com/h9")
        self.assertEqual(ws["K6"].font.underline, "single")
        self.assertEqual(ws["K6"].border.left.style, "thin")
        self.assertEqual(ws["K6"].border.right.style, "thin")
        self.assertEqual(ws["K6"].border.top.style, "thin")
        self.assertEqual(ws["K6"].border.bottom.style, "thin")

    def test_matches_real_bi_shop_names_to_excel_brand_names(self):
        excel_brands = ["索爱", "索尼官方", "索尼影音", "JBL", "倍思"]

        self.assertEqual(match_new_product_brand("索尼官方旗舰店", excel_brands), "索尼官方")
        self.assertEqual(match_new_product_brand("索尼影音旗舰店", excel_brands), "索尼影音")
        self.assertEqual(match_new_product_brand("JBL耳机旗舰店", excel_brands), "JBL")

    def test_brand_shop_aliases_can_make_a_brand_strict(self):
        excel_brands = ["索爱", "索尼官方", "索尼影音", "绿联", "万魔"]
        aliases = {
            "索爱": ["SOAIY旗舰店"],
            "索尼官方": ["索尼官方旗舰店"],
            "索尼影音": ["索尼影音旗舰店"],
            "绿联": ["绿联数码旗舰店"],
            "万魔": ["1MORE万魔官方旗舰店"],
        }

        self.assertEqual(match_new_product_brand("SOAIY旗舰店", excel_brands, aliases), "索爱")
        self.assertIsNone(match_new_product_brand("索爱麒麟专卖店", excel_brands, aliases))
        self.assertIsNone(match_new_product_brand("索爱数码旗舰店", excel_brands, aliases))
        self.assertEqual(match_new_product_brand("索尼官方旗舰店", excel_brands, aliases), "索尼官方")
        self.assertEqual(match_new_product_brand("索尼影音旗舰店", excel_brands, aliases), "索尼影音")
        self.assertEqual(match_new_product_brand("绿联数码旗舰店", excel_brands, aliases), "绿联")
        self.assertEqual(match_new_product_brand("1MORE万魔官方旗舰店", excel_brands, aliases), "万魔")

    def test_writes_to_exact_excel_brand_when_normalized_names_overlap(self):
        wb = Workbook()
        ws = wb.active
        ws["J1"] = "上新监控"
        ws.merge_cells("J1:S1")
        ws["J2"] = "上架日期"
        ws.merge_cells("J2:J3")
        ws["K2"] = "型号"
        ws.merge_cells("K2:K3")
        ws["L2"] = "形态"
        ws.merge_cells("L2:L3")
        ws["M2"] = "价格"
        ws.merge_cells("M2:M3")
        ws["N2"] = "直播价格"
        ws.merge_cells("N2:R2")
        ws["S2"] = "新卖点"
        ws.merge_cells("S2:S3")
        ws["A4"] = "索尼官方"
        ws.merge_cells("A4:A5")
        ws.merge_cells("J4:S5")
        ws["J4"] = "无上新"
        ws["A6"] = "索尼影音"
        ws.merge_cells("A6:A7")
        ws.merge_cells("J6:S7")
        ws["J6"] = "无上新"
        layout = detect_new_product_layout(ws)
        product = NewProduct(
            brand="索尼影音",
            on_sale_date=date(2026, 6, 30),
            model="LinkBuds Clip",
            shape="耳夹式",
            price="999",
            link="https://example.com/sony",
            selling_point="无",
            source_id="sony",
            raw_name="Sony LinkBuds Clip",
        )

        stats = write_new_products(ws, layout, {"索尼影音": [product]}, no_data_text="无上新")

        self.assertEqual(stats.written, 1)
        self.assertEqual(ws["J4"].value, "无上新")
        self.assertEqual(ws["K6"].value, "LinkBuds Clip")


if __name__ == "__main__":
    unittest.main()
