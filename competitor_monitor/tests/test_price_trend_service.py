import unittest
from decimal import Decimal

from openpyxl import Workbook

from excel_service import SheetLayout
from price_trend_service import (
    PriceTrendWriteStats,
    classify_price_trend,
    extract_lowest_price,
    update_price_trends,
)


class PriceTrendServiceTest(unittest.TestCase):
    def test_extracts_lowest_price_from_multi_sku_text(self):
        self.assertEqual(extract_lowest_price("黑286.36 紫281.26 金323.76"), Decimal("281.26"))
        self.assertEqual(extract_lowest_price("268.52/305.17"), Decimal("268.52"))
        self.assertIsNone(extract_lowest_price("/"))
        self.assertIsNone(extract_lowest_price(""))

    def test_classifies_core_weekly_price_trends(self):
        cases = [
            ([100, 100, 100, 100, 100], "价格稳定"),
            ([100, 101, 102, 102, 103], "价格上涨"),
            ([100, 99, 98, 98, 97], "价格下降"),
            ([100, 95, 95, 98, 101], "先降后升"),
            ([100, 105, 105, 103, 99], "先升后降"),
        ]
        for prices, expected in cases:
            with self.subTest(prices=prices):
                self.assertEqual(classify_price_trend(prices), expected)

    def test_ignores_missing_prices_and_describes_complex_movements(self):
        self.assertEqual(classify_price_trend(["/", 449.4, 533.8, 529.55, 495.3]), "先升后降")
        self.assertEqual(classify_price_trend(["/", 2499, 2499, 2499, 2499]), "价格稳定")
        self.assertEqual(classify_price_trend([100, 99, None, 98, 97]), "价格下降")
        self.assertEqual(classify_price_trend([100, "/", 98, 98, 97]), "价格下降")
        self.assertIsNone(classify_price_trend(["/", "/", "/", "/", "/"]))
        self.assertEqual(classify_price_trend([100, 98, 101, 99, 102]), "先跌、再涨、再跌、再涨")
        self.assertEqual(classify_price_trend([989.5, 998, 958.2, 958.2, 998]), "先涨、再跌、再涨")

    def test_updates_price_status_cells_without_overwriting_existing_values(self):
        wb = Workbook()
        ws = wb.active
        layout = SheetLayout(
            model_column=2,
            price_column=7,
            price_status_column=8,
            activity_column=9,
            price_header_row=3,
            data_start_row=4,
            right_start_column=10,
            date_price_columns=[3, 4, 5, 6, 7],
        )
        ws.cell(row=4, column=2).value = "稳定款"
        for column in layout.date_price_columns:
            ws.cell(row=4, column=column).value = "黑286.36 紫281.26 金323.76"
        ws.cell(row=5, column=2).value = "下降款"
        for column, value in zip(layout.date_price_columns, [100, 99, 98, 98, 97]):
            ws.cell(row=5, column=column).value = value
        ws.cell(row=6, column=2).value = "已有人工判断"
        for column, value in zip(layout.date_price_columns, [100, 101, 102, 102, 103]):
            ws.cell(row=6, column=column).value = value
        ws.cell(row=6, column=8).value = "人工备注"
        ws.cell(row=7, column=2).value = "首日缺失款"
        for column, value in zip(layout.date_price_columns, ["/", 100, 98, 98, 97]):
            ws.cell(row=7, column=column).value = value
        ws.cell(row=8, column=2).value = "复杂波动款"
        for column, value in zip(layout.date_price_columns, [100, 98, 101, 99, 102]):
            ws.cell(row=8, column=column).value = value
        ws.cell(row=9, column=2).value = "中间缺失款"
        for column, value in zip(layout.date_price_columns, [100, "/", 98, 98, 97]):
            ws.cell(row=9, column=column).value = value
        ws.cell(row=10, column=2).value = "全缺失款"
        for column, value in zip(layout.date_price_columns, ["/", "/", "/", "/", "/"]):
            ws.cell(row=10, column=column).value = value

        stats = update_price_trends(ws, layout, force_overwrite=False)

        self.assertEqual(
            stats,
            PriceTrendWriteStats(total=7, written=5, skipped_existing=1, skipped_incomplete=1, skipped_unclassified=0),
        )
        self.assertEqual(ws.cell(row=4, column=8).value, "价格稳定")
        self.assertEqual(ws.cell(row=5, column=8).value, "价格下降")
        self.assertEqual(ws.cell(row=6, column=8).value, "人工备注")
        self.assertEqual(ws.cell(row=7, column=8).value, "价格下降")
        self.assertEqual(ws.cell(row=8, column=8).value, "先跌、再涨、再跌、再涨")
        self.assertEqual(ws.cell(row=9, column=8).value, "价格下降")
        self.assertIsNone(ws.cell(row=10, column=8).value)


if __name__ == "__main__":
    unittest.main()
