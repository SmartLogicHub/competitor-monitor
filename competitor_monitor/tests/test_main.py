import logging
import io
import tempfile
import threading
import unittest
from datetime import date
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

import main
from bi_service import BICollectResult
from excel_service import ExcelService
from new_product_service import NewProduct
from new_product_writer import NewProductWriteStats


class MainCollectionOptionsTest(unittest.TestCase):
    def test_parse_row_filter_accepts_lists_and_ranges(self):
        self.assertEqual(main.parse_row_filter("24, 54,60-62"), {24, 54, 60, 61, 62})
        self.assertIsNone(main.parse_row_filter(None))
        self.assertIsNone(main.parse_row_filter(""))

    def test_parse_row_filter_rejects_invalid_values(self):
        for value in ["0", "-1", "abc", "10-8", "4,,5"]:
            with self.subTest(value=value):
                with self.assertRaises(ValueError):
                    main.parse_row_filter(value)

    def test_select_target_products_prefers_test_one_over_limit(self):
        products = [SimpleNamespace(row=i) for i in range(1, 6)]

        selected = main.select_target_products(products, test_one=True, limit=3)

        self.assertEqual([product.row for product in selected], [1])

    def test_select_target_products_filters_by_rows_before_limit(self):
        products = [SimpleNamespace(row=i) for i in range(1, 7)]

        selected = main.select_target_products(products, test_one=False, limit=2, row_filter={2, 4, 5})

        self.assertEqual([product.row for product in selected], [2, 4])

    def test_select_target_products_applies_limit_for_batch_testing(self):
        products = [SimpleNamespace(row=i) for i in range(1, 6)]

        selected = main.select_target_products(products, test_one=False, limit=3)

        self.assertEqual([product.row for product in selected], [1, 2, 3])

    def test_should_save_checkpoint_every_n_processed_items(self):
        self.assertFalse(main.should_save_checkpoint(1, 5))
        self.assertTrue(main.should_save_checkpoint(5, 5))
        self.assertFalse(main.should_save_checkpoint(6, 5))
        self.assertFalse(main.should_save_checkpoint(5, 0))

    def test_apply_cli_overrides_can_force_price_overwrite(self):
        config = {"force_overwrite": False}

        updated = main.apply_cli_overrides(config, SimpleNamespace(force_overwrite=True))

        self.assertIs(updated, config)
        self.assertTrue(config["force_overwrite"])

    def test_collect_with_retry_retries_until_success(self):
        browser = RetryBrowser(failures_before_success=2)
        collector = RetryCollector()

        result = main.collect_with_retry(browser, collector, "https://example.com/item", retry_count=2)

        self.assertEqual(result.price, "88.00")
        self.assertEqual(browser.open_count, 3)
        self.assertEqual(collector.collect_count, 1)

    def test_collect_with_retry_raises_after_all_attempts_fail(self):
        browser = RetryBrowser(failures_before_success=3)
        collector = RetryCollector()

        with self.assertRaises(TimeoutError):
            main.collect_with_retry(browser, collector, "https://example.com/item", retry_count=2)

        self.assertEqual(browser.open_count, 3)
        self.assertEqual(collector.collect_count, 0)


class RunCollectionCheckpointTest(unittest.TestCase):
    def test_run_collection_honors_stop_request_before_browser_startup(self):
        wb = Workbook()
        ws = wb.active
        layout = SimpleNamespace(price_column=3, activity_column=9)
        products = [SimpleNamespace(row=4, brand="A", name="P1", url="https://example.com/1")]
        excel = FakeExcel(reload_layout=layout)
        logger = logging.getLogger("test_run_collection_honors_stop_request_before_browser_startup")
        stop_event = threading.Event()
        stop_event.set()

        with patch.object(main, "BrowserService", side_effect=AssertionError("browser should not start")):
            stats = main.run_collection(
                config={
                    "min_delay_seconds": 0,
                    "max_delay_seconds": 0,
                    "retry_count": 0,
                    "enable_sku_matching": False,
                },
                logger=logger,
                excel=excel,
                workbook=wb,
                worksheet=ws,
                layout=layout,
                products=products,
                test_one=False,
                limit=None,
                row_filter=None,
                save_every=0,
                excel_path=None,
                target_date=date(2026, 7, 3),
                stop_event=stop_event,
            )

        self.assertEqual(stats.total, 1)
        self.assertEqual(stats.success, 0)

    def test_run_collection_collects_only_filtered_rows(self):
        wb = Workbook()
        ws = wb.active
        layout = SimpleNamespace(price_column=3, activity_column=9)
        products = [
            SimpleNamespace(row=4, brand="A", name="P1", url="https://example.com/1"),
            SimpleNamespace(row=5, brand="A", name="P2", url="https://example.com/2"),
            SimpleNamespace(row=6, brand="A", name="P3", url="https://example.com/3"),
        ]
        excel = FakeExcel(reload_layout=layout)
        logger = logging.getLogger("test_run_collection_collects_only_filtered_rows")

        with patch.object(main, "BrowserService", FakeBrowserService):
            with patch.object(main, "collect_with_retry", side_effect=fake_collect_with_retry) as collect_mock:
                with patch.object(main.time, "sleep") as sleep_mock:
                    stats = main.run_collection(
                        config={
                            "min_delay_seconds": 5,
                            "max_delay_seconds": 10,
                            "retry_count": 0,
                            "enable_sku_matching": False,
                        },
                        logger=logger,
                        excel=excel,
                        workbook=wb,
                        worksheet=ws,
                        layout=layout,
                        products=products,
                        test_one=False,
                        limit=None,
                        row_filter={5, 6},
                        save_every=0,
                        excel_path=None,
                        target_date=date(2026, 7, 3),
                    )

        self.assertEqual(stats.total, 2)
        self.assertEqual(stats.success, 2)
        self.assertEqual(collect_mock.call_count, 2)
        self.assertIsNone(ws.cell(row=4, column=3).value)
        self.assertEqual(ws.cell(row=5, column=3).value, "99.00")
        self.assertEqual(ws.cell(row=6, column=3).value, "99.00")

    def test_run_collection_saves_checkpoint_and_final_workbook(self):
        wb = Workbook()
        ws = wb.active
        layout = SimpleNamespace(price_column=3, activity_column=9)
        products = [
            SimpleNamespace(row=4, brand="A", name="P1", url="https://example.com/1"),
            SimpleNamespace(row=5, brand="A", name="P2", url="https://example.com/2"),
            SimpleNamespace(row=6, brand="A", name="P3", url="https://example.com/3"),
        ]
        excel = FakeExcel(reload_layout=layout)
        logger = logging.getLogger("test_run_collection_saves_after_each_checkpoint")

        with patch.object(main, "BrowserService", FakeBrowserService):
            with patch.object(main, "collect_with_retry", side_effect=fake_collect_with_retry):
                with patch.object(main.time, "sleep"):
                    stats = main.run_collection(
                        config={
                            "min_delay_seconds": 0,
                            "max_delay_seconds": 0,
                            "retry_count": 0,
                            "enable_sku_matching": False,
                        },
                        logger=logger,
                        excel=excel,
                        workbook=wb,
                        worksheet=ws,
                        layout=layout,
                        products=products,
                        test_one=False,
                        limit=None,
                        save_every=2,
                        excel_path="monitor.xlsx",
                        target_date=date(2026, 7, 3),
                    )

        self.assertEqual(stats.success, 3)
        self.assertEqual(excel.saved_paths, ["monitor.xlsx", "monitor.xlsx"])

    def test_run_collection_reloads_workbook_after_each_checkpoint(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "6.29-7.3"
        layout = SimpleNamespace(price_column=3, activity_column=9)
        products = [
            SimpleNamespace(row=4, brand="A", name="P1", url="https://example.com/1"),
            SimpleNamespace(row=5, brand="A", name="P2", url="https://example.com/2"),
            SimpleNamespace(row=6, brand="A", name="P3", url="https://example.com/3"),
        ]
        excel = FakeExcel(
            sheet_title=ws.title,
            reload_layout=layout,
            fail_on_repeated_workbook_save=True,
        )
        logger = logging.getLogger("test_run_collection_reloads_workbook_after_each_checkpoint")

        with patch.object(main, "BrowserService", FakeBrowserService):
            with patch.object(main, "collect_with_retry", side_effect=fake_collect_with_retry):
                with patch.object(main.time, "sleep"):
                    stats = main.run_collection(
                        config={
                            "min_delay_seconds": 0,
                            "max_delay_seconds": 0,
                            "retry_count": 0,
                            "enable_sku_matching": False,
                        },
                        logger=logger,
                        excel=excel,
                        workbook=wb,
                        worksheet=ws,
                        layout=layout,
                        products=products,
                        test_one=False,
                        limit=None,
                        save_every=1,
                        excel_path="monitor.xlsx",
                        target_date=date(2026, 7, 3),
                    )

        self.assertEqual(stats.success, 3)
        self.assertEqual(len(excel.saved_workbook_ids), 3)
        self.assertEqual(len(set(excel.saved_workbook_ids)), 3)

    def test_run_collection_skips_duplicate_item_urls_with_different_models(self):
        wb = Workbook()
        ws = wb.active
        layout = SimpleNamespace(price_column=3, activity_column=9)
        products = [
            SimpleNamespace(
                row=4,
                brand="Sanag",
                name="S6S proII",
                url="https://detail.tmall.com/item.htm?id=800417757444&pisk=aaa",
            ),
            SimpleNamespace(
                row=5,
                brand="Sanag",
                name="S6S ultra",
                url="https://detail.tmall.com/item.htm?id=800417757444&pisk=bbb",
            ),
        ]
        excel = FakeExcel(reload_layout=layout)
        logger = logging.getLogger("test_run_collection_skips_duplicate_item_urls")

        with patch.object(main, "BrowserService", FakeBrowserService):
            with patch.object(main, "collect_with_retry", side_effect=fake_collect_with_retry) as collect_mock:
                with patch.object(main.time, "sleep") as sleep_mock:
                    stats = main.run_collection(
                        config={
                            "min_delay_seconds": 5,
                            "max_delay_seconds": 10,
                            "retry_count": 0,
                            "enable_sku_matching": False,
                        },
                        logger=logger,
                        excel=excel,
                        workbook=wb,
                        worksheet=ws,
                        layout=layout,
                        products=products,
                        test_one=False,
                        limit=None,
                        save_every=0,
                        excel_path=None,
                        target_date=date(2026, 7, 3),
                    )

        collect_mock.assert_not_called()
        self.assertEqual(stats.success, 0)
        self.assertEqual(stats.skipped, 2)
        self.assertEqual(stats.failed, 0)
        self.assertIsNone(ws.cell(row=4, column=3).value)
        self.assertIsNone(ws.cell(row=5, column=3).value)
        self.assertIsNone(ws.cell(row=4, column=9).value)
        self.assertIsNone(ws.cell(row=5, column=9).value)

    def test_run_collection_uses_expected_model_for_duplicate_item_urls_when_sku_matching_enabled(self):
        wb = Workbook()
        ws = wb.active
        layout = SimpleNamespace(price_column=3, activity_column=9)
        products = [
            SimpleNamespace(
                row=4,
                brand="Sanag",
                name="S6S proII",
                url="https://detail.tmall.com/item.htm?id=800417757444&pisk=aaa",
            ),
            SimpleNamespace(
                row=5,
                brand="Sanag",
                name="S6S ultra",
                url="https://detail.tmall.com/item.htm?id=800417757444&pisk=bbb",
            ),
        ]
        excel = FakeExcel(reload_layout=layout)
        logger = logging.getLogger("test_run_collection_uses_expected_model")

        with patch.object(main, "BrowserService", FakeBrowserService):
            with patch.object(main, "collect_with_retry", side_effect=fake_collect_with_retry) as collect_mock:
                with patch.object(main.time, "sleep"):
                    stats = main.run_collection(
                        config={
                            "min_delay_seconds": 0,
                            "max_delay_seconds": 0,
                            "retry_count": 0,
                            "enable_sku_matching": True,
                        },
                        logger=logger,
                        excel=excel,
                        workbook=wb,
                        worksheet=ws,
                        layout=layout,
                        products=products,
                        test_one=False,
                        limit=None,
                        save_every=0,
                        excel_path=None,
                        target_date=date(2026, 7, 3),
                    )

        self.assertEqual(stats.success, 2)
        self.assertEqual([call.kwargs["expected_model"] for call in collect_mock.call_args_list], ["S6S proII", "S6S ultra"])
        self.assertEqual([call.kwargs["allow_fallback_price"] for call in collect_mock.call_args_list], [False, False])

    def test_run_collection_uses_expected_model_for_unique_item_urls_when_sku_matching_enabled(self):
        wb = Workbook()
        ws = wb.active
        layout = SimpleNamespace(price_column=3, activity_column=9)
        products = [
            SimpleNamespace(row=4, brand="Moondrop", name="\u592a\u7a7a\u6f2b\u6e382", url="https://detail.tmall.com/item.htm?id=730861790114"),
            SimpleNamespace(row=5, brand="Kinyon", name="A5", url="https://detail.tmall.com/item.htm?id=694593508978"),
        ]
        excel = FakeExcel(reload_layout=layout)
        logger = logging.getLogger("test_run_collection_uses_expected_model_for_unique_urls")

        with patch.object(main, "BrowserService", FakeBrowserService):
            with patch.object(main, "collect_with_retry", side_effect=fake_collect_with_retry) as collect_mock:
                with patch.object(main.time, "sleep"):
                    stats = main.run_collection(
                        config={
                            "min_delay_seconds": 0,
                            "max_delay_seconds": 0,
                            "retry_count": 0,
                            "enable_sku_matching": True,
                        },
                        logger=logger,
                        excel=excel,
                        workbook=wb,
                        worksheet=ws,
                        layout=layout,
                        products=products,
                        test_one=False,
                        limit=None,
                        save_every=0,
                        excel_path=None,
                        target_date=date(2026, 7, 3),
                    )

        self.assertEqual(stats.success, 2)
        self.assertEqual([call.kwargs["expected_model"] for call in collect_mock.call_args_list], ["\u592a\u7a7a\u6f2b\u6e382", "A5"])
        self.assertEqual([call.kwargs["allow_fallback_price"] for call in collect_mock.call_args_list], [True, True])

    def test_run_collection_does_not_skip_previous_slash_rows_by_default(self):
        wb = Workbook()
        ws = wb.active
        layout = SimpleNamespace(price_column=7, activity_column=9, date_price_columns=[3, 4, 5, 6, 7])
        for column in [3, 4, 5, 6]:
            ws.cell(row=4, column=column).value = "/"
        products = [SimpleNamespace(row=4, brand="Sony", name="WH-1000XM4", url="https://example.com/item")]
        excel = FakeExcel(reload_layout=layout)
        logger = logging.getLogger("test_run_collection_does_not_skip_previous_slash_rows_by_default")

        with patch.object(main, "BrowserService", FakeBrowserService):
            with patch.object(main, "collect_with_retry", side_effect=fake_collect_with_retry) as collect_mock:
                with patch.object(main.time, "sleep"):
                    stats = main.run_collection(
                        config={
                            "min_delay_seconds": 0,
                            "max_delay_seconds": 0,
                            "retry_count": 0,
                            "enable_sku_matching": False,
                        },
                        logger=logger,
                        excel=excel,
                        workbook=wb,
                        worksheet=ws,
                        layout=layout,
                        products=products,
                        test_one=False,
                        limit=None,
                        save_every=0,
                        excel_path=None,
                        target_date=date(2026, 7, 3),
                    )

        collect_mock.assert_called_once()
        self.assertEqual(stats.success, 1)
        self.assertEqual(ws.cell(row=4, column=7).value, "99.00")

    def test_previous_slash_skip_is_not_applied_on_monday_column(self):
        wb = Workbook()
        ws = wb.active
        layout = SimpleNamespace(price_column=3, activity_column=9, date_price_columns=[3, 4, 5, 6, 7])
        products = [SimpleNamespace(row=4, brand="Sony", name="WH-1000XM4", url="https://example.com/item")]
        excel = FakeExcel(reload_layout=layout)
        logger = logging.getLogger("test_previous_slash_skip_is_not_applied_on_monday_column")

        with patch.object(main, "BrowserService", FakeBrowserService):
            with patch.object(main, "collect_with_retry", side_effect=fake_collect_with_retry) as collect_mock:
                with patch.object(main.time, "sleep"):
                    stats = main.run_collection(
                        config={
                            "min_delay_seconds": 0,
                            "max_delay_seconds": 0,
                            "retry_count": 0,
                            "enable_sku_matching": False,
                            "skip_when_previous_prices_all_slash": True,
                        },
                        logger=logger,
                        excel=excel,
                        workbook=wb,
                        worksheet=ws,
                        layout=layout,
                        products=products,
                        test_one=False,
                        limit=None,
                        save_every=0,
                        excel_path=None,
                        target_date=date(2026, 6, 29),
                    )

        collect_mock.assert_called_once()
        self.assertEqual(stats.success, 1)
        self.assertEqual(ws.cell(row=4, column=3).value, "99.00")

    def test_previous_slash_skip_can_be_enabled_for_later_weekdays(self):
        wb = Workbook()
        ws = wb.active
        layout = SimpleNamespace(price_column=7, activity_column=9, date_price_columns=[3, 4, 5, 6, 7])
        for column in [3, 4, 5, 6]:
            ws.cell(row=4, column=column).value = "/"
        products = [SimpleNamespace(row=4, brand="Sony", name="WH-1000XM4", url="https://example.com/item")]
        excel = FakeExcel(reload_layout=layout)
        logger = logging.getLogger("test_previous_slash_skip_can_be_enabled_for_later_weekdays")

        with patch.object(main, "BrowserService", FakeBrowserService):
            with patch.object(main, "collect_with_retry", side_effect=fake_collect_with_retry) as collect_mock:
                with patch.object(main.time, "sleep"):
                    stats = main.run_collection(
                        config={
                            "min_delay_seconds": 0,
                            "max_delay_seconds": 0,
                            "retry_count": 0,
                            "enable_sku_matching": False,
                            "skip_when_previous_prices_all_slash": True,
                        },
                        logger=logger,
                        excel=excel,
                        workbook=wb,
                        worksheet=ws,
                        layout=layout,
                        products=products,
                        test_one=False,
                        limit=None,
                        save_every=0,
                        excel_path=None,
                        target_date=date(2026, 7, 3),
                    )

        collect_mock.assert_not_called()
        self.assertEqual(stats.success, 0)
        self.assertEqual(stats.skipped, 1)
        self.assertIsNone(ws.cell(row=4, column=7).value)

    def test_latest_previous_price_uses_rightmost_valid_price(self):
        wb = Workbook()
        ws = wb.active
        layout = SimpleNamespace(price_column=7, date_price_columns=[3, 4, 5, 6, 7])
        ws.cell(row=4, column=3).value = "/"
        ws.cell(row=4, column=4).value = ""
        ws.cell(row=4, column=5).value = "109.65"
        ws.cell(row=4, column=6).value = None

        self.assertEqual(main.latest_previous_price(ws, layout, 4), "109.65")

    def test_run_collection_uses_previous_price_when_collection_is_uncertain(self):
        wb = Workbook()
        ws = wb.active
        layout = SimpleNamespace(price_column=7, activity_column=9, date_price_columns=[3, 4, 5, 6, 7])
        ws.cell(row=4, column=5).value = "442"
        products = [SimpleNamespace(row=4, brand="Sanag", name="S7S", url="https://detail.tmall.com/item.htm?id=1")]
        excel = FakeExcel(reload_layout=layout)
        logger = logging.getLogger("test_run_collection_uses_previous_price_when_collection_is_uncertain")

        with patch.object(main, "BrowserService", FakeBrowserService):
            with patch.object(main, "collect_with_retry", return_value=SimpleNamespace(price=None, activities=[])):
                with patch.object(main.time, "sleep"):
                    stats = main.run_collection(
                        config={
                            "min_delay_seconds": 0,
                            "max_delay_seconds": 0,
                            "retry_count": 0,
                            "enable_sku_matching": True,
                            "fallback_to_previous_price_when_uncertain": True,
                        },
                        logger=logger,
                        excel=excel,
                        workbook=wb,
                        worksheet=ws,
                        layout=layout,
                        products=products,
                        test_one=False,
                        limit=None,
                        save_every=0,
                        excel_path=None,
                        target_date=date(2026, 7, 3),
                    )

        self.assertEqual(stats.success, 1)
        self.assertEqual(stats.failed, 0)
        self.assertEqual(ws.cell(row=4, column=7).value, "442")

    def test_run_collection_uses_previous_sheet_price_when_current_week_has_no_history(self):
        wb = Workbook()
        previous_ws = wb.active
        create_mature_price_sheet(
            previous_ws,
            title="6.23-6.26",
            date_headers=["6月23", "6月24", "6月25", "6月26"],
            row=4,
            brand="Sanag",
            model="S7S",
            prices=["438", "440", "441", "442"],
        )
        current_ws = wb.create_sheet("6.29-7.3")
        current_layout = SimpleNamespace(price_column=3, activity_column=9, date_price_columns=[3, 4, 5, 6, 7])
        current_ws.cell(row=4, column=1).value = "Sanag"
        current_ws.cell(row=4, column=2).value = "S7S"
        products = [SimpleNamespace(row=4, brand="Sanag", name="S7S", url="https://detail.tmall.com/item.htm?id=1")]
        excel = ExcelService("unused.xlsx")
        logger = logging.getLogger("test_run_collection_uses_previous_sheet_price")

        with patch.object(main, "BrowserService", FakeBrowserService):
            with patch.object(main, "collect_with_retry", return_value=SimpleNamespace(price=None, activities=[])):
                with patch.object(main.time, "sleep"):
                    stats = main.run_collection(
                        config={
                            "min_delay_seconds": 0,
                            "max_delay_seconds": 0,
                            "retry_count": 0,
                            "enable_sku_matching": True,
                            "fallback_to_previous_price_when_uncertain": True,
                        },
                        logger=logger,
                        excel=excel,
                        workbook=wb,
                        worksheet=current_ws,
                        layout=current_layout,
                        products=products,
                        test_one=False,
                        limit=None,
                        save_every=0,
                        excel_path=None,
                        target_date=date(2026, 6, 29),
                    )

        self.assertEqual(stats.success, 1)
        self.assertEqual(stats.failed, 0)
        self.assertEqual(current_ws.cell(row=4, column=3).value, "442")

    def test_manual_price_override_beats_wrong_recent_history_for_known_problem_product(self):
        wb = Workbook()
        ws = wb.active
        layout = SimpleNamespace(price_column=7, activity_column=9, date_price_columns=[3, 4, 5, 6, 7])
        ws.cell(row=4, column=6).value = "79.9"
        products = [SimpleNamespace(row=4, brand="金运", name="A5", url="https://detail.tmall.com/item.htm?id=1")]
        excel = FakeExcel(reload_layout=layout)
        logger = logging.getLogger("test_manual_price_override_beats_wrong_recent_history")

        with patch.object(main, "BrowserService", FakeBrowserService):
            with patch.object(main, "collect_with_retry", return_value=SimpleNamespace(price=None, activities=[])):
                with patch.object(main.time, "sleep"):
                    stats = main.run_collection(
                        config={
                            "min_delay_seconds": 0,
                            "max_delay_seconds": 0,
                            "retry_count": 0,
                            "enable_sku_matching": True,
                            "fallback_to_previous_price_when_uncertain": True,
                            "manual_price_overrides": [
                                {"brand": "金运", "model": "A5", "price": "109.65"},
                            ],
                        },
                        logger=logger,
                        excel=excel,
                        workbook=wb,
                        worksheet=ws,
                        layout=layout,
                        products=products,
                        test_one=False,
                        limit=None,
                        save_every=0,
                        excel_path=None,
                        target_date=date(2026, 7, 6),
                    )

        self.assertEqual(stats.success, 1)
        self.assertEqual(stats.failed, 0)
        self.assertEqual(ws.cell(row=4, column=7).value, "109.65")

    def test_detect_price_change_alert_for_large_price_drop(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "7.6-7.10"
        layout = SimpleNamespace(price_column=7, date_price_columns=[3, 4, 5, 6, 7])
        ws.cell(row=4, column=6).value = "109.65"
        product = SimpleNamespace(row=4, brand="金运", name="A5")

        alert = main.detect_price_change_alert(
            {"price_alert_enabled": True, "price_alert_threshold_percent": 20, "price_alert_min_abs_diff": 20},
            FakeExcel(reload_layout=layout),
            wb,
            ws,
            layout,
            product,
            date(2026, 7, 6),
            "79.9",
        )

        self.assertIsNotNone(alert)
        self.assertIn("价格异常波动", alert)
        self.assertIn("A5", alert)
        self.assertIn("109.65", alert)
        self.assertIn("79.9", alert)

    def test_run_collection_prefers_visible_discount_price_when_sku_price_hits_alert(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "7.6-7.10"
        layout = SimpleNamespace(
            price_column=3,
            activity_column=9,
            date_price_columns=[3, 4, 5, 6, 7],
            model_column=2,
            data_start_row=4,
        )
        previous_ws = wb.create_sheet("6.29-7.3")
        previous_ws.cell(row=82, column=2).value = "WH-1000XM5"
        previous_ws.cell(row=82, column=7).value = "1686.49"
        products = [
            SimpleNamespace(
                row=82,
                brand=None,
                name="WH-1000XM5",
                url="https://detail.tmall.com/item.htm?id=sony",
            )
        ]
        excel = FakeExcel(reload_layout=layout)
        logger = logging.getLogger("test_run_collection_prefers_visible_discount_price_when_sku_price_hits_alert")
        snapshot = SimpleNamespace(
            price="2699",
            activities=["\u8ddf\u4ef7\u534f\u8bae"],
            visible_text="Sony WH-1000XM5 \u5e97\u94fa\u4f18\u60e0\u540e \uffe51699 \u4f18\u60e0\u524d \uffe52999 \u8d85\u7ea7\u7206\u6b3e",
        )

        with patch.object(main, "BrowserService", FakeBrowserService):
            with patch.object(main, "collect_with_retry", return_value=snapshot):
                with patch.object(main.time, "sleep"):
                    stats = main.run_collection(
                        config={
                            "min_delay_seconds": 0,
                            "max_delay_seconds": 0,
                            "retry_count": 0,
                            "enable_sku_matching": True,
                            "price_alert_enabled": True,
                            "price_alert_threshold_percent": 20,
                            "price_alert_min_abs_diff": 20,
                        },
                        logger=logger,
                        excel=excel,
                        workbook=wb,
                        worksheet=ws,
                        layout=layout,
                        products=products,
                        test_one=False,
                        limit=None,
                        save_every=0,
                        excel_path=None,
                        target_date=date(2026, 7, 6),
                    )

        self.assertEqual(stats.success, 1)
        self.assertEqual(stats.failed, 0)
        self.assertEqual(ws.cell(row=82, column=3).value, "1699")
        self.assertEqual(ws.cell(row=82, column=9).value, "\u8ddf\u4ef7\u534f\u8bae")

    def test_run_collection_marks_slash_when_page_no_longer_contains_target_model(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "7.6-7.10"
        layout = SimpleNamespace(price_column=3, activity_column=9, date_price_columns=[3, 4, 5, 6, 7])
        products = [
            SimpleNamespace(
                row=91,
                brand="\u7eff\u8054",
                name="T3pro",
                url="https://detail.tmall.com/item.htm?id=old",
            )
        ]
        excel = FakeExcel(reload_layout=layout)
        logger = logging.getLogger("test_run_collection_marks_slash_when_page_no_longer_contains_target_model")
        snapshot = SimpleNamespace(
            price=None,
            activities=[],
            visible_text="\u7eff\u8054Q3\u84dd\u7259\u8033\u673a \u578b\u53f7 WS212 \uffe5439",
        )

        with patch.object(main, "BrowserService", FakeBrowserService):
            with patch.object(main, "collect_with_retry", return_value=snapshot):
                with patch.object(main.time, "sleep"):
                    stats = main.run_collection(
                        config={
                            "min_delay_seconds": 0,
                            "max_delay_seconds": 0,
                            "retry_count": 0,
                            "enable_sku_matching": True,
                        },
                        logger=logger,
                        excel=excel,
                        workbook=wb,
                        worksheet=ws,
                        layout=layout,
                        products=products,
                        test_one=False,
                        limit=None,
                        save_every=0,
                        excel_path=None,
                        target_date=date(2026, 7, 6),
                    )

        self.assertEqual(stats.success, 1)
        self.assertEqual(stats.failed, 0)
        self.assertEqual(ws.cell(row=91, column=3).value, "/")
        self.assertEqual(ws.cell(row=91, column=9).value, "/")

    def test_run_collection_does_not_mark_chinese_model_parts_as_delisted(self):
        wb = Workbook()
        ws = wb.active
        ws.title = "7.6-7.10"
        layout = SimpleNamespace(price_column=3, activity_column=9, date_price_columns=[3, 4, 5, 6, 7])
        products = [
            SimpleNamespace(
                row=57,
                brand="\u6c34\u6708\u96e8",
                name="\u7fbd\u7ffc\u5934\u6234\u5f0f",
                url="https://detail.tmall.com/item.htm?id=851486873910",
            )
        ]
        excel = FakeExcel(reload_layout=layout)
        logger = logging.getLogger("test_run_collection_does_not_mark_chinese_model_parts_as_delisted")
        stream = io.StringIO()
        handler = logging.StreamHandler(stream)
        logger.addHandler(handler)
        logger.setLevel(logging.WARNING)
        snapshot = SimpleNamespace(
            price="399",
            activities=["/"],
            visible_text=(
                "\u6c34\u6708\u96e8\u7fbd\u7ffc Edge ANC "
                "\u4e3b\u52a8\u964d\u566aHiFi\u97f3\u8d28\u771f\u65e0\u7ebf"
                "\u5934\u6234\u5f0f\u84dd\u72595.4\u8033\u673a "
                "\u5e97\u94fa\u4f18\u60e0\u540e \uffe5399"
            ),
        )

        try:
            with patch.object(main, "BrowserService", FakeBrowserService):
                with patch.object(main, "collect_with_retry", return_value=snapshot):
                    with patch.object(main.time, "sleep"):
                        stats = main.run_collection(
                            config={
                                "min_delay_seconds": 0,
                                "max_delay_seconds": 0,
                                "retry_count": 0,
                                "enable_sku_matching": True,
                            },
                            logger=logger,
                            excel=excel,
                            workbook=wb,
                            worksheet=ws,
                            layout=layout,
                            products=products,
                            test_one=False,
                            limit=None,
                            save_every=0,
                            excel_path=None,
                            target_date=date(2026, 7, 6),
                        )
        finally:
            logger.removeHandler(handler)

        self.assertEqual(stats.success, 1)
        self.assertEqual(stats.failed, 0)
        self.assertEqual(ws.cell(row=57, column=3).value, "399")
        self.assertNotIn("\u7591\u4f3c\u4e0b\u67b6", stream.getvalue())

    def test_run_collection_stops_before_starting_next_product(self):
        wb = Workbook()
        ws = wb.active
        layout = SimpleNamespace(price_column=3, activity_column=9)
        products = [
            SimpleNamespace(row=4, brand="A", name="P1", url="https://example.com/1"),
            SimpleNamespace(row=5, brand="A", name="P2", url="https://example.com/2"),
            SimpleNamespace(row=6, brand="A", name="P3", url="https://example.com/3"),
        ]
        excel = FakeExcel(reload_layout=layout)
        logger = logging.getLogger("test_run_collection_stops_before_starting_next_product")
        stop_event = threading.Event()
        progress_events = []

        def progress_callback(event):
            progress_events.append(event)
            if event["processed"] == 1:
                stop_event.set()

        with patch.object(main, "BrowserService", FakeBrowserService):
            with patch.object(main, "collect_with_retry", side_effect=fake_collect_with_retry) as collect_mock:
                with patch.object(main.time, "sleep") as sleep_mock:
                    stats = main.run_collection(
                        config={
                            "min_delay_seconds": 5,
                            "max_delay_seconds": 10,
                            "retry_count": 0,
                            "enable_sku_matching": False,
                        },
                        logger=logger,
                        excel=excel,
                        workbook=wb,
                        worksheet=ws,
                        layout=layout,
                        products=products,
                        test_one=False,
                        limit=None,
                        save_every=0,
                        excel_path=None,
                        stop_event=stop_event,
                        progress_callback=progress_callback,
                    )

        self.assertEqual(stats.total, 3)
        self.assertEqual(stats.success, 1)
        self.assertEqual(collect_mock.call_count, 1)
        self.assertEqual(ws.cell(row=4, column=3).value, "99.00")
        self.assertIsNone(ws.cell(row=5, column=3).value)
        self.assertEqual(progress_events[-1]["status"], "success")
        sleep_mock.assert_not_called()

    def test_run_collection_skips_startup_login_check_when_storage_state_exists(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            storage_state = Path(tmpdir) / "storage_state.json"
            storage_state.write_text('{"cookies":[{"name":"_tb_token_","value":"x","domain":".taobao.com","path":"/"}]}', encoding="utf-8")
            wb = Workbook()
            ws = wb.active
            layout = SimpleNamespace(price_column=3, activity_column=9)
            products = [SimpleNamespace(row=4, brand="A", name="P1", url="https://example.com/1")]
            excel = FakeExcel(reload_layout=layout)
            logger = logging.getLogger("test_run_collection_skips_startup_login_check_when_storage_state_exists")
            instances = []

            class BrowserWithLoginProbe(FakeBrowserService):
                def __init__(self, *args, **kwargs):
                    super().__init__(*args, **kwargs)
                    instances.append(self)

            with patch.object(main, "BrowserService", BrowserWithLoginProbe):
                with patch.object(main, "collect_with_retry", side_effect=fake_collect_with_retry):
                    with patch.object(main.time, "sleep"):
                        stats = main.run_collection(
                            config={
                                "min_delay_seconds": 0,
                                "max_delay_seconds": 0,
                                "retry_count": 0,
                                "enable_sku_matching": False,
                                "browser_storage_state_path": str(storage_state),
                            },
                            logger=logger,
                            excel=excel,
                            workbook=wb,
                            worksheet=ws,
                            layout=layout,
                            products=products,
                            test_one=False,
                            limit=None,
                            save_every=0,
                            excel_path=None,
                        )

        self.assertEqual(stats.success, 1)
        self.assertEqual(len(instances), 1)
        self.assertFalse(instances[0].login_checked)


class WeeklyNewProductModeTest(unittest.TestCase):
    def test_run_weekly_new_products_skips_bi_on_unconfigured_weekday(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            excel_path = Path(tmpdir) / "monitor.xlsx"
            create_weekly_new_workbook(excel_path)
            excel = ExcelService(excel_path)
            workbook = excel.load_workbook()
            logger = logging.getLogger("test_weekly_new_skips_weekday")

            with patch.object(main, "collect_bi_new_products") as collect_bi:
                stats = main.run_weekly_new_products(
                    config={"new_product_enabled": True, "new_product_no_data_text": "无上新"},
                    logger=logger,
                    excel=excel,
                    workbook=workbook,
                    run_date=date(2026, 7, 6),
                    dry_run=False,
                    excel_path=excel_path,
                )

            collect_bi.assert_not_called()
            self.assertEqual(stats, NewProductWriteStats())

    def test_run_weekly_new_products_writes_bi_products_and_no_new_brands(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            excel_path = Path(tmpdir) / "monitor.xlsx"
            create_weekly_new_workbook(excel_path)
            saved = load_workbook(excel_path)
            ws = saved["6.29-7.3"]
            for col in range(5, 8):
                ws.cell(row=4, column=col).value = "199.00"
            saved.save(excel_path)
            excel = ExcelService(excel_path)
            workbook = excel.load_workbook()
            product = NewProduct(
                brand="绿联旗舰店",
                on_sale_date=date(2026, 7, 1),
                model="S6PRO",
                shape="耳夹式",
                price="199",
                link="https://example.com/s6pro",
                selling_point="无",
                source_id="s6pro",
                raw_name="绿联 S6PRO 耳夹式蓝牙耳机",
            )
            unmatched_product = NewProduct(
                brand="十度旗舰店",
                on_sale_date=date(2026, 7, 1),
                model="充电宝",
                shape="未知",
                price="316",
                link="https://example.com/power",
                selling_point="无",
                source_id="power",
                raw_name="十度磁吸充电宝适用苹果17无线超薄小巧",
            )
            result = BICollectResult(products=[product, unmatched_product], total_shop_count=2, queried_shop_count=1, total_goods_count=2)
            logger = logging.getLogger("test_run_weekly_new_products")

            with patch.object(main, "collect_bi_new_products", return_value=result):
                stats = main.run_weekly_new_products(
                    config={
                        "new_product_enabled": True,
                        "new_product_no_data_text": "无上新",
                        "auto_backfill_missing_daily_prices_before_weekly_new": False,
                    },
                    logger=logger,
                    excel=excel,
                    workbook=workbook,
                    run_date=date(2026, 7, 5),
                    dry_run=False,
                    excel_path=excel_path,
                )

            saved = load_workbook(excel_path)
            ws = saved["6.29-7.3"]
            self.assertEqual(stats.written, 1)
            self.assertEqual(stats.failed, 0)
            self.assertEqual(ws["J4"].value, "7月1")
            self.assertEqual(ws["K4"].value, "S6PRO")
            self.assertEqual(ws["K4"].hyperlink.target, "https://example.com/s6pro")
            self.assertEqual(ws["J8"].value, "无上新")


    def test_run_weekly_new_products_backfills_missing_price_dates_before_bi(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            excel_path = Path(tmpdir) / "monitor.xlsx"
            create_weekly_new_workbook(excel_path)
            excel = ExcelService(excel_path)
            workbook = excel.load_workbook()
            result = BICollectResult(products=[], total_shop_count=2, queried_shop_count=2, total_goods_count=0)
            logger = logging.getLogger("test_weekly_new_backfills_missing_price_dates")
            attempted_dates = []

            def fake_backfill(**kwargs):
                missing_date = kwargs["run_date"]
                attempted_dates.append(missing_date)
                saved = load_workbook(excel_path)
                ws = saved["6.29-7.3"]
                column_by_date = {
                    date(2026, 7, 1): 5,
                    date(2026, 7, 2): 6,
                    date(2026, 7, 3): 7,
                }
                ws.cell(row=4, column=column_by_date[missing_date]).value = "299.00"
                saved.save(excel_path)
                return main.RunStats(total=1, success=1)

            with patch.object(main, "run_daily_price_job", side_effect=fake_backfill):
                with patch.object(main, "collect_bi_new_products", return_value=result):
                    main.run_weekly_new_products(
                        config={"new_product_enabled": True, "new_product_no_data_text": "无上新"},
                        logger=logger,
                        excel=excel,
                        workbook=workbook,
                        run_date=date(2026, 7, 5),
                        dry_run=False,
                        excel_path=excel_path,
                    )

            self.assertEqual(
                attempted_dates,
                [date(2026, 7, 1), date(2026, 7, 2), date(2026, 7, 3)],
            )


class PriceTrendModeTest(unittest.TestCase):
    def test_daily_dry_run_reports_target_products_to_progress_callback(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            excel_path = Path(tmpdir) / "monitor.xlsx"
            wb = Workbook()
            ws = wb.active
            create_mature_price_sheet(
                ws,
                "6.29-7.3",
                ["6月29", "6月30", "7月1", "7月2", "7月3"],
                row=4,
                brand="塞那",
                model="S6S proII",
                prices=[None, None, None, None, None],
            )
            wb.save(excel_path)
            excel = ExcelService(excel_path)
            workbook = excel.load_workbook()
            progress_events = []

            stats = main.run_daily_price_job(
                config={},
                logger=logging.getLogger("test_daily_dry_run_progress"),
                excel=excel,
                workbook=workbook,
                run_date=date(2026, 7, 3),
                dry_run=True,
                test_one=True,
                limit=None,
                row_filter=None,
                save_every=0,
                excel_path=excel_path,
                progress_callback=progress_events.append,
            )

            self.assertEqual(stats.total, 1)
            self.assertTrue(progress_events)
            self.assertEqual(progress_events[0]["event"], "phase")
            self.assertEqual(progress_events[0]["total"], 1)
            self.assertIn("1", progress_events[0]["step"])
            product_events = [event for event in progress_events if event.get("event") != "phase"]
            self.assertEqual(len(product_events), 1)
            self.assertEqual(product_events[0]["status"], "skipped")
            self.assertEqual(product_events[0]["error"], "dry-run")
            self.assertEqual(product_events[0]["model"], "S6S proII")

    def test_run_daily_price_job_updates_price_trend_after_friday_collection(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            excel_path = Path(tmpdir) / "monitor.xlsx"
            wb = Workbook()
            ws = wb.active
            create_mature_price_sheet(
                ws,
                "6.29-7.3",
                ["6月29", "6月30", "7月1", "7月2", "7月3"],
                row=4,
                brand="塞那",
                model="S6S proII",
                prices=[100, 99, 98, 98, 97],
            )
            wb.save(excel_path)
            excel = ExcelService(excel_path)
            workbook = excel.load_workbook()
            logger = logging.getLogger("test_daily_updates_price_trend")

            with patch.object(main, "run_collection", return_value=main.RunStats(total=1, success=1)):
                main.run_daily_price_job(
                    config={"price_trend_enabled": True},
                    logger=logger,
                    excel=excel,
                    workbook=workbook,
                    run_date=date(2026, 7, 3),
                    dry_run=False,
                    test_one=False,
                    limit=None,
                    row_filter=None,
                    save_every=0,
                    excel_path=excel_path,
                )

            saved = load_workbook(excel_path)
            self.assertEqual(saved["6.29-7.3"]["H4"].value, "价格下降")

    def test_run_price_trend_job_can_fill_previous_week_on_weekend(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            excel_path = Path(tmpdir) / "monitor.xlsx"
            wb = Workbook()
            ws = wb.active
            create_mature_price_sheet(
                ws,
                "6.29-7.3",
                ["6月29", "6月30", "7月1", "7月2", "7月3"],
                row=4,
                brand="塞那",
                model="S6S proII",
                prices=[100, 95, 95, 98, 101],
            )
            wb.save(excel_path)
            excel = ExcelService(excel_path)
            workbook = excel.load_workbook()
            logger = logging.getLogger("test_price_trend_weekend")

            stats = main.run_price_trend_job(
                config={"price_trend_enabled": True},
                logger=logger,
                excel=excel,
                workbook=workbook,
                run_date=date(2026, 7, 5),
                dry_run=False,
                excel_path=excel_path,
            )

            saved = load_workbook(excel_path)
            self.assertEqual(stats.written, 1)
            self.assertEqual(saved["6.29-7.3"]["H4"].value, "先降后升")


class FakeExcel:
    def __init__(self, sheet_title="Sheet", reload_layout=None, fail_on_repeated_workbook_save=False):
        self.saved_paths = []
        self.saved_workbook_ids = []
        self.sheet_title = sheet_title
        self.reload_layout = reload_layout
        self.fail_on_repeated_workbook_save = fail_on_repeated_workbook_save

    def write_product_result(self, worksheet, layout, product, price, activity_value, force_overwrite=False):
        worksheet.cell(row=product.row, column=layout.price_column).value = price
        worksheet.cell(row=product.row, column=layout.activity_column).value = activity_value
        return True

    def save_workbook(self, workbook, excel_path):
        workbook_id = id(workbook)
        if self.fail_on_repeated_workbook_save and workbook_id in self.saved_workbook_ids:
            raise ValueError("same workbook object saved twice")
        self.saved_workbook_ids.append(workbook_id)
        self.saved_paths.append(excel_path)

    def load_workbook(self):
        wb = Workbook()
        ws = wb.active
        ws.title = self.sheet_title
        return wb

    def detect_layout(self, worksheet, target_date):
        return self.reload_layout

    def list_mature_products(self, worksheet, model_column, data_start_row):
        products = []
        current_brand = None
        for row in range(data_start_row, worksheet.max_row + 1):
            brand = worksheet.cell(row=row, column=1).value or current_brand
            if worksheet.cell(row=row, column=1).value not in (None, ""):
                current_brand = worksheet.cell(row=row, column=1).value
            name = worksheet.cell(row=row, column=model_column).value
            if name in (None, ""):
                continue
            products.append(SimpleNamespace(row=row, brand=brand, name=name, url=""))
        return products


class FakeBrowserService:
    def __init__(self, *args, **kwargs):
        self.login_checked = False
        pass

    def __enter__(self):
        return self

    def __exit__(self, exc_type, exc, tb):
        return False

    def ensure_taobao_login(self):
        self.login_checked = True


def fake_collect_with_retry(browser, collector, url, retry_count, expected_model=None, allow_fallback_price=True):
    return SimpleNamespace(price="99.00", activities="国补")


class RetryBrowser:
    def __init__(self, failures_before_success):
        self.failures_before_success = failures_before_success
        self.open_count = 0

    def open_page(self, url):
        self.open_count += 1
        if self.open_count <= self.failures_before_success:
            raise TimeoutError("timeout")
        return RetryPage()


class RetryCollector:
    def __init__(self):
        self.collect_count = 0

    def collect_from_page(self, page, expected_model=None, allow_fallback_price=True):
        self.collect_count += 1
        return SimpleNamespace(price="88.00", activities="国补")


class RetryPage:
    def close(self):
        pass


def create_weekly_new_workbook(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "6.29-7.3"
    ws["A1"] = "品牌"
    ws["B2"] = "型号"
    ws["C2"] = "价格"
    ws["H3"] = "价格情况"
    ws["I2"] = "活动"
    for cell, value in zip(["C3", "D3", "E3", "F3", "G3"], ["6月29", "6月30", "7月1", "7月2", "7月3"]):
        ws[cell] = value
    ws["J1"] = "上新监控"
    ws.merge_cells("J1:S1")
    ws["J2"] = "上架日期"
    ws.merge_cells("J2:J3")
    ws["K2"] = "型号"
    ws.merge_cells("K2:K3")
    ws["L2"] = "形态"
    ws.merge_cells("L2:L3")
    ws["M2"] = "价格"
    ws.merge_cells("M2:M3")
    ws["N2"] = "直播价格"
    ws.merge_cells("N2:R2")
    ws["S2"] = "新卖点"
    ws.merge_cells("S2:S3")
    for cell, value in zip(["N3", "O3", "P3", "Q3", "R3"], ["6月29", "6月30", "7月1", "7月2", "7月3"]):
        ws[cell] = value
    ws["A4"] = "绿联"
    ws["B4"] = "S6S proII"
    ws["B4"].hyperlink = "https://example.com/item"
    ws["C4"] = "199.00"
    ws["D4"] = "199.00"
    ws.merge_cells("A4:A7")
    ws.merge_cells("J4:S7")
    ws["J4"] = "无上新"
    ws["A8"] = "华为"
    ws.merge_cells("A8:A10")
    ws.merge_cells("J8:S10")
    ws["J8"] = "无上新"
    wb.save(path)


def create_mature_price_sheet(worksheet, title, date_headers, row, brand, model, prices):
    worksheet.title = title
    worksheet["A1"] = "品牌"
    worksheet["B2"] = "型号"
    worksheet["C2"] = "价格"
    worksheet["H3"] = "价格情况"
    worksheet["I2"] = "活动"
    for column_offset, header in enumerate(date_headers, start=3):
        worksheet.cell(row=3, column=column_offset).value = header
    worksheet.cell(row=row, column=1).value = brand
    model_cell = worksheet.cell(row=row, column=2)
    model_cell.value = model
    model_cell.hyperlink = "https://example.com/item"
    for column_offset, price in enumerate(prices, start=3):
        worksheet.cell(row=row, column=column_offset).value = price


if __name__ == "__main__":
    unittest.main()
