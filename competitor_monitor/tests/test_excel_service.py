import tempfile
import unittest
import xml.etree.ElementTree as ET
import shutil
from copy import copy
from datetime import date
from pathlib import Path
from unittest.mock import patch
from zipfile import ZIP_DEFLATED, ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Border, Side

from excel_service import ExcelService, WorkbookLayoutError, WorkbookLockedError, backup_excel


def create_sample_workbook(path: Path) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "6.29-7.3"
    ws["A1"] = "品牌"
    ws["B1"] = "竞争品牌核心单品监控【成熟单品：销量高&销售金额高】"
    ws["J1"] = "上新监控"
    ws["B2"] = "型号"
    ws["C2"] = "价格"
    ws["I2"] = "活动"
    ws["J2"] = "上架日期"
    ws["K2"] = "型号"
    for cell, value in zip(["C3", "D3", "E3", "F3", "G3"], ["6月29", "6月30", "7月1", "7月2", "7月3"]):
        ws[cell] = value
    ws["H3"] = "价格情况"
    ws["A4"] = "塞那"
    ws["B4"] = "S6S proII"
    ws["B4"].hyperlink = "https://detail.tmall.com/item.htm?id=800417757444"
    ws["C4"] = "268.52"
    ws["I4"] = "国补"
    ws["J4"] = "无上新"
    ws["A5"] = None
    ws["B5"] = "S6S ultra"
    ws["B5"].hyperlink = "https://detail.tmall.com/item.htm?id=922811879396"
    ws["K5"] = "右侧上新内容"
    wb.save(path)


class ExcelServiceTest(unittest.TestCase):
    def test_finds_sheet_headers_and_mature_product_links(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            excel_path = Path(tmpdir) / "竞品监控.xlsx"
            create_sample_workbook(excel_path)
            service = ExcelService(excel_path)
            wb = service.load_workbook()

            ws = service.get_or_create_current_sheet(wb, date(2026, 7, 1))
            layout = service.detect_layout(ws, date(2026, 7, 1))
            products = service.list_mature_products(ws)

            self.assertEqual(ws.title, "6.29-7.3")
            self.assertEqual(layout.price_column, 5)
            self.assertEqual(layout.activity_column, 9)
            self.assertEqual(layout.price_status_column, 8)
            self.assertEqual(len(products), 2)
            self.assertEqual(products[0].row, 4)
            self.assertEqual(products[0].brand, "塞那")
            self.assertEqual(products[0].name, "S6S proII")
            self.assertIn("id=800417757444", products[0].url)

    def test_creates_current_sheet_from_latest_template_and_clears_runtime_cells(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            excel_path = Path(tmpdir) / "竞品监控.xlsx"
            create_sample_workbook(excel_path)
            service = ExcelService(excel_path)
            wb = service.load_workbook()

            ws = service.get_or_create_current_sheet(wb, date(2026, 7, 6))

            self.assertEqual(ws.title, "7.6-7.10")
            self.assertEqual([ws[cell].value for cell in ["C3", "D3", "E3", "F3", "G3"]], ["7月6", "7月7", "7月8", "7月9", "7月10"])
            self.assertEqual(ws["B4"].value, "S6S proII")
            self.assertIn("id=800417757444", ws["B4"].hyperlink.target)
            self.assertIsNone(ws["C4"].value)
            self.assertIsNone(ws["I4"].value)
            self.assertIsNone(ws["J4"].value)
            self.assertIsNone(ws["K5"].value)

    def test_refuses_to_create_sheet_for_weekend_date(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            excel_path = Path(tmpdir) / "竞品监控.xlsx"
            create_sample_workbook(excel_path)
            service = ExcelService(excel_path)
            wb = service.load_workbook()

            with self.assertRaises(WorkbookLayoutError) as ctx:
                service.get_or_create_current_sheet(wb, date(2026, 7, 5))

            self.assertIn("Monday-Friday", str(ctx.exception))
            self.assertNotIn("7.6-7.10", wb.sheetnames)

    def test_check_weekly_price_completeness_reports_blank_price_dates(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            excel_path = Path(tmpdir) / "竞品监控.xlsx"
            create_sample_workbook(excel_path)
            service = ExcelService(excel_path)
            wb = service.load_workbook()

            completeness = service.check_weekly_price_completeness(wb, date(2026, 7, 3))

            self.assertEqual(completeness.sheet_name, "6.29-7.3")
            self.assertFalse(completeness.is_complete)
            self.assertEqual(completeness.days[0].status, "partial")
            self.assertEqual(completeness.days[0].filled_count, 1)
            self.assertEqual(completeness.days[0].total_count, 2)
            self.assertEqual(
                completeness.missing_dates,
                (
                    date(2026, 6, 29),
                    date(2026, 6, 30),
                    date(2026, 7, 1),
                    date(2026, 7, 2),
                    date(2026, 7, 3),
                ),
            )

    def test_check_weekly_price_completeness_passes_when_all_weekdays_are_filled(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            excel_path = Path(tmpdir) / "竞品监控.xlsx"
            create_sample_workbook(excel_path)
            service = ExcelService(excel_path)
            wb = service.load_workbook()
            ws = wb["6.29-7.3"]
            for row in (4, 5):
                for col in range(3, 8):
                    ws.cell(row=row, column=col).value = "199.00"

            completeness = service.check_weekly_price_completeness(wb, date(2026, 7, 3))

            self.assertTrue(completeness.is_complete)
            self.assertEqual(completeness.missing_dates, ())
            self.assertEqual({day.status for day in completeness.days}, {"complete"})

    def test_check_weekly_price_completeness_reports_missing_sheet_without_creating_it(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            excel_path = Path(tmpdir) / "竞品监控.xlsx"
            create_sample_workbook(excel_path)
            service = ExcelService(excel_path)
            wb = service.load_workbook()

            completeness = service.check_weekly_price_completeness(wb, date(2026, 7, 6))

            self.assertIsNone(completeness.sheet_name)
            self.assertFalse(completeness.is_complete)
            self.assertEqual(completeness.days[0].status, "missing_sheet")
            self.assertEqual(completeness.missing_dates[0], date(2026, 7, 6))
            self.assertNotIn("7.6-7.10", wb.sheetnames)

    def test_backup_excel_copies_file_to_timestamped_backup_dir(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            excel_path = Path(tmpdir) / "竞品监控.xlsx"
            backup_dir = Path(tmpdir) / "backup"
            create_sample_workbook(excel_path)

            backup_path = backup_excel(excel_path, backup_dir)

            self.assertTrue(backup_path.exists())
            self.assertTrue(backup_path.name.startswith("竞品监控_"))
            self.assertEqual(load_workbook(backup_path).sheetnames, ["6.29-7.3"])

    def test_detect_layout_raises_clear_error_when_activity_header_missing(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            excel_path = Path(tmpdir) / "竞品监控.xlsx"
            create_sample_workbook(excel_path)
            wb = load_workbook(excel_path)
            ws = wb["6.29-7.3"]
            ws["I2"] = None
            wb.save(excel_path)
            service = ExcelService(excel_path)
            wb = service.load_workbook()

            with self.assertRaises(WorkbookLayoutError) as ctx:
                service.detect_layout(wb["6.29-7.3"], date(2026, 7, 1))

            self.assertIn("活动", str(ctx.exception))

    def test_write_product_result_preserves_existing_styles_and_links(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            excel_path = Path(tmpdir) / "竞品监控.xlsx"
            create_sample_workbook(excel_path)
            service = ExcelService(excel_path)
            wb = service.load_workbook()
            ws = wb["6.29-7.3"]
            layout = service.detect_layout(ws, date(2026, 7, 1))
            product = service.list_mature_products(ws)[0]
            price_cell = ws.cell(row=product.row, column=layout.price_column)
            activity_cell = ws.cell(row=product.row, column=layout.activity_column)
            price_style = copy(price_cell._style)
            activity_style = copy(activity_cell._style)
            model_link = ws.cell(row=product.row, column=layout.model_column).hyperlink.target

            service.write_product_result(
                ws,
                layout,
                product,
                "299.00",
                "国补、超级立减",
                force_overwrite=True,
            )

            self.assertEqual(copy(price_cell._style), price_style)
            self.assertEqual(copy(activity_cell._style), activity_style)
            self.assertEqual(ws.cell(row=product.row, column=layout.model_column).hyperlink.target, model_link)

    def test_save_workbook_keeps_existing_file_when_save_fails(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            excel_path = Path(tmpdir) / "monitor.xlsx"
            excel_path.write_bytes(b"original workbook bytes")
            service = ExcelService(excel_path)

            with self.assertRaises(RuntimeError):
                service.save_workbook(FailingWorkbook(), excel_path)

            self.assertEqual(excel_path.read_bytes(), b"original workbook bytes")

    def test_save_workbook_uses_fallback_when_target_is_locked(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            excel_path = Path(tmpdir) / "monitor.xlsx"
            fallback_path = Path(tmpdir) / "output" / "latest" / "monitor.xlsx"
            excel_path.write_bytes(b"original workbook bytes")
            fallback_calls = []

            def fallback(temp_path, target_path, error):
                fallback_path.parent.mkdir(parents=True)
                shutil.copy2(temp_path, fallback_path)
                fallback_calls.append((Path(target_path), str(error)))
                return fallback_path

            service = ExcelService(excel_path, save_fallback=fallback)
            with patch.object(
                ExcelService,
                "_replace_file_with_retry",
                side_effect=WorkbookLockedError("locked"),
            ) as replace_mock:
                actual_path = service.save_workbook(ByteWorkbook(b"updated workbook bytes"), excel_path)

            self.assertEqual(actual_path, fallback_path)
            self.assertEqual(service.excel_path, fallback_path)
            self.assertEqual(excel_path.read_bytes(), b"original workbook bytes")
            self.assertEqual(fallback_path.read_bytes(), b"updated workbook bytes")
            self.assertEqual(fallback_calls[0][0], excel_path)
            self.assertEqual(replace_mock.call_args.kwargs["attempts"], 3)

    def test_save_workbook_preserves_empty_sheet_tab_color_nodes(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            excel_path = Path(tmpdir) / "monitor.xlsx"
            create_sample_workbook(excel_path)
            rewrite_xlsx_member(excel_path, "xl/worksheets/sheet1.xml", add_empty_tab_color)
            service = ExcelService(excel_path)
            wb = service.load_workbook()

            wb["6.29-7.3"]["C4"] = "changed"
            service.save_workbook(wb, excel_path)

            with ZipFile(excel_path, "r") as workbook_zip:
                root = ET.fromstring(workbook_zip.read("xl/worksheets/sheet1.xml"))
            ns_uri = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
            tab_color = root.find(f"{{{ns_uri}}}sheetPr/{{{ns_uri}}}tabColor")
            self.assertIsNotNone(tab_color)
            self.assertEqual(tab_color.attrib, {})

    def test_save_workbook_restores_original_style_table(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            excel_path = Path(tmpdir) / "monitor.xlsx"
            create_sample_workbook(excel_path)
            wb = load_workbook(excel_path)
            wb["6.29-7.3"]["B4"].hyperlink = None
            wb["6.29-7.3"]["B5"].hyperlink = None
            wb.save(excel_path)
            original_styles = read_xlsx_member(excel_path, "xl/styles.xml")
            service = ExcelService(excel_path)
            wb = service.load_workbook()

            wb["6.29-7.3"]["C4"] = "changed"
            service.save_workbook(wb, excel_path)

            self.assertEqual(read_xlsx_member(excel_path, "xl/styles.xml"), original_styles)
            self.assertEqual(load_workbook(excel_path)["6.29-7.3"]["C4"].value, "changed")

    def test_save_workbook_preserves_style_changes_on_changed_existing_cells(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            excel_path = Path(tmpdir) / "monitor.xlsx"
            create_sample_workbook(excel_path)
            wb = load_workbook(excel_path)
            ws = wb["6.29-7.3"]
            thin = Side(style="thin")
            ws["K4"].border = Border(left=thin, right=thin, top=thin, bottom=thin)
            wb.save(excel_path)

            service = ExcelService(excel_path)
            wb = service.load_workbook()
            ws = wb["6.29-7.3"]
            cell = ws["K4"]
            original_border = copy(cell.border)
            cell.value = "H9"
            cell.hyperlink = "https://example.com/h9"
            font = copy(cell.font)
            font.color = "0563C1"
            font.underline = "single"
            cell.font = font

            service.save_workbook(wb, excel_path)

            saved_cell = load_workbook(excel_path)["6.29-7.3"]["K4"]
            self.assertEqual(saved_cell.value, "H9")
            self.assertEqual(saved_cell.hyperlink.target, "https://example.com/h9")
            self.assertEqual(saved_cell.font.underline, "single")
            self.assertEqual(saved_cell.font.color.rgb, "000563C1")
            self.assertEqual(saved_cell.border.top.style, original_border.top.style)

    def test_save_workbook_styles_mature_product_hyperlinks_without_losing_borders(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            excel_path = Path(tmpdir) / "monitor.xlsx"
            create_sample_workbook(excel_path)
            service = ExcelService(excel_path)
            wb = service.load_workbook()
            ws = wb["6.29-7.3"]
            original_border = copy(ws["B4"].border)

            ws["C4"] = "changed"
            service.save_workbook(wb, excel_path)

            saved_cell = load_workbook(excel_path)["6.29-7.3"]["B4"]
            self.assertEqual(saved_cell.hyperlink.target, "https://detail.tmall.com/item.htm?id=800417757444")
            self.assertEqual(saved_cell.font.underline, "single")
            self.assertIn(saved_cell.font.color.rgb, {"000563C1", "FF0563C1"})
            self.assertEqual(saved_cell.border.left.style, original_border.left.style)
            self.assertEqual(saved_cell.border.top.style, original_border.top.style)
            self.assertEqual(saved_cell.border.bottom.style, original_border.bottom.style)

    def test_save_workbook_preserves_existing_drawing_package_parts(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            excel_path = Path(tmpdir) / "monitor.xlsx"
            create_sample_workbook(excel_path)
            inject_fake_drawing(excel_path)
            service = ExcelService(excel_path)
            wb = service.load_workbook()

            wb["6.29-7.3"]["H4"] = "价格稳定"
            service.save_workbook(wb, excel_path)

            with ZipFile(excel_path, "r") as workbook_zip:
                names = set(workbook_zip.namelist())
                self.assertIn("xl/drawings/drawing1.xml", names)
                self.assertIn("xl/drawings/_rels/drawing1.xml.rels", names)
                self.assertIn("xl/media/image1.png", names)
                self.assertIn(b"drawing1.xml", workbook_zip.read("xl/worksheets/_rels/sheet1.xml.rels"))
                self.assertIn(b"<drawing", workbook_zip.read("xl/worksheets/sheet1.xml"))
                self.assertIn(b"/xl/drawings/drawing1.xml", workbook_zip.read("[Content_Types].xml"))
                self.assertEqual(workbook_zip.read("xl/media/image1.png"), TINY_PNG)
            self.assertEqual(load_workbook(excel_path)["6.29-7.3"]["H4"].value, "价格稳定")

    def test_replace_file_retries_transient_windows_permission_error(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            source = Path(tmpdir) / "source.txt"
            target = Path(tmpdir) / "target.txt"
            source.write_text("new", encoding="utf-8")
            target.write_text("old", encoding="utf-8")
            original_replace = Path.replace
            attempts = {"count": 0}

            def flaky_replace(path, destination):
                if path == source and attempts["count"] < 6:
                    attempts["count"] += 1
                    raise PermissionError("simulated transient lock")
                return original_replace(path, destination)

            with patch.object(Path, "replace", flaky_replace), patch("excel_service.time.sleep", lambda _seconds: None):
                ExcelService._replace_file_with_retry(source, target)

            self.assertEqual(target.read_text(encoding="utf-8"), "new")
            self.assertEqual(attempts["count"], 6)

    def test_assert_workbook_writable_reports_locked_excel_clearly(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            excel_path = Path(tmpdir) / "monitor.xlsx"
            excel_path.write_bytes(b"workbook bytes")

            with patch.object(ExcelService, "_open_file_exclusive", side_effect=PermissionError("locked")):
                with self.assertRaises(WorkbookLockedError) as ctx:
                    ExcelService.assert_workbook_writable(excel_path)

            self.assertIn("当前 Excel 文件被占用", str(ctx.exception))
            self.assertIn("Excel/WPS", str(ctx.exception))

    def test_wait_workbook_writable_retries_until_transient_lock_is_released(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            excel_path = Path(tmpdir) / "monitor.xlsx"
            excel_path.write_bytes(b"workbook bytes")
            calls = {"count": 0}
            sleeps: list[float] = []
            wait_events: list[tuple[int, int]] = []

            def flaky_check(path):
                calls["count"] += 1
                if calls["count"] <= 3:
                    raise WorkbookLockedError("simulated transient lock")

            with patch.object(ExcelService, "assert_workbook_writable", flaky_check), patch(
                "excel_service.time.sleep", lambda seconds: sleeps.append(seconds)
            ):
                ExcelService.wait_workbook_writable(
                    excel_path,
                    timeout_seconds=5,
                    delay_seconds=1,
                    on_wait=lambda attempt, remaining: wait_events.append((attempt, remaining)),
                )

            self.assertEqual(calls["count"], 4)
            self.assertEqual(sleeps, [1, 1, 1])
            self.assertEqual([event[0] for event in wait_events], [1, 2, 3])

    def test_wait_workbook_writable_raises_clear_error_after_timeout(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            excel_path = Path(tmpdir) / "monitor.xlsx"
            excel_path.write_bytes(b"workbook bytes")

            with patch.object(
                ExcelService,
                "assert_workbook_writable",
                side_effect=WorkbookLockedError("当前 Excel 文件被占用，请关闭 Excel/WPS/预览窗口后再运行：monitor.xlsx"),
            ), patch("excel_service.time.sleep", lambda _seconds: None):
                with self.assertRaises(WorkbookLockedError) as ctx:
                    ExcelService.wait_workbook_writable(excel_path, timeout_seconds=2, delay_seconds=1)

            self.assertIn("当前 Excel 文件被占用", str(ctx.exception))


class FailingWorkbook:
    def save(self, target):
        Path(target).write_bytes(b"partial workbook bytes")
        raise RuntimeError("simulated save failure")


class ByteWorkbook:
    def __init__(self, payload: bytes):
        self.payload = payload

    def save(self, target):
        Path(target).write_bytes(self.payload)


def rewrite_xlsx_member(path: Path, member_name: str, transform):
    rewritten_path = path.with_suffix(".rewritten.xlsx")
    with ZipFile(path, "r") as source, ZipFile(rewritten_path, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == member_name:
                data = transform(data)
            target.writestr(item, data)
    rewritten_path.replace(path)


def read_xlsx_member(path: Path, member_name: str) -> bytes:
    with ZipFile(path, "r") as workbook_zip:
        return workbook_zip.read(member_name)


TINY_PNG = (
    b"\x89PNG\r\n\x1a\n\x00\x00\x00\rIHDR\x00\x00\x00\x01\x00\x00\x00\x01"
    b"\x08\x06\x00\x00\x00\x1f\x15\xc4\x89\x00\x00\x00\nIDATx\x9cc\x00\x01"
    b"\x00\x00\x05\x00\x01\r\n-\xb4\x00\x00\x00\x00IEND\xaeB`\x82"
)


def inject_fake_drawing(path: Path) -> None:
    rel_ns = "http://schemas.openxmlformats.org/package/2006/relationships"
    spreadsheet_ns = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    office_rel_ns = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
    content_ns = "http://schemas.openxmlformats.org/package/2006/content-types"
    drawing_xml = (
        b'<wsDr xmlns="http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing">'
        b"<oneCellAnchor><from><col>18</col><colOff>0</colOff><row>78</row><rowOff>813948</rowOff></from>"
        b'<ext cx="581025" cy="239003"/><pic><nvPicPr><cNvPr id="1" name="1"/><cNvPicPr/></nvPicPr>'
        b'<blipFill><a:blip xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main" '
        b'xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships" '
        b'r:embed="rId1" r:link="rId0"/><a:stretch xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
        b"<a:fillRect/></a:stretch></blipFill><spPr/></pic><clientData/></oneCellAnchor></wsDr>"
    )
    rewritten_path = path.with_suffix(".with-drawing.xlsx")
    with ZipFile(path, "r") as source, ZipFile(rewritten_path, "w", ZIP_DEFLATED) as target:
        for item in source.infolist():
            data = source.read(item.filename)
            if item.filename == "xl/worksheets/sheet1.xml":
                root = ET.fromstring(data)
                ET.register_namespace("r", office_rel_ns)
                drawing = ET.Element(f"{{{spreadsheet_ns}}}drawing")
                drawing.attrib[f"{{{office_rel_ns}}}id"] = "rId99"
                root.append(drawing)
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            elif item.filename == "xl/worksheets/_rels/sheet1.xml.rels":
                root = ET.fromstring(data)
                relationship = ET.Element(f"{{{rel_ns}}}Relationship")
                relationship.attrib.update(
                    {
                        "Id": "rId99",
                        "Type": "http://schemas.openxmlformats.org/officeDocument/2006/relationships/drawing",
                        "Target": "/xl/drawings/drawing1.xml",
                    }
                )
                root.append(relationship)
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            elif item.filename == "[Content_Types].xml":
                root = ET.fromstring(data)
                override = ET.Element(f"{{{content_ns}}}Override")
                override.attrib.update(
                    {
                        "PartName": "/xl/drawings/drawing1.xml",
                        "ContentType": "application/vnd.openxmlformats-officedocument.drawing+xml",
                    }
                )
                root.append(override)
                if not any(child.attrib.get("Extension") == "png" for child in root):
                    default = ET.Element(f"{{{content_ns}}}Default")
                    default.attrib.update({"Extension": "png", "ContentType": "image/png"})
                    root.append(default)
                data = ET.tostring(root, encoding="utf-8", xml_declaration=True)
            target.writestr(item, data)
        target.writestr("xl/drawings/drawing1.xml", drawing_xml)
        target.writestr(
            "xl/drawings/_rels/drawing1.xml.rels",
            b'<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">'
            b'<Relationship Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/image" '
            b'Target="/xl/media/image1.png" Id="rId1"/></Relationships>',
        )
        target.writestr("xl/media/image1.png", TINY_PNG)
    rewritten_path.replace(path)


def add_empty_tab_color(data: bytes) -> bytes:
    ns_uri = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
    ET.register_namespace("", ns_uri)
    root = ET.fromstring(data)
    sheet_pr = root.find(f"{{{ns_uri}}}sheetPr")
    if sheet_pr is None:
        sheet_pr = ET.Element(f"{{{ns_uri}}}sheetPr")
        root.insert(0, sheet_pr)
    for tab_color in list(sheet_pr.findall(f"{{{ns_uri}}}tabColor")):
        sheet_pr.remove(tab_color)
    ET.SubElement(sheet_pr, f"{{{ns_uri}}}tabColor")
    return ET.tostring(root, encoding="utf-8", xml_declaration=True)


if __name__ == "__main__":
    unittest.main()
